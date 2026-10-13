# -*- coding: utf-8 -*-
"""Etapa B3 de Articulo 6, posterior a B1 (ECU_NDC, rama art6-escenarios).

B1 genera los escenarios desde B1_Scenario_Config.xlsx con su propio mecanismo: referencias (Modelo CLG y base
corregida congelada, Based_On = ref) y escenarios basados en CONA6 (A, B, C, estados, referencias de estado y
puentes V7). El motor lee cada escenario de Executables/<ESC>_0 y no necesita B2: ECUOutput.csv y ECUInput.csv son
archivos del Modelo CLG versionados en git y no se reescriben.

posterior(): informa el estado de glpsol por escenario, verifica cada escenario contra su base,
    escribe B2_Output_Art6/MANIFIESTO.csv, analiza volumenes, margen, sobreventa e ingresos, genera
    las graficas y el tablero, escribe los fragmentos de la documentacion, ejecuta
    B3_Visualizador_Art6.ipynb y compila la documentacion con Sphinx si esta instalado. Una falla de
    verificacion se informa en la terminal y en verificacion.txt y no detiene la etapa.
"""
import csv
import io
import shutil
import subprocess
import sys
import time
from pathlib import Path

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))

import analizar  # noqa: E402
import comun as C  # noqa: E402
import configuracion as K  # noqa: E402
import verificar  # noqa: E402

CUADERNO = C.ENERGY / "B3_Visualizador_Art6.ipynb"


def manifiesto():
    """B2_Output_Art6/MANIFIESTO.csv: datafile, md5 y composicion de cada escenario."""
    filas, bases = [], {}
    import pandas as pd
    ok = [e["codigo"] for e in K.escenarios() if C.estado_solucion(e["codigo"]) == "optimo"]
    salida = C.leer_ecuoutput(ok) if ok else None
    for e in K.escenarios():
        cod = e["codigo"]
        if not C.datafile(cod).exists():
            continue
        if e["base"] == "B1":
            n = 0
        else:
            if e["base"] not in bases and C.datafile(e["base"]).exists():
                bases[e["base"]] = C.leer_datafile(C.datafile(e["base"]))
            n = len(C.diferencias(bases[e["base"]], C.leer_datafile(C.datafile(cod)))) if e["base"] in bases else ""
        filas.append({"codigo": cod, "nombre": e["nombre"], "grupo": e["grupo"],
                      "datafile": "Executables/%s_0/%s_0.txt" % (cod, cod), "md5": C.md5(C.datafile(cod)),
                      "huella_datafile": C.huella_datafile(C.datafile(cod)),
                      "huella_resultados": C.huella_resultados(salida, cod) if cod in ok else "sin solución",
                      "output_md5": C.md5(C.salida(cod)) if cod in ok else "",
                      "base": e["base"], "palancas": e["palancas"] or "", "filas_distintas_de_la_base": n})
    C.RESULTADOS.mkdir(parents=True, exist_ok=True)
    if not filas:
        return
    with io.open(C.RESULTADOS / "MANIFIESTO.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(filas[0].keys()))
        w.writeheader()
        w.writerows(filas)


def ejecutar_cuaderno():
    try:
        import nbformat
        from nbclient import NotebookClient
    except ImportError:
        print("  falta nbclient: abra B3_Visualizador_Art6.ipynb y ejecute todas las celdas")
        return
    nb = nbformat.read(str(CUADERNO), as_version=4)
    try:
        NotebookClient(nb, timeout=900, kernel_name="python3",
                       resources={"metadata": {"path": str(C.ENERGY)}}).execute()
    except Exception as ex:          # el cuaderno no detiene la etapa: se informa
        print("  el cuaderno no se ejecutó completo: %s" % str(ex).splitlines()[0][:200])
        return
    nbformat.write(nb, str(CUADERNO))
    print("  gráficas listas en %s" % CUADERNO.name)


def _rst(texto):
    """Texto plano seguro para reStructuredText: escapa los guiones bajos finales (referencias)."""
    import re
    return re.sub(r"(\w)_(?=\W|$)", r"\1\\_", str(texto or ""))


def _figuras_rst(F, L, gen):
    """_generado/fig_<gid>.rst: cada grafica como figura numerada (PNG de la corrida), con el titulo y
    el subtitulo calculados, el enlace a su version interactiva y su guia de lectura."""
    import re
    import graficar
    import tablero
    for gid, _, _ in graficar.FIGURAS:
        if gid not in F:
            continue
        cabeza, sub, _ = tablero._separar_titulo(F[gid])
        limpio = lambda s: re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", s)).strip()
        partes = [".. figure:: /_static/graficas/%s.png" % gid, "   :name: fig-%s" % gid, "   :width: 100%",
                  "   :alt: %s" % _rst(limpio(cabeza)), "", "   %s" % _rst(limpio(cabeza)), ""]
        if sub:
            partes += ["   %s" % _rst(limpio(sub)), ""]
        partes += [".. raw:: html", "",
                   "   <p class='interactiva'><a data-grafica='%s' href='_static/interactivas/%s.html' "
                   "target='_blank' rel='noopener'>Abrir la versión interactiva</a></p>" % (gid, gid), ""]
        if gid in L:
            import guias
            for et, cuerpo in guias.partes(L[gid]):
                partes += ([".. rubric:: %s" % et, ""] if et else []) + [_rst(cuerpo) + "\n"]
        io.open(gen / ("fig_%s.rst" % gid), "w", encoding="utf-8").write("\n".join(partes) + "\n")


def _escenarios_html(gen):
    """Mapa, fichas y matriz de supuestos de los escenarios, para incluirlos con raw html."""
    import lectura
    import explicacion
    # cada fragmento lleva su hoja de estilos, porque la documentación lo inserta solo en su página
    io.open(gen / "lectura.html", "w", encoding="utf-8").write(
        explicacion._estilo(False) + lectura.html(fuentes=False, glosario=False))
    io.open(gen / "glosario.html", "w", encoding="utf-8").write(explicacion._estilo(False) + lectura.glosario_html())
    for nombre, f in (("mapa", explicacion.mapa_html), ("fichas", explicacion.fichas_html),
                      ("supuestos", explicacion.supuestos_html)):
        io.open(gen / ("escenarios_%s.html" % nombre), "w", encoding="utf-8").write(f(fuentes=False))


def _hoja_xlsx(ruta, hoja):
    import openpyxl
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    filas = list(wb[hoja].iter_rows(values_only=True))
    wb.close()
    return [dict(zip(filas[0], f)) for f in filas[1:]]


def _estructura(gen):
    """Tecnologias, combustibles y emisiones del modelo (conjuntos de B1_Model_Structure.xlsx) con
    sus nombres de los clasificadores y la parametrizacion de A1."""
    import openpyxl
    wb = openpyxl.load_workbook(C.ENERGY / "B1_Model_Structure.xlsx", read_only=True, data_only=True)
    filas = list(wb["sets"].iter_rows(values_only=True))
    wb.close()
    cab = filas[0]
    conj = {n: [f[i] for f in filas[3:] if f[i] is not None] for i, n in enumerate(cab) if n}
    a1i, a1o = C.ENERGY / "A1_Inputs", C.ENERGY / "A1_Outputs"
    tec, comb = {}, {}
    for hoja in ("PrimaryEnergy", "SecondaryEnergy"):
        for d in _hoja_xlsx(a1i / "A-I_Classifier_Modes_Supply.xlsx", hoja):
            if d.get("Tech - Code") not in (None, "Not required", "TEST ROW"):
                tec[d["Tech - Code"]] = d.get("Tech - Plain Spanish")
            for lado in ("Output", "Input"):
                cod = d.get("Fuel - Code (%s)" % lado)
                if cod and cod != "TEST ROW":
                    comb.setdefault(cod, d.get("Fuel - Plain Spanish (%s)" % lado))
    nombres_hoja = {"Primary Techs": "Primaria", "Secondary Techs": "Secundaria", "Demand Techs": "Demanda",
                    "Transport Fuel Distribution": "Distribución para transporte", "Vehicle Techs": "Vehículos",
                    "Vehicle Groups": "Grupos de vehículos", "Other_Techs": "Otras"}
    par = openpyxl.load_workbook(a1o / "A-O_Parametrization.xlsx", read_only=True, data_only=True)
    hojas = par.sheetnames[1:]
    par.close()
    grupo = {}
    for hoja in hojas:
        for d in _hoja_xlsx(a1o / "A-O_Parametrization.xlsx", hoja):
            t = d.get("Tech")
            if t:
                tec.setdefault(t, d.get("Tech.Name"))
                grupo.setdefault(t, nombres_hoja.get(hoja, hoja))
    sect = {d["Sector"]: d.get("Plain English") for d in _hoja_xlsx(a1i / "A-I_Classifier_Modes_Demand.xlsx", "Sectors")}
    fuel = {d["Code"]: d.get("Plain English") for d in _hoja_xlsx(a1i / "A-I_Classifier_Modes_Demand.xlsx", "Fuel_to_Code")}
    with io.open(gen / "tecnologias.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(["Código", "Nombre", "Grupo de la parametrización", "Tipo de planta"])
        for t in conj.get("TECHNOLOGY", []):
            nombre = tec.get(t)
            if not nombre and t.startswith("T5") and t[5:8] in sect and t[2:5] in fuel:
                nombre = "Demand %s for %s" % (fuel[t[2:5]], sect[t[5:8]])
            if t not in grupo and nombre and " - " in nombre:      # tecnologias del clasificador de oferta
                grupo[t] = nombre.split(" - ")[0]
            w.writerow([t, nombre or "", grupo.get(t, ""), C.familia(t) or ""])
    with io.open(gen / "combustibles.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(["Código", "Nombre"])
        for f in conj.get("FUEL", []):
            nombre = comb.get(f)
            if not nombre and f.startswith("E5") and f[2:5] in sect and f[5:8] in fuel:
                nombre = "Final - %s - %s" % (sect[f[2:5]], fuel[f[5:8]])
            w.writerow([f, nombre or ""])
    comp = {"DE": "demanda final", "PP": "generación eléctrica", "FUG": "fugitivas", "REF": "refinación",
            "sources": "todas las fuentes"}
    with io.open(gen / "emisiones.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(["Código", "Gas", "Componente", "En las emisiones del sector energía"])
        for e in conj.get("EMISSION", []):
            gas, _, cp = e.partition("_")
            w.writerow([e, gas, comp.get(cp, cp), "sí" if e in C.COMP_ENERGIA else ""])


def _fuentes(gen):
    with io.open(gen / "fuentes.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(["Fuente", "Uso", "Ubicación"])
        for f in K.hoja("Fuentes"):
            w.writerow([f.get("fuente") or "", f.get("uso") or "", f.get("ubicacion") or ""])
    if C.PROCEDENCIA.exists():
        shutil.copy2(C.PROCEDENCIA, gen / "procedencia.csv")
    if (C.RESULTADOS / "verificacion.txt").exists():         # informe de la verificacion de B3
        shutil.copy2(C.RESULTADOS / "verificacion.txt", gen / "verificacion.txt")
    with io.open(gen / "precios.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(["Precio", "US$/t", "Referencia", "Fuente", "Fecha del dato", "Estado"])
        for p in K.hoja("Precios"):
            w.writerow([p.get("escenario_precio"), p.get("usd_por_t"), p.get("referencia"), p.get("fuente"),
                        p.get("fecha_dato"), p.get("estado")])


def _corrida(gen):
    """_generado/corrida.rst: sustituciones con el alcance y las cifras de contexto de la corrida
    (metas, trayectorias, precios, lineas), para que ninguna pagina de la documentacion lleve una cifra
    escrita a mano. La fecha y el md5 de la corrida quedan en resultados_art6.json y en el manifiesto,
    no en las paginas. _generado/adicionalidad.csv: las lineas de Art6_Adicionalidad."""
    import graficar
    R = analizar.resultados()
    c = R.get("corrida", {})
    T = R.get("roles", {})
    cc = R.get("conciliacion_condicional_2035", {})
    m = R.get("metas_2035", {})
    H = {h["parametro"]: h["valor"] for h in K.hoja("Headroom")}
    P = sorted(K.hoja("Precios"), key=lambda p: float(p["usd_por_t"]))
    ad = K.hoja("Adicionalidad")
    grupos = R.get("grupos_de_lineas", {})
    por_cod = {x["codigo"]: x for x in ad}
    estados = T.get("estados", [])
    f = graficar.miles
    minus = lambda s: s[:1].lower() + s[1:] if len(s) > 1 and s[1:2].islower() else s
    pct = lambda v: (("%.0f" % v) if abs(v - round(v)) < 0.05 else ("%.1f" % v).replace(".", ",")) + " %"
    lineas = lambda clave: C.enumerar([str(por_cod[x]["linea"]) for x in grupos.get(clave, []) if x in por_cod]) or "ninguna"
    S = {
        "escenarios": str(len(c.get("escenarios", []))),
        "ventana": "%d-%d" % (C.VENTANA[0], C.VENTANA[-1]),
        "ventana_ini": str(C.VENTANA[0]), "ventana_fin": str(C.VENTANA[-1]),
        "inc": T.get("inc") or "s. d.", "con": T.get("con") or "s. d.",
        "escalera": C.enumerar(T.get("escalera", [])) or "ninguno", "n_escalera": str(len(T.get("escalera", []))),
        "pruebas": C.enumerar([minus(d["nombre"]) for d in estados if d["grupo"] == "Riesgo"]) or "ninguna",
        "sensibilidades": C.enumerar([minus(d["nombre"]) for d in estados if d["grupo"] != "Riesgo"]) or "ninguna",
        "meta_inc_kt": f(cc.get("reduccion_incondicional_publicada", 0.0)),
        "meta_con_kt": f(cc.get("reduccion_condicional_publicada", 0.0)),
        "pct_inc": pct(m.get("pct_meta_inc", 0.0)), "pct_con": pct(m.get("pct_meta_con", 0.0)),
        "anio_meta": str(H.get("anio_meta", "")), "anio_inicio_senda": str(H.get("anio_inicio_senda", "")),
        "tendencial_2035": f(cc.get("tendencial", 0.0), 1), "condicional_2035": f(m.get("condicional", 0.0), 1),
        "reduccion_modelada": f(cc.get("reduccion_modelada", 0.0), 1),
        "exceso_modelado": f(cc.get("exceso_modelado", 0.0), 1),
        "retro_energia": f(cc.get("tendencial_retroactivo_energia", 0.0)),
        "precios": C.enumerar(["%s %s" % (str(p["escenario_precio"]).lower(), f(float(p["usd_por_t"]))) for p in P]) + " US$/t",
        "regla_nacional": str(H.get("regla_nacional", "")), "regla_prudente": str(H.get("regla_prudente", "")),
        "archivo_nacional": str(H.get("archivo_nacional", "")), "n_lineas": str(len(ad)),
        "firmes": lineas("firme"), "sin_adicionalidad": lineas("sin_adicionalidad"),
        "condicionadas_declaradas": lineas("condicionadas_declaradas"),
    }
    io.open(gen / "corrida.rst", "w", encoding="utf-8").write(
        "".join(".. |%s| replace:: %s\n" % (k, v) for k, v in S.items()))
    with io.open(gen / "adicionalidad.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(["Código", "Línea", "Periodo", "Escenario de la NDC", "kt/año declarados", "Escenario Art. 6",
                    "Adicionalidad", "Veredicto de compradores"])
        for x in ad:
            kt = x.get("kt_anio_declarado")
            w.writerow([x.get("codigo"), x.get("linea"), x.get("periodo"), x.get("escenario_NDC"),
                        ("%.2f" % float(kt)).replace(".", ",") if isinstance(kt, (int, float)) else "",
                        x.get("escenario_art6"), x.get("adicionalidad"), x.get("veredicto_compradores") or ""])


def exportar_docs(F=None, L=None):
    """Copia a docs/ las tablas, las graficas PNG y los supuestos del libro de configuracion, y escribe
    los fragmentos que la documentacion incluye (figuras con su guia, fichas de escenarios, tablas de
    estructura y fuentes), para que la documentacion muestre siempre la ultima corrida."""
    gen, img = C.DOCS / "_generado", C.DOCS / "_static" / "graficas"
    gen.mkdir(parents=True, exist_ok=True)
    img.mkdir(parents=True, exist_ok=True)
    for f in C.TABLAS.glob("*.csv"):
        shutil.copy2(f, gen / f.name)
    for f in C.GRAFICAS.glob("*.png"):
        shutil.copy2(f, img / f.name)
    P = K.palancas()
    valor = {}          # 2030 y 2035 de cada palanca, leidos del datafile de un escenario que la usa
    for e in K.escenarios():
        if e["base"] == "B1" or not C.datafile(e["codigo"]).exists():
            continue
        T = C.tabla(C.leer_datafile(C.datafile(e["codigo"])))
        for pid in K.palancas_en(e["codigo"]):
            p = next((x for x in P if x["id"] == pid), None)
            if p is not None and (p["parametro"], str(p["clave"])) in T:
                valor.setdefault(pid, T[(p["parametro"], str(p["clave"]))][1])
    grupos = [g for g in dict.fromkeys(p["grupo"] for p in P) if any(p["grupo"] == g and p["id"] in valor for p in P)]
    RA = K.riesgos_agrupados(P)
    for viejo in gen.glob("palancas_*.csv"):          # se rehacen todos; los riesgos van agrupados
        viejo.unlink()
    with io.open(gen / "palancas_riesgos.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(["Riesgo", "Palancas", "Parámetros", "Tecnologías", "Qué cambia", "Fuente"])
        for g in grupos:
            if g in RA:
                r = RA[g]
                w.writerow([r["nombre"], "%s (%d)" % (r["rango"], r["n"]), r["parametros_txt"], r["tecnologias_txt"],
                            r["mecanismo"], r["fuente"]])
    for g in grupos:
        if g in RA:
            continue
        with io.open(gen / ("palancas_%s.csv" % g), "w", encoding="utf-8", newline="") as h:
            w = csv.writer(h)
            w.writerow(["Palanca", "Tecnología", "Parámetro", "2030", "2035", "Qué fija", "Iniciativa", "Adicionalidad", "Fuente"])
            for p in P:
                if p["grupo"] == g and p["id"] in valor:
                    w.writerow([p["id"], p["tecnologia"], p["parametro"].replace("TotalTechnologyAnnualActivity", "Actividad ")
                                .replace("LowerLimit", "piso").replace("UpperLimit", "techo"),
                                "%.3f" % valor[p["id"]][C.IDX[2030]], "%.3f" % valor[p["id"]][C.IDX[2035]], p["descripcion"],
                                p["iniciativa"], p["adicionalidad"], p["fuente"]])
    with io.open(gen / "escenarios.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(["Código", "Escenario", "Grupo", "Palancas aplicadas", "Narrativa"])
        for e in K.escenarios():
            w.writerow([e["codigo"], e["nombre"], e["grupo"], e["palancas"] or "(B1)", e["narrativa"]])
    with io.open(gen / "riesgos.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(["Riesgo", "Nombre", "Desde", "Mecanismo", "Derivado de", "Validación", "Nota"])
        for r in K.hoja("Riesgos"):
            w.writerow([r["grupo"], r["nombre"], r["desde"], r["mecanismo"], r["derivado_de"], r["validacion"], r["nota"]])
    for nombre, archivo, clave in (("AnoSeco", "anoseco_clases.csv", "Clase"),
                                   ("AnoSeco_Centrales", "anoseco_centrales.csv", "Clase del modelo")):
        try:                    # sandbox PNUD: sin filas por central (Art6_AnoSeco_Centrales no se publica)
            filas = [f for f in K.hoja(nombre) if str(f.get(clave) or "").startswith("PP")]
        except KeyError:
            continue
        if filas:
            cab = [c for c in filas[0] if c and c != "Fuente"]
            with io.open(gen / archivo, "w", encoding="utf-8", newline="") as h:
                w = csv.writer(h)
                w.writerow(cab)
                for f in filas:
                    w.writerow([("%.4f" % f[c]).replace(".", ",") if c.startswith("Razón") and isinstance(f[c], float)
                                else (("%.2f" % f[c]).replace(".", ",") if c.startswith("Caudal")
                                      and isinstance(f[c], float) else f[c])
                                for c in cab])
    import explicacion
    T = explicacion.supuestos_tabla()
    cols = explicacion.columnas()
    cab = ["Escenario"] + ["%s %s" % (g, et) for _, et, g in cols] + ["Qué cambia"]
    with io.open(gen / "supuestos_escenarios.csv", "w", encoding="utf-8", newline="") as h:
        w = csv.writer(h)
        w.writerow(cab)
        for cod, r in T.iterrows():
            celdas = []
            for k, _, _ in cols:
                v = r[k]
                if v is None or v != v:
                    celdas.append("–")
                elif v >= explicacion.SIN_TECHO:
                    celdas.append("sin techo")
                else:
                    celdas.append(("%.3f" if k.startswith("fp") else "%.2f") % v)
            w.writerow([cod] + celdas + [r["que_cambia"]])
    if F is None or L is None:
        import graficar
        F, L = graficar.generar(png=False)
    _figuras_rst(F, L, gen)
    _escenarios_html(gen)
    _estructura(gen)
    _fuentes(gen)
    _corrida(gen)
    _grupos_palancas(gen, [g for g in grupos if g not in RA], bool(RA))
    import portada
    portada.generar()               # gráfica de impacto de la portada (docs/_generado/portada.rst)
    print("  documentación actualizada en docs/_generado y docs/_static/graficas")


def _grupos_palancas(gen, grupos, riesgos=False):
    """_generado/palancas.rst: una tabla por grupo de palancas de los escenarios y una tabla de los
    riesgos hidrologicos con una fila por riesgo, para incluirlas en la documentacion."""
    nombre = {e["codigo"]: e["nombre"] for e in K.escenarios()}
    estados = {d["sufijo"]: d["nombre"] for d in analizar.resultados().get("roles", {}).get("estados", [])}
    partes = []
    for g in grupos:
        titulo = "Palancas del grupo %s" % g
        if g in nombre:
            titulo = "Palancas del escenario «%s» (%s)" % (nombre[g], g)
        elif g in estados:
            titulo = "Palancas de «%s» (%s)" % (estados[g], g)
        partes += [titulo, "^" * len(titulo), "", ".. csv-table::", "   :file: /_generado/palancas_%s.csv" % g,
                   "   :header-rows: 1", "   :widths: 6, 10, 12, 7, 7, 30, 10, 10, 18", ""]
    if riesgos:
        titulo = "Palancas de los riesgos hidrológicos"
        partes += [titulo, "^" * len(titulo), "",
                   "Cada riesgo agrupa las palancas que aplican el mismo choque a las clases hidroeléctricas; los valores "
                   "por clase están en la matriz de parámetros de entrada y en la hoja ``Art6_Palancas``.", "",
                   ".. csv-table::", "   :file: /_generado/palancas_riesgos.csv", "   :header-rows: 1",
                   "   :widths: 16, 12, 14, 14, 28, 16", ""]
    io.open(gen / "palancas.rst", "w", encoding="utf-8").write("\n".join(partes) + "\n")


def _sphinx():
    """Interprete con Sphinx: el de la cadena o, si no lo tiene, el python del sistema."""
    import importlib.util
    if importlib.util.find_spec("sphinx"):
        return sys.executable
    for cand in ("python", "python3", "py"):
        exe = shutil.which(cand)
        if exe and Path(exe).resolve() != Path(sys.executable).resolve():
            r = subprocess.run([exe, "-c", "import sphinx"], capture_output=True)
            if r.returncode == 0:
                return exe
    return None


def construir_docs():
    """Compila docs/ con Sphinx en docs/_build/html, con el interprete que tenga Sphinx. Si no hay
    ninguno, lo dice y sigue: el flujo documentation.yml compila en GitHub al publicar."""
    if not (C.DOCS / "conf.py").exists():
        print("  sin conf.py en %s: los fragmentos quedan generados; la documentación del Modelo CLG no se toca"
              % C.DOCS.relative_to(C.ENERGY))
        return False
    py = _sphinx()
    if not py:
        print("  Sphinx no está instalado: docs/_build queda sin actualizar; el flujo de GitHub compila al publicar")
        return False
    destino = C.DOCS / "_build" / "html"
    import tempfile                       # los doctrees (cache con rutas locales) quedan fuera de la entrega
    doctrees = Path(tempfile.gettempdir()) / "art6_doctrees"
    r = subprocess.run([py, "-m", "sphinx", "-b", "html", "-q", "-d", str(doctrees), str(C.DOCS), str(destino)],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    avisos = [l for l in (r.stderr or "").splitlines() if "WARNING" in l or "ERROR" in l]
    if r.returncode != 0:
        print("  Sphinx terminó con errores:\n    " + "\n    ".join((r.stderr or "").splitlines()[-12:]))
        return False
    print("  documentación compilada en %s%s" % (destino / "index.html",
                                                 (" (%d avisos)" % len(avisos)) if avisos else ""))
    for l in avisos[:12]:
        print("    " + l)
    return True


def resumen_solver():
    """Estado de la ultima corrida de B1 en cada escenario (comun.estado_solucion)."""
    print()
    print("6b: Resultado del solver por escenario")
    fallas = []
    for e in K.escenarios():
        if str(e["correr"]).upper() != "SI":
            continue
        est = C.estado_solucion(e["codigo"])
        ok = est == "optimo"
        fallas += [] if ok else [e["codigo"]]
        print("  %-16s %s" % (e["codigo"], "solución óptima" if ok else "SIN SOLUCIÓN ÓPTIMA (%s)" % est))
    if fallas:
        print("  ATENCIÓN: %d escenarios sin solución óptima quedan fuera del análisis: %s" % (len(fallas), fallas))
    return not fallas


def posterior():
    """Etapa de Artículo 6 como paso posterior a B1 en el Modelo Climate Lead Group (rama art6-escenarios). Solo lee
    los datafiles, el libro de configuración y las salidas de cada escenario en Executables; no modifica nada de lo que
    produce el modelo."""
    t0 = time.time()
    resumen_solver()
    print()
    print("8: Etapa de Artículo 6: verificación de los escenarios y manifiesto")
    if not verificar.main():
        print("  ATENCIÓN: la verificación encontró errores; ver B2_Output_Art6/verificacion.txt")
    manifiesto()
    print()
    print("9: Análisis, gráficas, tablero, cuaderno y documentación")
    analizar.main()
    # contabilidad con las decisiones del 13-10-2026, recalculada desde las tablas (se agrega a verificacion.txt)
    if not verificar.contabilidad():
        print("  ATENCIÓN: la contabilidad del Artículo 6 encontró errores; ver B2_Output_Art6/verificacion.txt")
    import graficar
    F, L = graficar.generar(png=True)
    graficar.exportar_html(F, L)
    import tablero
    tablero.exportar(F, L)
    exportar_docs(F, L)
    ejecutar_cuaderno()
    construir_docs()
    print("  etapa de Artículo 6 en %.0f s" % (time.time() - t0))

