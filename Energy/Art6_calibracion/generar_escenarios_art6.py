# -*- coding: utf-8 -*-
"""Genera los escenarios del Artículo 6 sobre la base congelada CONA6 (rama art6-escenarios de ECU_NDC).

Uso, desde ECU_NDC/Energy:
    python Art6_calibracion/generar_escenarios_art6.py              escribe filas, hojas documentales y carpetas
    python Art6_calibracion/generar_escenarios_art6.py --comprobar  no escribe; sale con 1 si el disco no coincide
    python Art6_calibracion/generar_escenarios_art6.py --quitar     retira del libro las filas del Artículo 6 (para
                                                                     volver a correr generar_referencias.py, que exige
                                                                     que las filas de los clones vayan al final)

Qué genera (decisiones del 11 y 12-10-2026):
- Referencias de estado, basadas en CONA6: CONA6P (sin los pisos térmicos del Modelo CLG desde 2026) y CONA6I
  (importación desde Colombia en cero desde 2026).
- Familias A, B y C basadas en CONA6 (Based_On = CONA6), con los estados intacto, CCOFF, CCSOLO, CCDER, ANOSECO,
  ANOSEC3, ANOSECM, CCOFFANOSECO, FOI, IMP0, ACT y SIN (solo A); C sin FOI ni SIN.
- Decisiones del 13-10-2026: respaldo de emergencia (grupo RESP: tope de PPFOI liberado desde 2026, CapitalCost y
  FixedCost de PPFOI por 1.000, STOELE cerrado en 2026-2035) en los estados de riesgo; familia GAS (A-GAS, B-GAS y
  C-GAS: las palancas de A, B y C con Based_On = CONA6G, la base sin la expansión a gas declarada; comprobar_base_gas
  falla si las filas que tocan difieren entre CONA6 y CONA6G); estados ACT y ACTDEM (ACT-DEM) anclados a 2026, después
  de la verificación independiente de K16: en 2026, demanda eléctrica, hidrología por clase, Coca Codo e importación
  observadas (act_2026_observado.csv y act_2026_centrales.csv, de situacion_act_2026.py); desde 2027, un año seco cada
  tres (2029, 2032 y 2035) e importación cero; la demanda vuelve a la del Modelo CLG desde 2027 (ACT) o persiste
  (ACTDEM). Con --solo ACT,ACTDEM solo esos escenarios quedan activos para una corrida parcial de B1.
- Puentes A-V7, B-V7 y C-V7: las palancas del sandbox PNUD del 08-10-2026 con sus valores exactos desde 2026, sobre la
  historia de CONA6.

Cada palanca se recalcula desde el datafile y la salida de CONA6 de la corrida congelada (12-10-2026, 16:46; el
script falla si el datafile de CONA6 no es el de tests/base_congelada.json). Ninguna fila copia valores del sandbox
salvo los puentes V7, que existen para eso, y las razones de demanda del Escenario C (misma regla del sandbox:
razón C/NDCCON año a año, aplicada a la demanda de CONA6).

B1 no hereda las filas de Electrical, SmartGrid ni Efficiency del escenario base (B1_Base_Scenarios_Adj_Parallel.py,
l. 3807-3819), así que cada escenario lleva una copia de las filas de CONA6 en esas hojas, salvo las claves que
reemplazan sus palancas (B1 aplica solo la primera fila de cada clave). Las palancas de Electrical se escriben como
vector completo 2010-2035 (Exact_Years hasta 2034 e hito 2035): «intact» leería la carpeta A2 de CONA6, que no trae
las correcciones de K11 y K12 hechas en filas de Electrical, y reescribiría la historia.

El libro B1_Scenario_Config.xlsx se edita en su XML con las utilidades de generar_referencias.py: las filas del
Modelo CLG y de los clones quedan byte a byte; las filas del Artículo 6 van al final de cada hoja y se reescriben en
cada corrida. Sandbox PNUD: las hojas documentales (Art6_*) van al final del mismo B1_Scenario_Config.xlsx (B1 no las lee) y las filas por central quedan fuera del repositorio: el año seco, Coca Codo y el estado ACT leen los parámetros por clase de parametros_hidro_por_clase.csv (agregar_por_clase.py).
"""
import argparse
import csv
import hashlib
import io
import json
import re
import shutil
import sys
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape as xml_escape
from pathlib import Path

import openpyxl
import pandas as pd

AQUI = Path(__file__).resolve().parent
ENERGY = AQUI.parent
sys.path.insert(0, str(AQUI))
sys.path.insert(0, str(ENERGY / "Art6_motor"))
import generar_referencias as GR  # noqa: E402  utilidades XML del libro
import agregar_por_clase as AG  # noqa: E402  parámetros hidroeléctricos por clase (sin filas por central)
import comun as C  # noqa: E402  lectura de datafiles

# versión del sandbox del 08-10-2026 (puentes V7, razones de demanda del C y hojas Art6 de partida), congelada en el
# repositorio con la misma estructura de carpetas
SANDBOX = AQUI / "version_2026-10-08"          # solo lectura
LIBRO = ENERGY / "B1_Scenario_Config.xlsx"
ART6_LIBRO = LIBRO                     # sandbox PNUD: las hojas Art6_* van en el libro de B1
RESPALDOS = ENERGY.parent / "_stale_respaldos"       # copias del libro antes de escribirlo (fuera de git)
REGISTRO = AQUI / "registro_escenarios_art6.json"
BASE = "CONA6"
ANIOS = list(range(2010, 2036))
IX = {a: i for i, a in enumerate(ANIOS)}
HOJAS_FILAS = ("SmartGrid", "Electrical", "Efficiency")
PISO, TECHO, FP = C.PISO, C.TECHO, C.FP
INV, TOPE = "TotalAnnualMinCapacityInvestment", "TotalAnnualMaxCapacity"
CAU = 31.536            # CapacityToActivityUnit de PPWNDON, PPPVT, PPBIM y PPGEO en el datafile de CONA6 (PJ/GW-año)
HIDRO = ["PPHD", "PPHDAMAB", "PPHDAMAM", "PPHDAMAS", "PPHDAMPM", "PPHDAMPS",
         "PPHRORAB", "PPHRORAM", "PPHRORAS", "PPHRORPM", "PPHRORPS"]
METODO = "Overwrite ; Interpolate ; Fix_Last"
METODO_NO = "Write ; Interpolate ; Fix_Last"

# ------------------------------------------------------------------------------------------ supuestos con fuente
B9_MW = 39.3                      # ESIN-CELEC EP (2020), cuatro plantas de palma y cascarilla (decisión 11-10-2026)
B9_FP_DECLARADO = 0.56            # ESIN, lámina 31; el CF del modelo es FP / AF
B_EOL_BLOQUE, B_EOL_ANIOS = 0.117, range(2028, 2033)     # 585 MW en cinco bloques (PEG 2023-2032, tabla 4-7)
B_SOL_BLOQUE, B_SOL_ANIOS = 0.198, range(2028, 2036)     # 1.584 MWp en ocho bloques (CELEC EP y AFD, 2024)
GEO_ANIO, GEO_GW, GEO_PJ = 2029, 0.051, 1.27             # Chachimbiro (PEG 2023-2032, tabla 4-4); techo del sandbox
TUFINO_GW, TUFINO_ANIO = 0.138, 2035                      # Tufiño-Chiles-Cerro Negro (supuesto del análisis)
GEO_TOPE = 0.5
SECO_MODERADO = 0.908             # caída de la hidro de 2024: -9,2 % (BEN 2025, tabla 2.20)
CC_DERRATEO = 0.20                # derrateo crónico de Coca Codo (supuesto del análisis; CENACE 2025)
ANIOS_SECOS_3 = (2026, 2029, 2032, 2035)                 # un año seco cada tres, 2026 observado como seco

FUENTE_SECO = "CENACE (2025), Informe Anual 2024, tabla 1.14, p. 73; razón por clase ponderada por la potencia " \
              "nominal de la capa CENTRALES (Art6_calibracion/parametros_hidro_por_clase.csv, %s)" % AG.NOTA
FUENTE_CC = "CENACE, Informe Anual 2025, Parte 1, figuras 1.4 y 1.5, p. 31-32 (S26-008): Coca Codo 7.695,99 GWh y " \
            "Sopladora 2.884,03 GWh brutos (Art6_calibracion/parametros_hidro_por_clase.csv)"
FUENTE_CC_OFF = "Supuesto del análisis: salida desde 2029 por el avance de la erosión regresiva del río Coca " \
                "(frente a 3,6 km de la captación en mayo de 2026)"
FUENTE_PEG = "Plan de Expansión de la Generación 2023-2032, tabla 4-4 (Chachimbiro 50 MW en 2029) y tabla 4-7 " \
             "(585 MW eólicos)"
FUENTE_P3 = "Climate Lead Group (2025), P3 Informe de modelamiento de la Segunda NDC, tabla 6, p. 19 (677 MW solares " \
            "y 247,5 MW eólicos); capacidad existente y factor de planta de CONA6 (K11 y K12, " \
            "Art6_calibracion/correcciones_K.csv)"
# respaldo de emergencia en los estados de riesgo (decisión del 13-10-2026: «respaldo de fuel oil caro solo
# en los estados de riesgo»): los mismos valores que las filas RE01 de INCA6G
CARO = 1000.0
FUENTE_RESP = ("Decisión del 13-10-2026: "
               "respaldo de fuel oil caro solo en los estados de riesgo, supuesto de cierre como la térmica arrendada que "
               "el país contrata hoy; regla de la variante de potencia mínima (Art6_calibracion/variante_suficiencia.py)")
NOTA_RE01 = ("En INCA6G las mismas palancas son las filas RE01 de correcciones_K.csv: un supuesto de cierre introducido en la "
             "corrida K16 del 13-10-2026, fuera del acuerdo original (respaldo solo en estados de riesgo); se aprobó "
             "el 13-10-2026, después de la verificación independiente, conservarlo rotulado «con cierre de fuel oil de "
             "emergencia», con el peso del respaldo al lado.")
# estado ACT anclado a 2026 (decisión del 13-10-2026): datos observados de CENACE y XM leídos de los
# documentos primarios por situacion_act_2026.py
ACT_OBSERVADO = AQUI / "act_2026_observado.csv"
ACT_CENTRALES = AQUI / "act_2026_centrales.csv"
ANIOS_SECOS_ACT = (2029, 2032, 2035)        # un año seco cada tres después de 2026, que entra con su hidrología observada
# la potencia efectiva de Coca Codo Sinclair (1.500 MW, CENACE, Informe Anual 2025, Parte 1, p. 36) se lee de
# act_2026_observado.csv (clave cc_potencia_efectiva_MW), que situacion_act_2026.py toma del informe
PJ_POR_GWH = 0.0036
DEMANDA_ELE = ("E5COMELE", "E5INDELE", "E5RESELE", "E5TROELE", "E5ALLTRNELE", "E5CONELE")
# centrales del tablero de CENACE en cada clase (nombre en centrales_clases.csv: nombre en el tablero); el resto de la
# potencia de la clase y las clases sin central del tablero toman la razón de «Otras Hidro»
CLASE_CENTRALES = {"PPHDAMAB": {"Molino": "Paute"}, "PPHDAMAM": {"Mazar": "Mazar", "Agoyán": "Agoyán"},
                   "PPHRORPM": {"San Francisco": "San Francisco", "Minas San Francisco": "Minas San Francisco"},
                   "PPHRORAM": {"Delsitanisagua": "Delsitanisagua"}}
FUENTE_ACT = ("CENACE: Informes Ejecutivos de Gestión de agosto y diciembre de 2025 y agosto de 2026 (secciones 1, 2 y 4), "
              "indicadores.xlsx (hoja datos), Informe Anual 2025 Parte 1 (figura 1.4 y p. 36) e Información Operativa al "
              "10-10-2026 (SCADA preliminar); XM, Exportaciones kWh 2026. Leídos por Art6_calibracion/situacion_act_2026.py "
              "en act_2026_observado.csv; razones por clase en parametros_hidro_por_clase.csv")
BASE_GAS = "CONA6G"                 # base de la familia GAS: CONA6 sin la expansión a gas declarada (G01)


class ErrorGenerador(Exception):
    pass


def md5(b):
    return hashlib.md5(b).hexdigest()


def r4(x):
    return round(float(x) + 0.0, 4)


# ------------------------------------------------------------------------------------------ lectura de la base
def tabla(cod, raiz=ENERGY):
    ruta = raiz / "Executables" / f"{cod}_0" / f"{cod}_0.txt"
    return C.tabla(C.leer_datafile(ruta)), ruta


def capacidad_2025(tec):
    o = pd.read_csv(ENERGY / "Executables" / f"{BASE}_0" / f"{BASE}_0_Output.csv",
                    usecols=["Technology", "Year", "TotalCapacityAnnual"], low_memory=False)
    x = o[(o.Technology == tec) & (o.Year == 2025)]["TotalCapacityAnnual"].dropna()
    if x.empty:
        raise ErrorGenerador(f"{BASE}_0_Output.csv no trae TotalCapacityAnnual de {tec} en 2025")
    return float(x.max())


def comprobar_base():
    fijo = json.loads((ENERGY / "tests" / "base_congelada.json").read_text(encoding="utf-8"))
    ruta = ENERGY / "Executables" / f"{BASE}_0" / f"{BASE}_0.txt"
    lineas = ruta.read_bytes().replace(b"\r", b"").split(b"\n")
    if md5(b"\n".join(sorted(lineas))) != fijo["referencias"][BASE]["datafile"]:
        raise ErrorGenerador(f"el datafile de {BASE} no es el de la base congelada (tests/base_congelada.json)")
    salida = ENERGY / "Executables" / f"{BASE}_0" / f"{BASE}_0_Output.csv"
    md5_salida = md5(salida.read_bytes())
    if md5_salida != fijo["corrida"]["CONA6_0_Output_md5"]:
        raise ErrorGenerador(f"{salida.name} tiene md5 {md5_salida}, no el de la corrida congelada")
    return {"datafile_ordenado": fijo["referencias"][BASE]["datafile"], "output_md5": md5_salida}


def comprobar_base_gas(T, P, E):
    """La familia GAS lleva las mismas palancas de A, B y C sobre CONA6G. Se calculan sobre CONA6: falla si alguna fila
    que tocan, o la capacidad eólica y solar de 2025 que usan, difiere entre CONA6 y CONA6G (entonces la palanca no sería
    la misma sobre las dos bases)."""
    grupos = {g for e in E if e.get("base") == BASE_GAS for g in e["palancas"].split(",")}
    if not grupos:
        return None
    TG, ruta_g = tabla(BASE_GAS)
    distintas = []
    for p in P:
        if p.grupo not in grupos:
            continue
        k = (p.parametro, p.clave)
        a, b = T.get(k), TG.get(k)
        if (a is None) != (b is None) or (a and any(abs(x - y) > 1e-12 for x, y in zip(a[1], b[1]))):
            distintas.append(f"{p.id} {k}")
    o = {}
    for esc in (BASE, BASE_GAS):
        x = pd.read_csv(ENERGY / "Executables" / f"{esc}_0" / f"{esc}_0_Output.csv",
                        usecols=["Technology", "Year", "TotalCapacityAnnual"], low_memory=False)
        o[esc] = x[(x.Year == 2025) & x.Technology.isin(["PPWNDON", "PPPVT"])].groupby("Technology") \
            .TotalCapacityAnnual.max().to_dict()
    if o[BASE] != o[BASE_GAS]:
        distintas.append(f"capacidad 2025 {o[BASE]} frente a {o[BASE_GAS]}")
    if distintas:
        raise ErrorGenerador(f"familia GAS: {BASE} y {BASE_GAS} difieren en filas que tocan las palancas: {distintas[:5]}")
    return {"base": BASE_GAS, "grupos": sorted(grupos), "datafile_md5": md5(ruta_g.read_bytes()),
            "filas_iguales_en_las_dos_bases": len([p for p in P if p.grupo in grupos])}


# ------------------------------------------------------------------------------------------ palancas
class Palanca:
    def __init__(self, pid, grupo, parametro, tecnologia, vector, unidad, descripcion, iniciativa, adicionalidad,
                 fuente, hoja="Electrical", clave=None, builtin="YES", familia=""):
        self.id, self.grupo, self.parametro, self.tecnologia = pid, grupo, parametro, tecnologia
        self.vector = [r4(v) for v in vector] if hoja == "Electrical" else [float(v) for v in vector]
        self.unidad, self.descripcion, self.iniciativa = unidad, descripcion, iniciativa
        self.adicionalidad, self.fuente, self.hoja, self.builtin = adicionalidad, fuente, hoja, builtin
        self.clave = clave or (tecnologia + ",All" if parametro == FP else tecnologia)
        self.familia = familia
        if len(self.vector) != len(ANIOS):
            raise ErrorGenerador(f"{pid}: el vector no tiene 26 años")

    def doc(self):
        return {"grupo": self.grupo, "id": self.id, "tecnologia": self.tecnologia, "clave": self.clave,
                "familia": self.familia, "parametro": self.parametro, "unidad": self.unidad,
                "descripcion": self.descripcion, "iniciativa": self.iniciativa, "adicionalidad": self.adicionalidad,
                "fuente": self.fuente}


def vec(T, parametro, clave):
    if (parametro, clave) not in T:
        raise ErrorGenerador(f"el datafile de {BASE} no trae {parametro} {clave}")
    return list(T[(parametro, clave)][1])


def desde(base, anio, valor_o_funcion):
    """Copia de base con los años >= anio reemplazados."""
    out = list(base)
    for a in ANIOS:
        if a >= anio:
            out[IX[a]] = valor_o_funcion(a) if callable(valor_o_funcion) else valor_o_funcion
    return out


def potencial(existente, inversion, cf, af, anio_ini=2026):
    """Potencial anual (PJ) de la capacidad existente más la inversión acumulada desde anio_ini."""
    out, acum = [], 0.0
    for a in ANIOS:
        if a >= anio_ini:
            acum += inversion[IX[a]]
        out.append((existente + acum) * cf[IX[a]] * af[IX[a]] * CAU)
    return out


def razones_secas():
    """Razón de caudales 2024 por clase (sandbox PNUD: del archivo agregado parametros_hidro_por_clase.csv, que
    agregar_por_clase.py deriva de los datos por central con las mismas fórmulas; el repositorio no lleva filas por
    central). La versión con los archivos por central queda abajo como _razones_secas_por_central."""
    d = AG.leer()
    razon = {c: d[("razon_anoseco_2024", c)][0] for c in HIDRO}
    origen = {c: d[("razon_anoseco_2024", c)][1] for c in HIDRO}
    return razon, origen, []


def _razones_secas_por_central():
    """Razón de caudales 2024 por clase, recompuesta con centrales_clases.csv (ponderada por potencia nominal)."""
    cen = {r["serie_cenace"]: float(r["razon"]) for r in csv.DictReader(io.open(AQUI / "anoseco_cenace_2024.csv",
                                                                                   encoding="utf-8-sig"))}
    # central del modelo -> serie del CENACE (tabla 1.14). Sopladora con la razón del Paute y San Francisco con la de
    # Agoyán (decisión del 12-10-2026); Pucará con la de Pisayambo, el embalse que la alimenta según el
    # campo de río de centrales_clases.csv («Pastaza (embalse Pisayambo)»)
    serie = {"Molino": "Paute, ingreso Amaluza", "Mazar": "Mazar", "Sopladora": "Paute, ingreso Amaluza",
             "Coca Codo Sinclair": "Coca Codo Sinclair", "Agoyán": "Agoyán", "San Francisco": "Agoyán",
             "Pucará": "Pisayambo", "Delsitanisagua": "Delsitanisagua",
             "Minas San Francisco": "Minas San Francisco", "Marcel Laniado": "Marcel Laniado"}
    filas = list(csv.DictReader(io.open(AQUI / "centrales_clases.csv", encoding="utf-8-sig")))
    clases, detalle = {}, []
    for f in filas:
        cl = f["clase_modelo"]
        if not cl:
            continue
        s = serie.get(f["central"])
        p = float(f["potencia_nominal_MW"])
        detalle.append({"central": f["central"], "clase": cl, "MW_nominal": p, "serie_cenace": s or "sin serie",
                        "razon": cen[s] if s else None})
        if s:
            clases.setdefault(cl, []).append((p, cen[s], f["central"]))
    razon, origen = {}, {}
    for cl, xs in clases.items():
        razon[cl] = sum(p * r for p, r, _ in xs) / sum(p for p, _, _ in xs)
        origen[cl] = "medida: " + ", ".join(f"{c} ({p:g} MW, {r:.4f})" for p, r, c in xs)
    # clases sin central con serie en la tabla: herencia del sandbox; PPHD con el promedio nacional por caudal
    razon["PPHDAMAS"], origen["PPHDAMAS"] = cen["Pisayambo"], "heredada: Pisayambo (embalse amazónico pequeño)"
    razon["PPHRORAS"], origen["PPHRORAS"] = razon["PPHDAMAS"], "heredada de PPHDAMAS"
    razon["PPHDAMPS"], origen["PPHDAMPS"] = razon["PPHDAMPM"], \
        "heredada de PPHDAMPM (Toachi Pilatón-Sarapullo sin serie)"
    razon["PPHRORPS"], origen["PPHRORPS"] = razon["PPHRORPM"], "heredada de PPHRORPM"
    rows = list(csv.DictReader(io.open(AQUI / "anoseco_cenace_2024.csv", encoding="utf-8-sig")))
    q24 = sum(float(r["caudal_2024_m3s"]) for r in rows)
    qh = sum(float(r["caudal_historico_m3s"]) for r in rows)
    razon["PPHD"], origen["PPHD"] = q24 / qh, "promedio nacional de las 8 series, ponderado por caudal"
    faltan = [c for c in HIDRO if c not in razon]
    if faltan:
        raise ErrorGenerador(f"sin razón de año seco para {faltan}")
    return {c: round(razon[c], 4) for c in HIDRO}, origen, detalle


def participacion_coca_codo():
    """Participación de Coca Codo en la energía de PPHRORAB en 2025 (archivo agregado por clase)."""
    return AG.leer()[("participacion_coca_codo_2025", "PPHRORAB")][0]


def _participacion_coca_codo_por_central():
    filas = {r["central"]: float(r["produccion_bruta_2025_GWh"])
             for r in csv.DictReader(io.open(AQUI / "participacion_pphrorab_2025.csv", encoding="utf-8-sig"))}
    return filas["Coca Codo Sinclair"] / (filas["Coca Codo Sinclair"] + filas["Sopladora"])


def act_2026(T, s_cc):
    """Estado ACT anclado a 2026 (decisión del 13-10-2026), desde act_2026_observado.csv y
    act_2026_centrales.csv (situacion_act_2026.py los lee de los documentos de CENACE y XM).

    Hidrología de 2026 por clase. La actividad de 2026 de cada clase es la de CONA6 en 2025 (año anclado al balance) por
    la razón observada 2026/2025 de sus centrales: la participación de la central en la hidro acumulada al 10-10-2026
    sobre su participación en 2025, por la razón de la hidro total de enero a agosto. Las centrales se ponderan por su
    participación en la potencia de la clase (centrales_clases.csv) y el resto de la clase toma la razón de «Otras Hidro».
    PPHRORAB es la excepción: Coca Codo entra con su factor de planta de 2026 estimado (banda de 0,55 a 0,64)
    sobre 1.500 MW y Sopladora con su razón. El factor de 2026 que multiplica el factor de planta, el piso y el techo de
    CONA6 es esa actividad sobre el techo de CONA6 en 2026; en una clase sin actividad en 2025 es la razón misma.

    Demanda: la eléctrica final de 2026 (E5*ELE) es la de CONA6 en 2025 por el crecimiento observado de enero a agosto; m
    es esa demanda sobre la de CONA6 en 2026 y multiplica cada combustible (en ACT solo 2026; en ACT-DEM desde 2026).
    Importación: en 2026 la importada antes del 07-09-2026; cero desde 2027."""
    leer = lambda p: list(csv.DictReader(io.open(p, encoding="utf-8-sig")))
    obs = {r["clave"]: float(r["valor"]) for r in leer(ACT_OBSERVADO)}
    agr = AG.leer()            # sandbox PNUD: razones 2026/2025 por clase (parametros_hidro_por_clase.csv)
    o = pd.read_csv(ENERGY / "Executables" / f"{BASE}_0" / f"{BASE}_0_Output.csv",
                    usecols=["Technology", "Year", "TotalTechnologyAnnualActivity"], low_memory=False)
    act25 = o[(o.Year == 2025) & o.Technology.isin(HIDRO)].groupby("Technology").TotalTechnologyAnnualActivity.max()
    fp_cc = obs["cc_fp_2026_estimado"]
    CC_MW = obs["cc_potencia_efectiva_MW"]
    if not 0.55 <= fp_cc <= 0.64:
        raise ErrorGenerador(f"factor de planta de Coca Codo de 2026 ({fp_cc:.4f}) fuera de la banda 0,55-0,64")
    clases, f26 = [], {}
    for cl in HIDRO:
        a25 = float(act25.get(cl, 0.0)) if pd.notna(act25.get(cl, 0.0)) else 0.0
        up26, cf26 = vec(T, TECHO, cl)[IX[2026]], vec(T, FP, cl + ",All")[IX[2026]]
        if cl == "PPHRORAB":
            cc_pj = fp_cc * CC_MW * 8.76 * PJ_POR_GWH
            raz_sop = agr[("razon_2026_2025_resto", "PPHRORAB")][0]
            sop_pj = (1.0 - s_cc) * a25 * raz_sop
            a26, r_cl = cc_pj + sop_pj, (cc_pj + sop_pj) / a25
            origen = ("Coca Codo con factor de planta %.4f (estimado de 2026: acumulado al 10-10 más el resto del año al "
                      "ritmo del 1 al 10 de octubre) sobre %s MW: %.4f PJ; Sopladora, participación %.4f de CONA6 2025 "
                      "por su razón %.4f: %.4f PJ" % (fp_cc, "{:,.0f}".format(CC_MW).replace(",", "."), cc_pj,
                                                     1.0 - s_cc, raz_sop, sop_pj))
        else:
            r_cl, origen = agr[("razon_2026_2025", cl)]
            a26 = a25 * r_cl
        if a25 < 0.05 * up26:                 # clase sin actividad en 2025: la razón se aplica al valor de CONA6 de 2026
            f, a26 = r_cl, up26 * r_cl
            origen += "; sin actividad en 2025: razón sobre CONA6 2026"
        else:
            f = a26 / up26
        if cf26 * f >= 1.0:
            raise ErrorGenerador(f"{cl}: el factor de 2026 ({f:.4f}) lleva el factor de planta a {cf26 * f:.4f}")
        f26[cl] = round(f, 6)
        clases.append({"clase": cl, "actividad_CONA6_2025_PJ": round(a25, 4), "razon_2026_2025": round(r_cl, 6),
                       "actividad_ACT_2026_PJ": round(a26, 4), "techo_CONA6_2026_PJ": round(up26, 4),
                       "factor_2026": f26[cl], "fp_CONA6_2026": round(cf26, 4), "fp_ACT_2026": round(cf26 * f, 4),
                       "origen": origen})
    dem = {c: vec(T, "SpecifiedAnnualDemand", c) for c in DEMANDA_ELE}
    d25, d26 = (sum(v[IX[a]] for v in dem.values()) for a in (2025, 2026))
    g = obs["crecimiento_demanda_2026"]
    m = d25 * g / d26
    imp_pj = obs["importacion_2026_antes_07_09_GWh"] * PJ_POR_GWH
    hid25 = sum(x["actividad_CONA6_2025_PJ"] for x in clases)
    hid26 = sum(x["actividad_ACT_2026_PJ"] for x in clases)
    return {"factor_2026": f26, "clases": clases, "observado": obs, "multiplicador_demanda": round(m, 6),
            "demanda_CONA6_2025_PJ": round(d25, 4), "demanda_CONA6_2026_PJ": round(d26, 4),
            "demanda_ACT_2026_PJ": round(d25 * g, 4), "crecimiento_demanda_2026": g,
            "importacion_2026_PJ": round(imp_pj, 6), "fp_coca_codo_2026": fp_cc,
            "hidro_CONA6_2025_PJ": round(hid25, 4), "hidro_ACT_2026_PJ": round(hid26, 4),
            "razon_hidro_modelo_2026_2025": round(hid26 / hid25, 4) if hid25 else None,
            "anios_secos": list(ANIOS_SECOS_ACT),
            "insumos_md5": {p.name: md5(p.read_bytes()) for p in (ACT_OBSERVADO, AG.SALIDA)}}


def palancas(T, SB):
    """Todas las palancas, recalculadas desde CONA6 (T) y, para los puentes V7, desde el sandbox (SB)."""
    P = []
    af_w, af_s, af_b = vec(T, "AvailabilityFactor", "PPWNDON"), vec(T, "AvailabilityFactor", "PPPVT"), \
        vec(T, "AvailabilityFactor", "PPBIM")
    cf_w, cf_s = vec(T, FP, "PPWNDON,All"), vec(T, FP, "PPPVT,All")
    inv_w, inv_s = vec(T, INV, "PPWNDON"), vec(T, INV, "PPPVT")
    up_w, up_s, lo_s = vec(T, TECHO, "PPWNDON"), vec(T, TECHO, "PPPVT"), vec(T, PISO, "PPPVT")
    lo_w = vec(T, PISO, "PPWNDON")
    ex_w, ex_s = capacidad_2025("PPWNDON"), capacidad_2025("PPPVT")
    ex_s_sin = ex_s - inv_s[IX[2025]]               # sin los 67 MW de 2025 del condicional
    pot = lambda ex, inv, cf, af: potencial(ex, inv, cf, af)
    techo = lambda base_v, p: desde(base_v, 2026, lambda a: max(p[IX[a]], base_v[IX[a]]))

    # ---- A: NDC financiada
    pA1 = pot(ex_w, inv_w, cf_w, af_w)
    pA3 = pot(ex_s, inv_s, cf_s, af_s)
    txt_w = ("Techo de generación eólica igual al potencial de la capacidad existente de CONA6 (%.2f MW en 2025, K11) "
             "más la inversión mínima del condicional (230,4 MW en 2026-2030), con el factor de planta de CONA6 "
             "(%.4f, medio observado 2023-2025) y el factor de disponibilidad de 0,95: %.4f PJ en 2026 y %.4f PJ "
             "desde 2030. Nunca baja del techo de CONA6.") % (ex_w * 1000, cf_w[IX[2026]], max(pA1[IX[2026]], up_w[IX[2026]]),
                                                             max(pA1[IX[2035]], up_w[IX[2035]]))
    P.append(Palanca("A1", "A", TECHO, "PPWNDON", techo(up_w, pA1), "PJ/año", txt_w, "EN10, EN12", "Por demostrar",
                     FUENTE_P3, familia="Eólica"))
    txt_s = ("Techo de generación solar igual al potencial de la capacidad existente de CONA6 (%.2f MW en 2025, K12, "
             "con los 67 MW del condicional de 2025) más la inversión mínima del condicional de 2026-2034, con el factor "
             "de planta de CONA6 (%.4f) y el factor de disponibilidad de 0,95; donde el techo de CONA6 es mayor, rige "
             "el de CONA6.") % (ex_s * 1000, cf_s[IX[2026]])
    P.append(Palanca("A3", "A", TECHO, "PPPVT", techo(up_s, pA3), "PJ/año", txt_s, "EN20, EN25", "Por demostrar",
                     FUENTE_P3, familia="Solar"))
    geo_lo, geo_up = vec(T, PISO, "PPGEO"), vec(T, TECHO, "PPGEO")
    geo_tope, geo_inv = vec(T, TOPE, "PPGEO"), vec(T, INV, "PPGEO")
    P.append(Palanca("A5", "A", PISO, "PPGEO", desde(geo_lo, GEO_ANIO, GEO_PJ), "PJ/año",
                     "Piso de generación geotérmica de 1,27 PJ desde 2029: Chachimbiro (50 MW) despacha desde su "
                     "entrada adelantada.", "EN08", "Firme", FUENTE_PEG, familia="Geotermia"))
    P.append(Palanca("A7", "A", TOPE, "PPGEO", desde(geo_tope, GEO_ANIO, GEO_TOPE), "GW",
                     "Tope de capacidad geotérmica de 0,5 GW desde 2029 en lugar de 2032: habilita la entrada "
                     "adelantada de Chachimbiro.", "EN08", "Firme", FUENTE_PEG, familia="Geotermia"))
    inv_a8 = [GEO_GW if a == GEO_ANIO else 0.0 for a in ANIOS]
    P.append(Palanca("A8", "A", INV, "PPGEO", inv_a8, "GW",
                     "Inversión mínima geotérmica de 0,051 GW en 2029 en lugar de tres módulos de 0,017 GW en "
                     "2032-2034: Chachimbiro entra en la fecha del Plan de Expansión 2023-2032.", "EN08", "Firme",
                     FUENTE_PEG, familia="Geotermia"))
    P.append(Palanca("A9", "A", TECHO, "PPGEO", desde(geo_up, GEO_ANIO, GEO_PJ), "PJ/año",
                     "Techo de generación geotérmica de 1,27 PJ desde 2029 en lugar de la rampa 2032-2034 del "
                     "condicional.", "EN08", "Firme", FUENTE_PEG, familia="Geotermia"))

    # ---- B: ambición ampliada
    inv_b7 = [v + (B_EOL_BLOQUE if a in B_EOL_ANIOS else 0.0) for a, v in zip(ANIOS, inv_w)]
    inv_b8 = [v + (B_SOL_BLOQUE if a in B_SOL_ANIOS else 0.0) for a, v in zip(ANIOS, inv_s)]
    pB4 = pot(ex_w, inv_b7, cf_w, af_w)
    pB2 = pot(ex_s, inv_b8, cf_s, af_s)
    P.append(Palanca("B2", "B", TECHO, "PPPVT", techo(up_s, pB2), "PJ/año",
                     "Techo de generación solar igual al potencial de la capacidad existente de CONA6, la inversión del "
                     "condicional y los 1.584 MWp de CELEC EP y AFD (B8), con el factor de planta de CONA6 y el factor "
                     "de disponibilidad de 0,95: %.4f PJ en 2035." % techo(up_s, pB2)[IX[2035]],
                     "fuera de la NDC", "Por demostrar", "Regla de construcción del A aplicada a la capacidad del B",
                     familia="Solar"))
    P.append(Palanca("B4", "B", TECHO, "PPWNDON", techo(up_w, pB4), "PJ/año",
                     "Techo de generación eólica igual al potencial de la capacidad existente de CONA6, la inversión del "
                     "condicional y los 585 MW del Plan de Expansión (B7), con el factor de planta de CONA6 y el factor "
                     "de disponibilidad de 0,95: %.4f PJ desde 2032." % techo(up_w, pB4)[IX[2035]],
                     "fuera de la NDC", "Por demostrar", "Regla de construcción del A aplicada a la capacidad del B",
                     familia="Eólica"))
    cf_b10 = r4(B9_FP_DECLARADO / af_b[IX[2031]])
    techo_bio = r4(B9_MW / 1000.0 * cf_b10 * af_b[IX[2031]] * CAU)
    bim_up = vec(T, TECHO, "PPBIM")
    P.append(Palanca("B6", "B", TECHO, "PPBIM", desde(bim_up, 2031, techo_bio), "PJ/año",
                     "Techo de generación de biomasa igual al potencial de los 39,3 MW nuevos (B9) con el factor de "
                     "planta declarado de 0,56 (CF del modelo %.4f con AF 0,95): %.4f PJ (%.1f GWh) al año desde 2031."
                     % (cf_b10, techo_bio, techo_bio / 0.0036), "B-BIO", "Por demostrar",
                     "ESIN-CELEC EP (2020), portafolio de biomasa de 39,3 MW en cuatro plantas de palma y cascarilla; "
                     "factor de planta 0,56 (lámina 31)", familia="Biomasa y biogás"))
    P.append(Palanca("B7", "B", INV, "PPWNDON", inv_b7, "GW",
                     "Inversión mínima eólica: la de CONA6 más 585 MW del Plan de Expansión 2023-2032 en cinco bloques de "
                     "117 MW entre 2028 y 2032 (calendario supuesto). Conserva la historia de CONA6 (sin la eólica "
                     "ficticia de 2013-2015 del Modelo CLG).", "B-EOL", "Por demostrar", FUENTE_PEG, familia="Eólica"))
    P.append(Palanca("B8", "B", INV, "PPPVT", inv_b8, "GW",
                     "Inversión mínima solar: la de CONA6 más 1.584 MWp del portafolio de CELEC EP y AFD en ocho bloques "
                     "de 198 MW entre 2028 y 2035 (calendario supuesto).", "B-SOL", "Por demostrar",
                     "CELEC EP y AFD-TRACTEBEL (2024), 1.584 MWp", familia="Solar"))
    P.append(Palanca("B9", "B", INV, "PPBIM", [B9_MW / 1000.0 if a == 2031 else 0.0 for a in ANIOS], "GW",
                     "Inversión mínima de biomasa de 0,0393 GW en 2031: el portafolio de 39,3 MW de ESIN-CELEC EP "
                     "(2020). CONA6 no tiene capacidad de PPBIM desde 2011 (K05).", "B-BIO", "Por demostrar",
                     "ESIN-CELEC EP (2020), 39,3 MW en cuatro plantas", builtin="NO", familia="Biomasa y biogás"))
    P.append(Palanca("B10", "B", FP, "PPBIM", desde(vec(T, FP, "PPBIM,All"), 2031, cf_b10), "fracción",
                     "Factor de planta del modelo para la biomasa nueva desde 2031: %.4f, el 0,56 declarado dividido por el "
                     "factor de disponibilidad de 0,95 (CONA6 trae 0,4592 sin capacidad de PPBIM)." % cf_b10, "B-BIO",
                     "Por demostrar", "ESIN-CELEC EP (2020), lámina 31", familia="Biomasa y biogás"))

    # ---- C: ambición integral
    inv_c1 = [GEO_GW if a == GEO_ANIO else (TUFINO_GW if a == TUFINO_ANIO else 0.0) for a in ANIOS]
    rend_geo = GEO_PJ / GEO_GW
    up_c2 = desde(geo_up, GEO_ANIO, lambda a: GEO_PJ + (TUFINO_GW * rend_geo if a >= TUFINO_ANIO else 0.0))
    P.append(Palanca("C-GEO1", "C", INV, "PPGEO", inv_c1, "GW",
                     "Inversión mínima geotérmica: Chachimbiro 0,051 GW en 2029 (como A8) más 0,138 GW del Proyecto "
                     "Geotérmico Binacional Tufiño-Chiles-Cerro Negro en 2035 (entrada supuesta).", "C-GEO",
                     "Por demostrar", "Primicias (2025) y ThinkGeoEnergy (2025): 138 MW; año 2035 supuesto",
                     familia="Geotermia"))
    P.append(Palanca("C-GEO2", "C", TECHO, "PPGEO", up_c2, "PJ/año",
                     "Techo de generación geotérmica: 1,27 PJ de Chachimbiro desde 2029 (como A9) más %.4f PJ de Tufiño "
                     "en 2035 con el mismo rendimiento por GW (%.2f PJ/GW)." % (TUFINO_GW * rend_geo, rend_geo),
                     "C-GEO", "Por demostrar", "Primicias (2025) y ThinkGeoEnergy (2025): 138 MW; año 2035 supuesto",
                     familia="Geotermia"))
    # demanda del Plan de Acción: razón C/NDCCON del sandbox, año a año, aplicada a CONA6 desde 2026
    SN, _ = tabla("NDCCON", SANDBOX)
    SC, _ = tabla("C", SANDBOX)
    ind = ["E5INDBIM", "E5INDCOA", "E5INDCOK", "E5INDDSL", "E5INDELE", "E5INDFIR", "E5INDFOI", "E5INDGSL",
           "E5INDLPG", "E5INDNGS", "E5INDOID", "E5INDRDF"]
    tro = [("C-EM1", "E5TROGSL", "gasolina carretera de pasajeros"), ("C-EM2", "E5TRODSL", "diésel carretero de pasajeros"),
           ("C-EM3", "E5TROELE", "electricidad del transporte carretero de pasajeros")]
    dem = "SpecifiedAnnualDemand"

    def por_razon(comb):
        b, n, c = vec(T, dem, comb), SN[(dem, comb)][1], SC[(dem, comb)][1]
        return [b[i] if a < 2026 or n[i] == 0 else b[i] * c[i] / n[i] for i, a in enumerate(ANIOS)]
    for k, comb in enumerate(ind, 1):
        P.append(Palanca("C-EI%d" % k, "C", dem, comb, por_razon(comb), "PJ",
                         "Eficiencia energética industrial del Plan de Acción del PLANMICC: %s con 2,45 %% menos en 2035, "
                         "rampa lineal 2026-2035 (razón del sandbox aplicada a la demanda de CONA6)." % comb, "C-EI",
                         "Por demostrar", "PLANMICC, Plan de Acción Fase I (MAE, 2026), tabla 36; modelo OSeMOSYS del "
                         "Plan de Acción, hoja Demandas_PJ", hoja="Efficiency", clave=comb, familia="Industria"))
    for pid, comb, nom in tro:
        P.append(Palanca(pid, "C", dem, comb, por_razon(comb), "PJ",
                         "Electromovilidad del Plan de Acción del PLANMICC: %s; en 2035 el combustible del transporte "
                         "cae 2,8 %% del tendencial y la electricidad añadida es el 1,4 %%, descontado lo que ya hace el "
                         "condicional; rampa lineal 2026-2035 (razón del sandbox aplicada a CONA6)." % nom, "C-EM",
                         "Por demostrar", "PLANMICC, Plan de Acción Fase I (MAE, 2026), tablas 35 y 36", hoja="Efficiency",
                         clave=comb, familia="Transporte"))

    # ---- contrafactual del A (SIN)
    p_sin_w = pot(ex_w, [0.0] * 26, cf_w, af_w)
    p_sin_s = pot(ex_s_sin, [0.0] * 26, cf_s, af_s)
    P.append(Palanca("SIN1", "SIN", TECHO, "PPWNDON", desde(up_w, 2026, lambda a: p_sin_w[IX[a]]), "PJ/año",
                     "Contrafactual: techo eólico igual al potencial de la capacidad existente de CONA6 (%.2f MW), sin "
                     "la capacidad nueva del condicional." % (ex_w * 1000), "contrafactual", "no aplica",
                     "regla de A1 sin inversión", familia="Eólica"))
    P.append(Palanca("SIN2", "SIN", TECHO, "PPPVT", desde(up_s, 2026, lambda a: p_sin_s[IX[a]]), "PJ/año",
                     "Contrafactual: techo solar igual al potencial de la capacidad existente de CONA6 sin el condicional "
                     "(%.2f MW)." % (ex_s_sin * 1000), "contrafactual", "no aplica", "regla de A3 sin inversión",
                     familia="Solar"))
    P.append(Palanca("SIN4", "SIN", INV, "PPWNDON", desde(inv_w, 2026, 0.0), "GW",
                     "Contrafactual: sin la inversión eólica del condicional (230,4 MW de 2026-2030 a cero).",
                     "contrafactual", "no aplica", "retiro de la inversión del condicional", familia="Eólica"))
    P.append(Palanca("SIN5", "SIN", INV, "PPPVT", desde(inv_s, 2025, 0.0), "GW",
                     "Contrafactual: sin la inversión solar del condicional (700 MW de 2025-2034 a cero).",
                     "contrafactual", "no aplica", "retiro de la inversión del condicional", familia="Solar"))
    P.append(Palanca("SIN6", "SIN", PISO, "PPPVT", desde(lo_s, 2026, 0.98 * p_sin_s[IX[2026]]), "PJ/año",
                     "Contrafactual: piso solar igual a 0,98 por la generación de la capacidad existente (regla de K12c "
                     "sin el incremento del condicional).", "contrafactual", "no aplica", "regla de K12c",
                     familia="Solar"))
    P.append(Palanca("SIN7", "SIN", TECHO, "PPGEO", [0.0] * 26, "PJ/año", "Contrafactual: sin techo de generación "
                     "geotérmica (cero).", "contrafactual", "no aplica", "retiro de la geotermia del condicional",
                     familia="Geotermia"))
    P.append(Palanca("SIN8", "SIN", TOPE, "PPGEO", [0.0] * 26, "GW", "Contrafactual: sin capacidad geotérmica (tope "
                     "cero).", "contrafactual", "no aplica", "retiro de la geotermia del condicional", familia="Geotermia"))
    P.append(Palanca("SIN9", "SIN", INV, "PPGEO", [0.0] * 26, "GW", "Contrafactual: sin inversión geotérmica (cero).",
                     "contrafactual", "no aplica", "retiro de la geotermia del condicional", familia="Geotermia"))

    # ---- estados de riesgo y sensibilidad
    razon, _, _ = razones_secas()
    s_cc = participacion_coca_codo()
    f_solo, f_der = 1.0 - s_cc, 1.0 - CC_DERRATEO * s_cc
    pars_h = [(FP, "%s,All", "fracción", "R01"), (PISO, "%s", "PJ/año", "R02"), (TECHO, "%s", "PJ/año", "R03")]

    def hidro(grupo, prefijo, clases, factor, descr, fuente, n0=1):
        n = n0
        for par, fmt, uni, _ in pars_h:
            for cl in clases:
                base_v = vec(T, par, fmt % cl)
                v = [base_v[i] * factor(cl, a) for i, a in enumerate(ANIOS)]
                P.append(Palanca("%s%02d" % (prefijo, n), grupo, par, cl, v, uni, descr(cl), "riesgo", "", fuente,
                                 clave=fmt % cl, familia="Hidroeléctrica"))
                n += 1
    hidro("CCOFF", "R-CC", ["PPHRORAB"], lambda cl, a: 0.0 if a >= 2029 else 1.0,
          lambda cl: "Salida de la clase PPHRORAB (Coca Codo y Sopladora) desde 2029: factor de planta y límites de "
                     "actividad en cero. Cota superior del riesgo de Coca Codo.", FUENTE_CC_OFF)
    hidro("CCSOLO", "R-CS", ["PPHRORAB"], lambda cl, a: f_solo if a >= 2029 else 1.0,
          lambda cl: "Salida de Coca Codo sola desde 2029: factor de planta y límites de actividad de PPHRORAB por "
                     "%.4f, la participación de Sopladora en la energía de la clase en 2025." % f_solo,
          FUENTE_CC_OFF + "; " + FUENTE_CC)
    hidro("CCDER", "R-CD", ["PPHRORAB"], lambda cl, a: f_der if a >= 2027 else 1.0,
          lambda cl: "Derrateo del 20 %% de Coca Codo desde 2027, medido por su participación en la energía de la clase "
                     "(%.4f): PPHRORAB por %.4f." % (s_cc, f_der),
          "Supuesto del análisis: derrateo crónico del 20 %; " + FUENTE_CC)
    hidro("ANOSECO", "R-AS", HIDRO, lambda cl, a: razon[cl] if a >= 2026 else 1.0,
          lambda cl: "Año seco severo: valor de CONA6 por la razón de caudales de 2024 de la clase (%.4f) en todos los "
                     "años desde 2026. Cota." % razon[cl], FUENTE_SECO)
    hidro("ANOSEC3", "R-A3", HIDRO, lambda cl, a: razon[cl] if a in ANIOS_SECOS_3 else 1.0,
          lambda cl: "Un año seco cada tres: valor de CONA6 por la razón de caudales de 2024 de la clase (%.4f) en 2026 "
                     "(año seco observado), 2029, 2032 y 2035; años normales en los demás." % razon[cl], FUENTE_SECO)
    hidro("ANOSECM", "R-AM", HIDRO, lambda cl, a: SECO_MODERADO if a >= 2026 else 1.0,
          lambda cl: "Año seco moderado: toda la hidroelectricidad por 0,908 desde 2026 (caída de 9,2 % de 2024).",
          "BEN 2025 (MAE), tabla 2.20, p. 66: generación hidroeléctrica 2024 frente a 2023")
    hidro("CCOFF_ANOSECO", "R-CA", ["PPHRORAB"],
          lambda cl, a: 0.0 if a >= 2029 else (razon[cl] if a >= 2026 else 1.0),
          lambda cl: "Coca Codo y Sopladora fuera desde 2029 sobre el año seco severo (se aplica después del grupo "
                     "ANOSECO).", FUENTE_CC_OFF + "; " + FUENTE_SECO)
    for n, tec in enumerate(("PPFOI", "PPDSL", "PPNGS"), 1):
        origen_piso = ("piso K16, mínimo observado 2015-2025" if tec == "PPFOI"
                       else "piso del Modelo CLG, valor de 2022 congelado, sin fuente")
        P.append(Palanca("S%d" % n, "FOI", PISO, tec, desde(vec(T, PISO, tec), 2026, 0.0), "PJ/año",
                         "Sin el piso térmico desde 2026 (%s): piso de %s en cero; el despacho lo decide el modelo "
                         "dentro de su tope K15b." % (origen_piso, tec), "ninguna",
                         "No es actividad", "Decisiones del 12 y 13-10-2026: sensibilidad de los pisos "
                         "(familia FOI); A2_Output_Params/NDCCON y K16 de correcciones_K.csv", familia="Térmica"))
    for n, par in enumerate((PISO, TECHO), 1):
        P.append(Palanca("I%d" % n, "IMP0", par, "IMPELE", desde(vec(T, par, "IMPELE"), 2026, 0.0), "PJ/año",
                         "Importación de electricidad desde Colombia en cero desde 2026 (%s de IMPELE)."
                         % ("piso" if par == PISO else "techo"), "ninguna", "No es actividad",
                         "Resolución MME 40064 y Circular CREG 343; CENACE 2026 (S26-077): menos de 1 GWh por mes",
                         familia="Importación"))

    # ---- estado ACT anclado a 2026 (decisión del 13-10-2026)
    act = act_2026(T, s_cc)
    f26 = act["factor_2026"]
    hidro("ACTH", "R-AT", HIDRO,
          lambda cl, a: f26[cl] if a == 2026 else (razon[cl] if a in ANIOS_SECOS_ACT else 1.0),
          lambda cl: "Situación actual: valor de CONA6 por %.4f en 2026 (hidrología observada de 2026 de la clase; %s) y "
                     "por la razón de caudales de 2024 (%.4f) en 2029, 2032 y 2035, un año seco cada tres; años normales "
                     "en los demás." % (f26[cl], next(x["origen"] for x in act["clases"] if x["clase"] == cl), razon[cl]),
          FUENTE_ACT + "; " + FUENTE_SECO)
    imp26 = act["importacion_2026_PJ"]
    for n, par in enumerate((PISO, TECHO), 3):
        P.append(Palanca("I%d" % n, "IMPACT", par, "IMPELE", desde(vec(T, par, "IMPELE"), 2026,
                                                                   lambda a: imp26 if a == 2026 else 0.0), "PJ/año",
                         "Importación desde Colombia de la situación actual (%s de IMPELE): en 2026, la importada antes del "
                         "07-09-2026 (%.2f GWh = %.4f PJ: CENACE enero a agosto y XM 1 a 6 de septiembre); cero desde 2027, "
                         "supuesto de estrés declarado." % ("piso" if par == PISO else "techo",
                                                             act["observado"]["importacion_2026_antes_07_09_GWh"], imp26),
                         "ninguna", "No es actividad", FUENTE_ACT + "; Resolución MME 40064 y Circular CREG 343",
                         familia="Importación"))
    m_dem = act["multiplicador_demanda"]
    c_dem = {"E5INDELE": por_razon("E5INDELE"), "E5TROELE": por_razon("E5TROELE")}      # demanda del Escenario C
    for grupo, pref, persiste, base_c in (("DEMACT", "D-A", False, False), ("DEMACT_C", "D-C", False, True),
                                          ("DEMACTDEM", "D-P", True, False), ("DEMACTDEM_C", "D-Q", True, True)):
        for k, comb in enumerate(DEMANDA_ELE, 1):
            b = c_dem[comb] if base_c and comb in c_dem else vec(T, "SpecifiedAnnualDemand", comb)
            v = [b[i] * m_dem if (a == 2026 or (persiste and a > 2026)) else b[i] for i, a in enumerate(ANIOS)]
            texto = ("el nivel de 2026 vuelve a la trayectoria del Modelo CLG desde 2027 (choque transitorio)" if not persiste
                     else "el nivel de 2026 persiste y desde ahí crece con las tasas del Modelo CLG")
            P.append(Palanca("%s%d" % (pref, k), grupo, "SpecifiedAnnualDemand", comb, v, "PJ",
                             "Demanda eléctrica observada de 2026 (%s%s): %s por %.6f en %s, para que la demanda eléctrica "
                             "final de 2026 sea la de CONA6 en 2025 por el crecimiento observado de enero a agosto "
                             "(%.4f); %s. Supuesto declarado." % (comb, ", sobre la demanda del Escenario C" if base_c and
                                                                   comb in c_dem else "", comb, m_dem,
                                                                   "2026" if not persiste else "2026-2035",
                                                                   act["crecimiento_demanda_2026"], texto),
                             "ninguna", "No es actividad", FUENTE_ACT, hoja="Efficiency", clave=comb, familia="Demanda"))

    # ---- respaldo de emergencia (decisión del 13-10-2026): estados de riesgo, incluidos ACT y ACT-DEM
    tope_foi, cap_foi, fij_foi = vec(T, TOPE, "PPFOI"), vec(T, "CapitalCost", "PPFOI"), vec(T, "FixedCost", "PPFOI")
    rot = "Respaldo de emergencia (supuesto de cierre, como la térmica arrendada que el país contrata hoy): "
    P.append(Palanca("RE1", "RESP", TOPE, "PPFOI", desde(tope_foi, 2026, 99999.0), "GW",
                     rot + "tope de capacidad de PPFOI liberado desde 2026 (99999, el valor por omisión) en lugar del tope "
                     "K15b (potencia existente más expansión declarada).", "ninguna", "No es actividad", FUENTE_RESP,
                     familia="Térmica"))
    P.append(Palanca("RE2", "RESP", "CapitalCost", "PPFOI", desde(cap_foi, 2026, lambda a: cap_foi[IX[a]] * CARO),
                     "costo del modelo", rot + "CapitalCost de PPFOI 2026-2035 por 1.000: el modelo construye solo la "
                     "potencia sin la cual no cubre la demanda.", "ninguna", "No es actividad", FUENTE_RESP,
                     familia="Térmica"))
    P.append(Palanca("RE3", "RESP", "FixedCost", "PPFOI", desde(fij_foi, 2026, lambda a: fij_foi[IX[a]] * CARO),
                     "costo del modelo", rot + "FixedCost de PPFOI 2026-2035 por 1.000.", "ninguna", "No es actividad",
                     FUENTE_RESP, familia="Térmica"))
    P.append(Palanca("RE4", "RESP", TOPE, "STOELE", desde(vec(T, TOPE, "STOELE"), 2026, 0.0), "GW",
                     rot + "STOELE cerrado en 2026-2035 (tope de capacidad cero): es un atajo sin correlato físico que el "
                     "modelo usa cuando falta potencia y que subestima las emisiones.", "ninguna", "No es actividad",
                     FUENTE_RESP, familia="Almacenamiento"))

    # ---- puentes V7: valores exactos del sandbox del 08-10-2026 desde 2026, sobre la historia de CONA6
    P += puentes_v7(T, SB)
    return P, {"razon_seca": razon, "participacion_coca_codo_2025": s_cc, "factor_ccsolo": f_solo,
               "factor_ccder": f_der, "cf_b10": cf_b10, "techo_b6": techo_bio, "existente_eolica_GW": ex_w,
               "existente_solar_GW": ex_s, "existente_solar_sin_condicional_GW": ex_s_sin,
               "act_2026": act, "respaldo_multiplicador_costo": CARO}


def puentes_v7(T, SB):
    sbn, _ = tabla("NDCCON", SANDBOX)
    doc = {p["id"]: p for p in leer_hoja_sandbox("Art6_Palancas")}
    salida = []
    for fam, grupos in (("A", ("A",)), ("B", ("B",)), ("C", ("C",))):
        datos, _ = tabla(fam, SANDBOX)
        for pid, p in doc.items():
            if p["grupo"] not in grupos:
                continue
            par, cl = p["parametro"], str(p["clave"])
            hoja = "Efficiency" if par == "SpecifiedAnnualDemand" else "Electrical"
            if (par, cl) not in datos:
                raise ErrorGenerador(f"el datafile {fam} del sandbox no trae {par} {cl} ({pid})")
            vs = datos[(par, cl)][1]
            vn = sbn.get((par, cl), (None, [None] * 26))[1]
            vb = T[(par, cl)][1] if (par, cl) in T else [0.0] * 26
            v = [vs[i] if a >= 2026 and (vn[i] is None or abs(vs[i] - vn[i]) > 1e-9) else vb[i]
                 for i, a in enumerate(ANIOS)]
            builtin = "NO" if (par, cl) not in T else "YES"
            salida.append(Palanca(pid + "-V7", fam + "_V7", par, p["tecnologia"], v, p["unidad"],
                                  "Puente V7: valores de la palanca %s del sandbox PNUD (corrida del 08-10-2026) desde "
                                  "2026, sobre la historia de CONA6." % pid, p["iniciativa"], p["adicionalidad"],
                                  "Art6_calibracion/version_2026-10-08/Executables/%s_0/%s_0.txt (copia de solo lectura)" % (fam, fam),
                                  hoja=hoja, clave=cl, builtin=builtin, familia=p["familia"]))
    return salida


def leer_hoja_sandbox(nombre):
    wb = openpyxl.load_workbook(SANDBOX / "B1_Scenario_Config.xlsx", data_only=True, read_only=True)
    filas = list(wb[nombre].iter_rows(values_only=True))
    wb.close()
    cab = [str(x) if x is not None else "" for x in filas[0]]
    return [dict(zip(cab, f)) for f in filas[1:] if f and f[0] is not None]


# ------------------------------------------------------------------------------------------ escenarios
FAMILIAS = {"A": ("NDC financiada", "#0E7C6B", "A"), "B": ("Ambición ampliada", "#C08A2E", "A,B"),
            "C": ("Ambición integral", "#2F5597", "A,B,C")}
ESTADOS = [  # sufijo, nombre, grupo, grupos de palancas que se agregan, familias
    # RESP: respaldo de emergencia en los estados de riesgo (decisión del 13-10-2026); los escenarios A, B y
    # C intactos, las referencias y las sensibilidades tienen solución sin él
    ("CCOFF", "Salida de Coca Codo y Sopladora (clase completa, cota)", "Riesgo", "CCOFF,RESP", "ABC"),
    ("CCSOLO", "Salida de Coca Codo sola", "Riesgo", "CCSOLO,RESP", "ABC"),
    ("CCDER", "Derrateo de Coca Codo", "Riesgo", "CCDER,RESP", "ABC"),
    ("ANOSECO", "Año seco severo (cota)", "Riesgo", "ANOSECO,RESP", "ABC"),
    ("ANOSEC3", "Un año seco cada tres", "Riesgo", "ANOSEC3,RESP", "ABC"),
    ("ANOSECM", "Año seco moderado", "Riesgo", "ANOSECM,RESP", "ABC"),
    ("CCOFFANOSECO", "Salida de Coca Codo y año seco severo", "Riesgo", "ANOSECO,CCOFF_ANOSECO,RESP", "ABC"),
    ("FOI", "Sin pisos térmicos", "Sensibilidad", "FOI", "AB"),
    ("IMP0", "Sin importación desde Colombia", "Riesgo", "IMP0,RESP", "ABC"),
    # ACT y ACT-DEM (decisión del 13-10-2026): 2026 observado, un año seco cada tres desde 2029, importación
    # cero desde 2027 y respaldo; la demanda de 2026 vuelve a la del Modelo CLG desde 2027 (ACT) o persiste (ACTDEM). En la
    # familia C los grupos de demanda llevan la demanda del Escenario C (sufijo _C, ver escenarios())
    ("ACT", "Situación actual 2026 con choque de demanda transitorio", "Riesgo", "ACTH,IMPACT,DEMACT,RESP", "ABC"),
    ("ACTDEM", "Situación actual 2026 con demanda persistente (ACT-DEM)", "Riesgo", "ACTH,IMPACT,DEMACTDEM,RESP",
     "ABC"),
    ("GAS", "Sin la expansión a gas declarada (sobre CONA6G)", "Sensibilidad", None, "ABC"),
    ("SIN", "Sin sus proyectos (contrafactual)", "Contrafactual", None, "A"),
]
BASE_DE_ESTADO = {"GAS": BASE_GAS}          # estados construidos sobre otra referencia que CONA6
NARRATIVA = {
    "A": "El Ecuador financia con mercado de carbono una parte de su NDC condicional y esas actividades rinden lo que "
         "declaran: techos de generación iguales al potencial de la capacidad existente y declarada, con el factor de "
         "planta observado de la base corregida, y Chachimbiro adelantado a la fecha del Plan de Expansión. Se mide "
         "contra CONA6.",
    "B": "Sobre el A, un paquete fuera de la Segunda NDC entra como capacidad construida: la eólica del Plan de "
         "Expansión, el portafolio solar de CELEC EP con la AFD y 39,3 MW de biomasa del portafolio ESIN-CELEC EP.",
    "C": "Sobre el B, la demanda del Plan de Acción del PLANMICC (electromovilidad y eficiencia industrial) y la "
         "geotermia de Tufiño-Chiles-Cerro Negro en 2035.",
    "CCOFF": "la clase PPHRORAB (Coca Codo Sinclair y Sopladora) sale de operación desde 2029; cota superior.",
    "CCSOLO": "Coca Codo Sinclair sale de operación desde 2029 y Sopladora sigue; caso central del riesgo de la central.",
    "CCDER": "Coca Codo pierde el 20 % de su generación desde 2027 por sedimentos, medido por su participación en la "
             "energía de la clase.",
    "ANOSECO": "los caudales de 2024 se repiten todos los años desde 2026; cota.",
    "ANOSEC3": "un año seco cada tres, con 2026 como año seco observado (2026, 2029, 2032 y 2035); caso central del "
               "riesgo hidrológico.",
    "ANOSECM": "toda la hidroelectricidad cae 9,2 % desde 2026, como en 2024.",
    "CCOFFANOSECO": "los dos choques a la vez: año seco severo y salida de la clase de Coca Codo.",
    "FOI": "sin pisos térmicos desde 2026 (K16 en PPFOI; Modelo CLG en PPDSL y PPNGS); se mide contra CONA6P.",
    "IMP0": "sin importación desde Colombia desde 2026; las capas se miden contra CONA6I.",
    "ACT": "estrés con la situación actual anclada a 2026: en 2026, la demanda eléctrica observada (crecimiento de enero a "
           "agosto), la hidrología observada de cada clase, Coca Codo con su factor de planta de 2026 y la importación "
           "hasta el 06-09-2026; desde 2027, un año seco cada tres (razones de ANOSEC3 en 2029, 2032 y 2035) e importación "
           "cero como supuesto de estrés. La demanda vuelve a la trayectoria del Modelo CLG desde 2027 (choque "
           "transitorio, supuesto declarado)",
    "ACTDEM": "la situación actual de ACT con la demanda de 2026 persistente: el nivel observado de 2026 se mantiene y "
              "desde ahí crece con las tasas del Modelo CLG (supuesto declarado)",
    "GAS": "las palancas del escenario sobre CONA6G, sin la expansión a gas declarada (PPNGS en su piso CLG, G01); sus "
           "capas se miden contra CONA6G.",
    "SIN": "el mismo sistema del A sin la eólica, la solar y la geotermia nuevas del condicional; la diferencia con el "
           "A es el aporte atribuible por retiro.",
}


TEXTO_RESP = (" Con respaldo de emergencia (supuesto de cierre, como la térmica arrendada que el país contrata hoy): tope "
              "de PPFOI liberado desde 2026 con su potencia nueva cara (CapitalCost y FixedCost por 1.000) y STOELE "
              "cerrado en 2026-2035, de modo que solo se construye la potencia que falta.")


def escenarios(info=None):
    """Lista de escenarios nuevos: código, nombre, grupo, composición de palancas, color, narrativa y base."""
    act = (info or {}).get("act_2026")
    coma = lambda v, d=2: (("%." + str(d) + "f") % v).replace(".", ",")
    txt_act = "" if act is None else (
        " Valores de 2026: demanda eléctrica final %s PJ (CONA6 %s PJ; multiplicador %s), Coca Codo con factor de planta "
        "%s, hidro de las clases %s veces la de CONA6 en 2025 e importación de %s GWh."
        % (coma(act["demanda_ACT_2026_PJ"], 4), coma(act["demanda_CONA6_2026_PJ"], 4),
           coma(act["multiplicador_demanda"], 4), coma(act["fp_coca_codo_2026"], 4),
           coma(act["razon_hidro_modelo_2026_2025"], 4), coma(act["observado"]["importacion_2026_antes_07_09_GWh"])))
    E = [{"codigo": "CONA6P", "nombre": "CONA6 sin pisos térmicos", "grupo": "Referencia del estado",
          "palancas": "FOI", "color": "#6B4F3A", "base": BASE,
          "narrativa": "CONA6 sin pisos térmicos desde 2026 (PPFOI con el piso K16; PPDSL y PPNGS con el del Modelo "
                       "CLG). Referencia de la familia FOI."},
         {"codigo": "CONA6I", "nombre": "CONA6 sin importación", "grupo": "Referencia del estado",
          "palancas": "IMP0", "color": "#7F7F7F", "base": BASE,
          "narrativa": "CONA6 con la importación desde Colombia en cero desde 2026. Referencia de la familia IMP0."}]
    for fam, (nom, color, comp) in FAMILIAS.items():
        E.append({"codigo": fam, "nombre": nom, "grupo": "Escenario Art. 6", "palancas": comp, "color": color,
                  "narrativa": NARRATIVA[fam], "base": BASE})
        for suf, nom_e, grupo, extra, fams in ESTADOS:
            if fam not in fams:
                continue
            if suf == "SIN":
                comp_e = "SIN"
            elif extra is None:
                comp_e = comp
            else:
                comp_e = comp + "," + extra
            if fam == "C":                     # la demanda de ACT y ACT-DEM sobre la demanda del Escenario C
                comp_e = ",".join(g + "_C" if g in ("DEMACT", "DEMACTDEM") else g for g in comp_e.split(","))
            narr = NARRATIVA[suf] + ("." + txt_act if suf in ("ACT", "ACTDEM") else "")
            if "RESP" in comp_e.split(","):
                narr += TEXTO_RESP
            E.append({"codigo": f"{fam}-{suf}", "nombre": f"{nom} · {nom_e}", "grupo": grupo, "palancas": comp_e,
                      "color": "#8C8C8C" if suf == "SIN" else ("#6B4F3A" if suf in ("FOI", "GAS") else color),
                      "base": BASE_DE_ESTADO.get(suf, BASE),
                      "narrativa": f"{('Prueba de estrés' if grupo == 'Riesgo' else 'Variante')} sobre el Escenario "
                                   f"{fam}: {narr}"})
    for fam, (nom, color, comp) in FAMILIAS.items():
        comp_v = ",".join(g + "_V7" for g in comp.split(","))
        E.append({"codigo": f"{fam}-V7", "nombre": f"{nom} · Definición del 08-10-2026 (puente)", "grupo": "Puente",
                  "palancas": comp_v, "color": color, "base": BASE,
                  "narrativa": f"El Escenario {fam} con las palancas exactas del sandbox PNUD del 08-10-2026 (biomasa de "
                               f"35,6 MW y techos de esa fecha) sobre CONA6. Separa el efecto de la base del efecto de "
                               f"las decisiones nuevas en el caso intacto."})
    return E


def composicion(e, P):
    """{(hoja, parametro, clave): palanca}: los grupos se aplican en orden y manda el último."""
    final = {}
    for g in [x.strip() for x in e["palancas"].split(",") if x.strip()]:
        for p in P:
            if p.grupo == g:
                final[(p.hoja, p.parametro, p.clave)] = p
    return final


# ------------------------------------------------------------------------------------------ libro de B1
def fila_electrical(hoja, plantilla, esc, p):
    celdas = {k: v for k, v in plantilla.items()}
    pon = lambda col, tipo, texto: hoja.poner(celdas, col, tipo, texto)
    v = p.vector
    pon("Scenario", "s", esc)
    pon("Parameter", "s", p.parametro)
    pon("Tech_Set", "s", p.tecnologia)
    pon("Description Parameter-Set", "s", "[%s] %s" % (p.id, p.descripcion))
    pon("Built-in Parameter-Set", "s", p.builtin)
    pon("Linear", "s", "YES")
    pon("Exact_Years", "s", " ; ".join(str(a) for a in ANIOS[:-1]))
    pon("Exact_Values", "s", " ; ".join(repr(x) for x in v[:-1]))
    pon("y_ini", "n", "2034")
    pon("Milestone_Years", "n", "2035")
    pon("Milestone_Value", "n", repr(v[-1]))
    pon("Unit", "s", "PJ" if "PJ" in p.unidad else ("GW" if p.unidad == "GW" else
                                                     ("fraction" if p.unidad == "fracción" else p.unidad)))
    pon("Security_Multiplier", "n", "1")
    pon("Method", "s", METODO if p.builtin == "YES" else METODO_NO)
    return celdas


def fila_efficiency(hoja, plantilla, esc, p):
    celdas = {k: v for k, v in plantilla.items()}
    hoja.poner(celdas, "Scenario", "s", esc)
    hoja.poner(celdas, "Set", "s", p.clave)
    hoja.poner(celdas, "Description_Set", "s", "[%s] %s" % (p.id, p.descripcion))
    for a, x in zip(ANIOS, p.vector):
        hoja.poner(celdas, a, "n", repr(float(x)))
    return celdas


class HojaArt6:
    """Hoja del libro de B1 con las filas del Artículo 6 al final: las demás filas no se tocan."""

    def __init__(self, nombre, xml, compartidas, propios):
        self.nombre, self.xml = nombre, xml.decode("utf-8")
        filas = GR.leer_hoja(xml, compartidas)
        self.cabecera = {col: GR.nombre_cabecera(c) for col, c in filas[0][1].items() if c[0] is not None}
        self.col_de = {v: k for k, v in self.cabecera.items()}
        self.datos, self.propias = [], []
        for r, celdas in filas[1:]:
            a = celdas.get("A", (None, None, None))[1]
            if a in propios:
                self.propias.append(r)
            elif any(t is not None for t, _, _ in celdas.values()):
                if self.propias:
                    raise ErrorGenerador(f"{nombre}: la fila {r} ({a}) está después de las filas del Artículo 6")
                self.datos.append((r, celdas))
        self.ultima = self.datos[-1][0] if self.datos else 1

    filas_de = GR.Hoja.filas_de
    valor = GR.Hoja.valor
    poner = GR.Hoja.poner

    def xml_nuevo(self, nuevas):
        import re
        texto = self.xml
        for r in self.propias:
            texto, n = re.subn(rf'<row r="{r}"(?=[\s>/])[^>]*?(?:/>|>.*?</row>)', "", texto, count=1, flags=re.S)
            if n != 1:
                raise ErrorGenerador(f"{self.nombre}: no se encontró la fila {r}")
        r = self.ultima
        partes = []
        for celdas in nuevas:
            r += 1
            partes.append(GR.fila_xml(r, celdas))
        texto = texto.replace("</sheetData>", "".join(partes) + "</sheetData>")
        m = re.search(r'<dimension ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"/>', texto)
        if m:
            texto = texto.replace(m.group(0), f'<dimension ref="{m.group(1)}{m.group(2)}:{m.group(3)}{r}"/>')
        return texto.encode("utf-8")


def construir_libro(original, E, P, quitar=False, solo=None):
    """solo: códigos de los escenarios del Artículo 6 que quedan activos (Activated = YES) para una corrida parcial de
    B1; los demás se escriben igual con Activated = NO. Sin solo, todos activos (estado normal del libro)."""
    propios = {e["codigo"] for e in E}
    with zipfile.ZipFile(io.BytesIO(original)) as z:
        rutas = GR.rutas_hojas(z)
        comp = GR.leer_compartidas(z)
        hojas = {n: HojaArt6(n, z.read(rutas[n]), comp, propios) for n in ("Scenarios",) + HOJAS_FILAS}
    nuevas = {n: [] for n in ("Scenarios",) + HOJAS_FILAS}
    resumen = {}
    if not quitar:
        sc = hojas["Scenarios"]
        plant = sc.filas_de(BASE)
        if len(plant) != 1:
            raise ErrorGenerador("Scenarios: falta CONA6")
        bases = sorted({e.get("base", BASE) for e in E})
        for b in bases:
            if len(sc.filas_de(b)) != 1:
                raise ErrorGenerador(f"Scenarios: la base {b} no está una sola vez")
        for e in E:
            celdas = dict(plant[0])
            sc.poner(celdas, "Name", "s", e["codigo"])
            sc.poner(celdas, "Activated", "s", "YES" if solo is None or e["codigo"] in solo else "NO")
            sc.poner(celdas, "Base", "s", "NO")
            sc.poner(celdas, "Reference", "s", "based")
            sc.poner(celdas, "Based_On", "s", e.get("base", BASE))
            sc.poner(celdas, "Description", "s", "%s. %s" % (e["nombre"], e["narrativa"]))
            nuevas["Scenarios"].append(celdas)
        filas_base = {b: {n: hojas[n].filas_de(b) for n in HOJAS_FILAS} for b in bases}
        for e in E:
            base_filas = filas_base[e.get("base", BASE)]
            final = composicion(e, P)
            cuenta = {}
            for n in HOJAS_FILAS:
                h = hojas[n]
                clave_col = "Set" if n == "Efficiency" else "Tech_Set"
                tocadas = {(par, str(cl).split(",")[0]) for (hj, par, cl) in final if hj == n}
                copias = []
                for celdas in base_filas[n]:
                    k = (h.valor(celdas, "Parameter"), str(h.valor(celdas, clave_col)))
                    if n != "SmartGrid" and k in tocadas:
                        continue
                    c2 = dict(celdas)
                    h.poner(c2, "Scenario", "s", e["codigo"])
                    copias.append(c2)
                propias = []
                for (hj, par, cl), p in sorted(final.items(), key=lambda x: (x[0][1], x[0][2])):
                    if hj != n:
                        continue
                    if n == "Electrical":
                        plantilla = base_filas[n][0]
                        propias.append(fila_electrical(h, plantilla, e["codigo"], p))
                    else:
                        plantilla = base_filas[n][0]
                        propias.append(fila_efficiency(h, plantilla, e["codigo"], p))
                nuevas[n] += copias + propias
                cuenta[n] = {"copias_de_la_base": len(copias), "base": e.get("base", BASE), "palancas": len(propias)}
            resumen[e["codigo"]] = cuenta
    datos = {rutas[n]: hojas[n].xml_nuevo(nuevas[n]) for n in nuevas}
    salida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as zin, zipfile.ZipFile(salida, "w") as zout:
        for info in zin.infolist():
            d = datos.get(info.filename, zin.read(info.filename))
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = info.compress_type
            zi.external_attr = info.external_attr
            zi.create_system = info.create_system
            zout.writestr(zi, d)
    return salida.getvalue(), resumen


def comparar_libros(antes, despues, propios):
    """Las filas que no son del Artículo 6 quedan iguales (valores en caché)."""
    a = pd.read_excel(io.BytesIO(antes), sheet_name=None)
    d = pd.read_excel(io.BytesIO(despues), sheet_name=None)
    if list(a) != list(d):
        raise ErrorGenerador("el libro nuevo no tiene las mismas hojas")
    for n in a:
        col = a[n].columns[0] if len(a[n].columns) else None
        if col is None:
            continue
        x = a[n][~a[n][col].astype(str).isin(propios)].reset_index(drop=True)
        y = d[n][~d[n][col].astype(str).isin(propios)].reset_index(drop=True)
        pd.testing.assert_frame_equal(x, y, check_dtype=False, obj=f"{n} (filas que no son del Artículo 6)")


# ------------------------------------------------------------------------------------------ hojas documentales
def libro_documental(E, P, info):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def hoja(nombre, cab, filas):
        ws = wb.create_sheet(nombre)
        ws.append(cab)
        for f in filas:
            ws.append([f.get(c) for c in cab] if isinstance(f, dict) else list(f))

    hoja("Art6_LEEME", ["texto"], [[t] for t in LEEME])
    hoja("Art6_Roles", ["rol", "escenario", "descripcion"], ROLES)
    ref_E = [{"codigo": "NDCINC", "nombre": "NDC incondicional (Modelo CLG)", "grupo": "Modelo CLG", "palancas": None,
              "color": "#BFC3C7", "narrativa": "Escenario incondicional del Modelo Climate Lead Group, sin cambios."},
             {"codigo": "NDCCON", "nombre": "NDC condicional (Modelo CLG)", "grupo": "Modelo CLG", "palancas": None,
              "color": "#9AA3B5", "narrativa": "Escenario condicional del Modelo Climate Lead Group, sin cambios."},
             {"codigo": "INCA6", "nombre": "NDC incondicional corregida", "grupo": "Referencia corregida",
              "palancas": None, "color": "#8A8F94",
              "narrativa": "NDCINC del Modelo CLG con las correcciones K01-K16 (base congelada del 13-10-2026)."},
             {"codigo": "CONA6", "nombre": "NDC condicional corregida", "grupo": "Referencia corregida",
              "palancas": None, "color": "#4E5D7E",
              "narrativa": "NDCCON del Modelo CLG con las correcciones K01-K16 (base congelada del 13-10-2026). Base de "
                           "A, B y C."},
             {"codigo": BASE_GAS, "nombre": "CONA6 sin la expansión a gas declarada", "grupo": "Referencia del estado",
              "palancas": None, "color": "#9C6B30",
              "narrativa": "CONA6 con PPNGS en su piso CLG 2026-2035 (G01): sin la expansión a gas declarada del Anexo A. "
                           "Base y referencia de la familia GAS (sensibilidad de gas de la base congelada)."}]
    cab_e = ["codigo", "nombre", "grupo", "palancas", "color", "narrativa"]
    hoja("Art6_Escenarios", cab_e, ref_E + E)
    cab_p = ["grupo", "id", "tecnologia", "clave", "familia", "parametro", "unidad", "descripcion", "iniciativa",
             "adicionalidad", "fuente"]
    hoja("Art6_Palancas", cab_p, [p.doc() for p in P])
    hoja("Art6_Riesgos", ["grupo", "nombre", "desde", "mecanismo", "derivado_de", "validacion", "nota"],
         riesgos(info))
    razon, origen, detalle = razones_secas()
    hoja("Art6_AnoSeco", ["Clase", "Origen de la razón", "Razón aplicada", "Desde", "Años secos ANOSEC3"],
         [[c, origen[c], razon[c], 2026, "2026, 2029, 2032, 2035"] for c in HIDRO])
    # sandbox PNUD: sin la hoja Art6_AnoSeco_Centrales (filas por central)
    act = info["act_2026"]
    cab_a = ["clase", "actividad_CONA6_2025_PJ", "razon_2026_2025", "actividad_ACT_2026_PJ", "techo_CONA6_2026_PJ",
             "factor_2026", "fp_CONA6_2026", "fp_ACT_2026", "origen"]
    hoja("Art6_ACT2026", cab_a, act["clases"])
    filas_d = [dict(r) for r in csv.DictReader(io.open(ACT_OBSERVADO, encoding="utf-8-sig"))]
    filas_d += [{"clave": "multiplicador_demanda_2026", "valor": act["multiplicador_demanda"], "unidad": "razón",
                 "periodo": "2026", "fuente": "derivación propia: demanda E5*ELE de CONA6 2025 x crecimiento / CONA6 2026",
                 "nota": "ACT: solo 2026; ACT-DEM: desde 2026"},
                {"clave": "importacion_2026_PJ", "valor": act["importacion_2026_PJ"], "unidad": "PJ", "periodo": "2026",
                 "fuente": "importacion_2026_antes_07_09_GWh x 0,0036", "nota": "piso y techo de IMPELE en 2026"}]
    hoja("Art6_ACT2026_Datos", ["clave", "valor", "unidad", "periodo", "fuente", "ubicacion", "archivos", "md5",
                                "id_situacion", "nota"], filas_d)
    for nombre in ("Art6_Headroom", "Art6_Precios", "Art6_Adicionalidad", "Art6_Reporte", "Art6_Fuentes",
                   "Art6_Changelog"):
        filas = leer_hoja_sandbox(nombre)
        cab = list(filas[0].keys()) if filas else []
        filas = AJUSTES.get(nombre, lambda x: x)(filas)
        hoja(nombre, cab, filas)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# proyecciones del margen (Art6_Headroom, parámetro proyecciones; headroom_proyecciones.csv): la del Modelo CLG mide
# cada escenario frente a la referencia corregida (CONA6 o INCA6) y la corregida frente a NDCCON o NDCINC. Corrección
# del 13-10-2026 tras la verificación independiente del paso 15: las descripciones estaban intercambiadas.
ROLES = [["inc", "INCA6", "Referencia incondicional corregida: referencia de la proyección del Modelo CLG con la regla "
                          "prudente (lectura de autorización)"],
         ["con", "CONA6", "Referencia condicional corregida: base de A, B y C, referencia de sus capas y de la proyección "
                          "del Modelo CLG con la regla condicional"],
         ["inc_clg", "NDCINC", "Incondicional del Modelo CLG: referencia de la proyección corregida con la regla "
                               "prudente"],
         ["con_clg", "NDCCON", "Condicional del Modelo CLG: referencia de la proyección corregida con la regla "
                               "condicional"],
         ["tendencial_clg", "BAU1NDC", "Tendencial retroactivo del Modelo CLG que usa la agregación nacional"],
         ["ref_estado:FOI", "CONA6P", "Referencia de la familia FOI (CONA6 sin pisos térmicos: K16 en PPFOI; Modelo CLG "
                                      "en PPDSL y PPNGS)"],
         ["ref_estado:IMP0", "CONA6I", "Referencia de la familia IMP0 (CONA6 sin importación)"],
         ["ref_estado:GAS", BASE_GAS, "Referencia y base de la familia GAS (CONA6 sin la expansión a gas declarada, G01)"]]

LEEME = [
    "Escenarios del Artículo 6 sobre la base corregida CONA6 (ECU_NDC, rama art6-escenarios)",
    "Las filas que B1 aplica (Scenarios, Electrical, SmartGrid y Efficiency) están en Energy/B1_Scenario_Config.xlsx, al "
    "final de cada hoja; las genera Art6_calibracion/generar_escenarios_art6.py desde el datafile y la salida de CONA6 "
    "de la corrida congelada del 12-10-2026. Este libro guarda la documentación (Art6_*), que B1 no lee.",
    "Referencias: INCA6 y CONA6 (base corregida congelada del 13-10-2026, K01-K16, G01 y RE01); NDCINC y NDCCON del "
    "Modelo Climate Lead Group (proyección del Modelo CLG); CONA6P, CONA6I y CONA6G, referencias de las familias FOI, "
    "IMP0 y GAS (Art6_Roles).",
    "Escenarios: A, B y C basados en CONA6, con sus estados por el patrón <familia>-<estado>; contrafactual A-SIN; "
    "puentes A-V7, B-V7 y C-V7 con las palancas del sandbox del 08-10-2026 sobre CONA6; familia GAS (A-GAS, B-GAS y "
    "C-GAS) con las mismas palancas sobre CONA6G.",
    "Los estados de riesgo (CCOFF, CCSOLO, CCDER, ANOSECO, ANOSEC3, ANOSECM, CCOFFANOSECO, IMP0, ACT y ACTDEM) llevan "
    "el respaldo de emergencia (grupo RESP): tope de PPFOI liberado desde 2026, CapitalCost y FixedCost de PPFOI por "
    "1.000 y STOELE cerrado en 2026-2035. Es un supuesto de cierre, como la térmica arrendada que el país contrata hoy "
    "(decisión del 13-10-2026). " + NOTA_RE01,
    "ACT y ACTDEM (ACT-DEM) están anclados a 2026 (decisión del 13-10-2026): demanda eléctrica, hidrología "
    "por clase, Coca Codo e importación observadas en 2026 (hojas Art6_ACT2026 y Art6_ACT2026_Datos); desde 2027, un año "
    "seco cada tres e importación cero. En ACT la demanda vuelve a la del Modelo CLG desde 2027; en ACTDEM el nivel de "
    "2026 persiste. Las dos variantes de demanda son supuestos declarados.",
    "Cada palanca empieza su descripción en Electrical o Efficiency con su código entre corchetes ([A1]). Las "
    "palancas de Electrical son vectores completos 2010-2035 que conservan la historia de CONA6.",
    "Para cambiar una palanca: editar generar_escenarios_art6.py (nunca el libro a mano), correr el generador, "
    "tests/validar_libro.py, B1 y las pruebas, y anotar el cambio en Art6_Changelog.",
    "Para volver a generar la base (generar_referencias.py) hay que retirar antes estas filas con "
    "generar_escenarios_art6.py --quitar: generar_referencias exige que las filas de los clones vayan al final.",
]


def riesgos(info):
    s, fs, fd = info["participacion_coca_codo_2025"], info["factor_ccsolo"], info["factor_ccder"]
    return [
        ["CCOFF", "Salida de la clase de Coca Codo y Sopladora", 2029, "Factor de planta y límites de actividad de "
         "PPHRORAB en cero desde 2029.", FUENTE_CC_OFF, "verificar.py: cada estado difiere de su escenario solo en sus "
         "palancas.", "Cota: retira la clase completa (1.987 MW: Coca Codo 1.500 y Sopladora 487)."],
        ["CCSOLO", "Salida de Coca Codo sola", 2029, "PPHRORAB por %.4f desde 2029 (sale la participación de Coca Codo en "
         "la energía de la clase, %.4f)." % (fs, s), FUENTE_CC, "verificar.py", "Caso central del riesgo de la central."],
        ["CCDER", "Derrateo del 20 % de Coca Codo", 2027, "PPHRORAB por 1 - 0,20 x %.4f = %.4f desde 2027." % (s, fd),
         FUENTE_CC, "verificar.py", "Medido por participación en energía (decisión del 12-10-2026)."],
        ["ANOSECO", "Año seco severo con la hidrología de 2024", 2026, "Factor de planta, piso y techo de las 11 clases "
         "hidroeléctricas por la razón de caudales de 2024 de cada clase, en todos los años desde 2026.", FUENTE_SECO,
         "verificar.py: cociente frente a CONA6 igual a la razón de Art6_AnoSeco.", "Cota; desde 2026 porque la "
         "historia ya es observada."],
        ["ANOSEC3", "Un año seco cada tres", 2026, "Las mismas razones solo en 2026, 2029, 2032 y 2035.", FUENTE_SECO,
         "verificar.py", "Caso central del riesgo hidrológico (decisión del 12-10-2026); 2026 es año seco observado."],
        ["ANOSECM", "Año seco moderado", 2026, "Toda la hidroelectricidad por 0,908 desde 2026.",
         "BEN 2025 (MAE), tabla 2.20: caída de 9,2 % de la hidro en 2024", "verificar.py", "Variante moderada."],
        ["FOI", "Sin pisos térmicos", 2026, "Piso de PPFOI, PPDSL y PPNGS en cero desde 2026.",
         "Decisiones del 12 y 13-10-2026: en el caso central el piso de PPFOI es K16 (mínimo observado 2015-2025, "
         "5,0567 PJ) y los de PPDSL y PPNGS los del Modelo CLG (valor de 2022 congelado, sin fuente)",
         "verificar.py", "Sensibilidad; las capas se miden contra CONA6P."],
        ["IMP0", "Sin importación desde Colombia", 2026, "Piso y techo de IMPELE en cero desde 2026.",
         "Resolución MME 40064, Circular CREG 343 y CENACE 2026 (S26-077)", "verificar.py",
         "Se trata como riesgo para la sobreventa; las capas se miden contra CONA6I."],
        ["ACT", "Situación actual 2026 con choque de demanda transitorio", 2026,
         "2026 observado: factor de planta, piso y techo de cada clase hidroeléctrica por su factor de 2026 (Art6_ACT2026; "
         "Coca Codo con factor de planta %s), demanda eléctrica final por %s e IMPELE en %s PJ (importado antes del "
         "07-09-2026). Desde 2027: razones de ANOSEC3 en 2029, 2032 y 2035, IMPELE en cero y la demanda del Modelo CLG."
         % (("%.4f" % info["act_2026"]["fp_coca_codo_2026"]).replace(".", ","),
            ("%.4f" % info["act_2026"]["multiplicador_demanda"]).replace(".", ","),
            ("%.4f" % info["act_2026"]["importacion_2026_PJ"]).replace(".", ",")),
         FUENTE_ACT + "; " + FUENTE_SECO, "verificar.py: cociente frente a CONA6 igual al factor de Art6_ACT2026 en 2026 y "
         "a la razón de Art6_AnoSeco en 2029, 2032 y 2035", "Decisión del 13-10-2026; supuesto declarado: "
         "choque de demanda transitorio; importación cero desde 2027 como estrés; con respaldo de emergencia."],
        ["ACTDEM", "Situación actual 2026 con demanda persistente (ACT-DEM)", 2026,
         "Igual que ACT, con la demanda eléctrica final por el mismo multiplicador en todos los años desde 2026: el nivel "
         "de 2026 persiste y crece con las tasas del Modelo CLG.", FUENTE_ACT + "; " + FUENTE_SECO,
         "verificar.py", "Decisión del 13-10-2026; supuesto declarado: demanda persistente; con respaldo de "
         "emergencia."],
        ["RESP", "Respaldo de emergencia", 2026, "Tope de PPFOI liberado desde 2026, CapitalCost y FixedCost de PPFOI por "
         "1.000 y STOELE cerrado en 2026-2035, en todos los estados de riesgo (incluidos ACT y ACT-DEM).",
         FUENTE_RESP, "B2_Output_Art6/tablas/respaldo_por_escenario.csv", "Supuesto de cierre, no es un riesgo: solo se "
         "construye la potencia que falta; la energía de respaldo se reporta con sus emisiones, la energía no suministrada "
         "equivalente y sus horas de corte equivalentes. " + NOTA_RE01],
        ["GAS", "Sin la expansión a gas declarada", 2026, "Palancas de A, B y C sobre CONA6G (PPNGS en su piso CLG "
         "2026-2035, G01).", "Decisión del 13-10-2026: el gas como rango", "verificar.py",
         "Sensibilidad; las capas se miden contra CONA6G."],
    ]


# textos de la hoja del sandbox que nombran «oficial» al Modelo Climate Lead Group o describen la proyeccion corregida
# como si fuera la unica (12-10-2026): el analisis usa por omision la proyeccion del Modelo CLG, frente a CONA6 e INCA6
TEXTOS_HEADROOM = [
    ("Agregación oficial de los cinco sectores", "Agregación nacional del Modelo CLG de los cinco sectores"),
    ("la serie modelada de la agregación oficial", "la serie modelada de la agregación nacional del Modelo CLG"),
    ("más la diferencia del sector energía frente a NDCCON.",
     "más la diferencia del sector energía frente a la referencia condicional: CONA6 en la proyección del Modelo CLG, "
     "que usan las gráficas, y NDCCON en la proyección corregida."),
    ("más la diferencia del sector energía frente a NDCINC.",
     "más la diferencia del sector energía frente a la referencia incondicional: INCA6 en la proyección del Modelo CLG "
     "y NDCINC en la proyección corregida."),
]


def _ajuste_headroom(filas):
    out = []
    for f in filas:
        f = dict(f)
        if f.get("parametro") == "archivo_nacional":
            f["valor"] = "Emisiones Totales.xlsx"
            f["fuente"] = "ECU_NDC/PowerBI/Nacionales/Emisiones Totales.xlsx (md5 5954c21d0c5a6c86a5aa23ce8fef10e0, igual " \
                          "a la copia del sandbox)"
        if isinstance(f.get("descripcion"), str):
            for viejo, nuevo in TEXTOS_HEADROOM:
                f["descripcion"] = f["descripcion"].replace(viejo, nuevo)
        out.append(f)
    out.append({"parametro": "proyecciones", "valor": "CLG;corregida", "unidad": None,
                "descripcion": "Proyección del Modelo CLG: la trayectoria nacional de un escenario es la senda de la meta "
                               "más su diferencia de energía frente a la referencia corregida (CONA6 o INCA6), sin el "
                               "puente de la corrección. Proyección corregida: más su diferencia frente a NDCCON o "
                               "NDCINC, con el puente.",
                "fuente": "Decisión del 11-10-2026"})
    return out


def _ajuste_adicionalidad(filas):
    out = []
    for f in filas:
        f = dict(f)
        if f.get("codigo") == "B-BIO":
            f["linea"] = "Biomasa del portafolio ESIN-CELEC EP (39,3 MW, 192,8 GWh/año)"
            f["argumento"] = "Despachable; cuatro plantas de palma y cascarilla identificadas por ESIN-CELEC EP (2020). " \
                             "Debe demostrar que el residuo no tiene uso competitivo y que hay barrera de inversión."
            f["fuente"] = "ESIN-CELEC EP (2020), lámina 31"
        if f.get("codigo") == "FOI":
            f["linea"] = "Sin pisos térmicos del Modelo CLG"
            f["periodo"] = "2026-2035"
        out.append(f)
    return out


def _ajuste_fuentes(filas):
    out = [dict(f) for f in filas]
    for f in out:
        if str(f.get("ubicacion") or "").endswith("Emisiones_Totales_nacionales.xlsx"):
            f["ubicacion"] = "ECU_NDC/PowerBI/Nacionales/Emisiones Totales.xlsx"
        if f.get("ubicacion") == "Energy/B1_Scenario_Config.xlsx":
            f["uso"] = "todos los escenarios (Modelo CLG, referencias corregidas, A, B, C, estados y puentes)"
    out += [{"fuente": "CENACE (2026), Informe Anual 2025, Parte 1, figuras 1.4 y 1.5, p. 31-32",
             "uso": "participación de Coca Codo en la energía de PPHRORAB (CCSOLO y CCDER)",
             "ubicacion": "Energy/Art6_calibracion/parametros_hidro_por_clase.csv"},
            {"fuente": "ESIN-CELEC EP (2020), portafolio de biomasa (39,3 MW, factor de planta 0,56)",
             "uso": "palancas B6, B9 y B10", "ubicacion": "nota técnica de biomasa del equipo consultor"},
            {"fuente": "Base corregida congelada CONA6, INCA6, CONA6G e INCA6G con K16 y RE01 (13-10-2026)",
             "uso": "base de todas las palancas", "ubicacion": "Energy/Art6_calibracion; tests/base_congelada.json"},
            {"fuente": FUENTE_ACT, "uso": "estado ACT y ACT-DEM anclado a 2026: demanda, hidrología por clase, Coca Codo "
                                          "e importación",
             "ubicacion": "Energy/Art6_calibracion/act_2026_observado.csv (situacion_act_2026.py) y "
                          "parametros_hidro_por_clase.csv (agregar_por_clase.py)"},
            {"fuente": "Resultados del 12-10-2026 con los pisos del Modelo CLG (sensibilidad)",
             "uso": "nivel robusto y potencial de la oferta 2031-2035",
             "ubicacion": "Energy/Art6_calibracion/sensibilidad_pisos_CLG_2026-10-12 (MANIFIESTO_md5.csv)"}]
    return out


def _ajuste_changelog(filas):
    return [dict(f) for f in filas] + [{"fecha": "2026-10-12", "autor": "generar_escenarios_art6.py",
        "cambio": "Traslado a ECU_NDC (rama art6-escenarios): A, B y C basados en CONA6 con estados CCOFF, CCSOLO, CCDER, "
                  "ANOSECO, ANOSEC3, ANOSECM, CCOFFANOSECO, FOI, IMP0 y SIN; referencias CONA6P y CONA6I; puentes V7. "
                  "Techos A1, A3, B2 y B4 con la capacidad existente y el factor de planta de CONA6; B9 a 39,3 MW con "
                  "B10 (CF 0,5895) y B6 0,6941 PJ; razones de año seco recompuestas con centrales_clases.csv. Se retira "
                  "A6 (CONA6 ya trae el factor de planta eólico observado, K11b).",
        "motivo": "Decisiones del 11 y 12-10-2026; base congelada del 12-10-2026."},
        {"fecha": "2026-10-13", "autor": "generar_escenarios_art6.py",
         "cambio": "Base con K16 (piso de PPFOI 2026-2035 = 5,0567 PJ, mínimo observado 2015-2025) y RE01 en INCA6G. "
                   "Respaldo de emergencia (grupo RESP: tope de PPFOI liberado, costos de PPFOI por 1.000 y STOELE "
                   "cerrado en 2026-2035) en CCOFF, CCOFFANOSECO, CCSOLO, CCDER, ANOSECO, ANOSEC3, ANOSECM e IMP0. Estado "
                   "nuevo ACT (un año seco cada tres, importación cero y Coca Codo observado) y familia GAS (palancas de "
                   "A, B y C sobre CONA6G). Los resultados con los pisos del Modelo CLG quedan como sensibilidad en "
                   "_stale_ECU_NDC_2026-10-11/resultados_con_pisos_CLG_2026-10-12.",
         "motivo": "Decisiones del 13-10-2026."},
        {"fecha": "2026-10-13", "autor": "generar_escenarios_art6.py",
         "cambio": "ACT rehecho y anclado a 2026: en 2026, demanda eléctrica observada, hidrología observada por clase "
                   "(grupo ACTH), Coca Codo con su factor de planta de 2026 e importación hasta el 06-09-2026 (grupo "
                   "IMPACT); desde 2027, razones de ANOSEC3 en 2029, 2032 y 2035, importación cero y respaldo. Variante "
                   "nueva ACTDEM (ACT-DEM) con la demanda de 2026 persistente. Se retira el grupo CCACT. La atribución de "
                   "RE01 en INCA6G pasa a supuesto de cierre aprobado después de la verificación.",
         "motivo": "Decisión del 13-10-2026 tras la verificación independiente de K16."},
        {"fecha": "2026-10-13", "autor": "generar_escenarios_art6.py",
         "cambio": "Reservas de la verificación del paso 15, sin cambio en las filas de B1: descripciones de Art6_Roles "
                   "corregidas (la proyección del Modelo CLG se mide frente a CONA6 e INCA6 y la corregida frente a "
                   "NDCCON y NDCINC); Art6_ACT2026_Datos con la participación térmica observada con la importación en el "
                   "denominador, como el modelo (enero a agosto de 2026 y 2025), y con la potencia efectiva y los "
                   "factores de planta de Coca Codo leídos del Informe Anual 2025 de CENACE.",
         "motivo": "Verificación independiente del paso 15 (13-10-2026)."}]


AJUSTES = {"Art6_Headroom": _ajuste_headroom, "Art6_Adicionalidad": _ajuste_adicionalidad,
           "Art6_Fuentes": _ajuste_fuentes, "Art6_Changelog": _ajuste_changelog}


def contenido_xlsx(b):
    """Contenido de un xlsx sin metadatos de fecha (para comparar dos versiones)."""
    hojas = pd.read_excel(io.BytesIO(b), sheet_name=None)
    return md5("".join(n + h.to_csv(index=False) for n, h in hojas.items()).encode("utf-8"))


def contenido_hojas(b, art6):
    """md5 del contenido (valores) de las hojas Art6_* (art6=True) o de las demás hojas del libro (art6=False)."""
    hojas = pd.read_excel(io.BytesIO(b), sheet_name=None)
    return md5("".join(n + h.to_csv(index=False) for n, h in hojas.items()
                       if n.startswith("Art6_") == art6).encode("utf-8"))


def _col(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


_CONTROL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _hoja_xml(filas):
    """XML de una hoja con cadenas en línea y números con repr (sin estilos)."""
    partes = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>']
    for r, fila in enumerate(filas, 1):
        celdas = []
        for c, v in enumerate(fila):
            if v is None:
                continue
            ref = _col(c) + str(r)
            if isinstance(v, bool):
                celdas.append('<c r="%s" t="b"><v>%d</v></c>' % (ref, int(v)))
            elif isinstance(v, (int, float)):
                celdas.append('<c r="%s"><v>%s</v></c>' % (ref, repr(v)))
            else:
                celdas.append('<c r="%s" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>'
                              % (ref, xml_escape(_CONTROL.sub("", str(v)))))
        partes.append('<row r="%d">%s</row>' % (r, "".join(celdas)))
    partes.append("</sheetData></worksheet>")
    return "".join(partes).encode("utf-8")


def unir_hojas_art6(libro, doc):
    """Reemplaza las hojas Art6_* del libro de B1 por las del libro documental, al final, editando el XML del libro: las
    hojas que lee B1 quedan byte a byte como estaban (openpyxl reescribiría todo el libro y redondearía los números a 16
    cifras). Falla si el libro tiene fórmulas (pandas, y por tanto B1, solo leen su valor en caché)."""
    d = openpyxl.load_workbook(io.BytesIO(doc), read_only=True)
    hojas = [(ws.title, [list(f) for f in ws.iter_rows(values_only=True)]) for ws in d.worksheets]
    d.close()
    zin = zipfile.ZipFile(io.BytesIO(libro))
    for n in zin.namelist():
        if n.startswith("xl/worksheets/") and n.endswith(".xml") and re.search(rb"<f[ >/]", zin.read(n)):
            raise ErrorGenerador(f"el libro tiene fórmulas en {n}: fíjelas a valores antes de escribir las hojas Art6_*")
    wb_xml = zin.read("xl/workbook.xml").decode("utf-8")
    rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    ct = zin.read("[Content_Types].xml").decode("utf-8")
    destino = {}
    for m in re.finditer(r"<Relationship\b[^>]*/>", rels):
        el = m.group(0)
        destino[re.search(r'\bId="([^"]+)"', el).group(1)] = re.search(r'\bTarget="([^"]+)"', el).group(1)
    quitar = set()
    for m in list(re.finditer(r"<sheet\b[^>]*/>", wb_xml)):
        el = m.group(0)
        if not re.search(r'\bname="([^"]+)"', el).group(1).startswith("Art6_"):
            continue
        rid = re.search(r'\br:id="([^"]+)"', el).group(1)
        t = destino[rid]
        parte = t.lstrip("/") if t.startswith("/") else "xl/" + t
        quitar.add(parte)
        wb_xml = wb_xml.replace(el, "")
        rels = re.sub(r'<Relationship\b[^>]*\bId="%s"[^>]*/>' % re.escape(rid), "", rels)
        ct = re.sub(r'<Override PartName="/%s"[^>]*/>' % re.escape(parte), "", ct)
    ids = [int(x) for x in re.findall(r'\bsheetId="(\d+)"', wb_xml)]
    rnum = [int(x) for x in re.findall(r'\bId="rId(\d+)"', rels)]
    nuevas, sheets, rel_add, ct_add = {}, [], [], []
    for k, (nombre, filas) in enumerate(hojas, 1):
        parte = "xl/worksheets/art6_hoja%d.xml" % k
        rid = "rId%d" % (max(rnum) + k)
        nuevas[parte] = _hoja_xml(filas)
        sheets.append('<sheet name="%s" sheetId="%d" r:id="%s"/>' % (xml_escape(nombre), max(ids) + k, rid))
        rel_add.append('<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                       'worksheet" Target="worksheets/art6_hoja%d.xml"/>' % (rid, k))
        ct_add.append('<Override PartName="/%s" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.'
                      'worksheet+xml"/>' % parte)
    wb_xml = wb_xml.replace("</sheets>", "".join(sheets) + "</sheets>")
    rels = rels.replace("</Relationships>", "".join(rel_add) + "</Relationships>")
    ct = ct.replace("</Types>", "".join(ct_add) + "</Types>")
    cambia = {"xl/workbook.xml": wb_xml.encode("utf-8"), "xl/_rels/workbook.xml.rels": rels.encode("utf-8"),
              "[Content_Types].xml": ct.encode("utf-8")}
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w") as zout:
        for info in zin.infolist():
            if info.filename in quitar:
                continue
            datos = cambia.get(info.filename, zin.read(info.filename))
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = info.compress_type
            zi.external_attr = info.external_attr
            zi.create_system = info.create_system
            zout.writestr(zi, datos)
        for parte, datos in nuevas.items():
            zi = zipfile.ZipInfo(parte, date_time=(2026, 10, 13, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, datos)
    return salida.getvalue()


# ------------------------------------------------------------------------------------------ principal
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--comprobar", action="store_true")
    ap.add_argument("--quitar", action="store_true")
    ap.add_argument("--solo", default=None,
                    help="estados que quedan activos para una corrida parcial de B1, separados por coma (ACT,ACTDEM): "
                         "sus escenarios A, B y C quedan con Activated = YES y los demás del Artículo 6 con NO. Después "
                         "de B1 se vuelve a correr sin --solo para dejar el libro completo")
    args = ap.parse_args()
    rama = GR.rama_actual(ENERGY)
    if GR.RAMA_ESPERADA and rama != GR.RAMA_ESPERADA:
        raise ErrorGenerador(f"rama {rama}; este script solo corre en {GR.RAMA_ESPERADA}")
    huella_base = comprobar_base()
    T, ruta_base = tabla(BASE)
    P, info = palancas(T, None)
    E = escenarios(info)
    info["familia_GAS"] = comprobar_base_gas(T, P, E)
    ids = [p.id for p in P]
    if len(ids) != len(set(ids)):
        raise ErrorGenerador("códigos de palanca repetidos")
    # coherencia: ningún piso sobre su techo en la composición final de cada escenario
    for e in E:
        final = composicion(e, P)
        for (hj, par, cl), p in final.items():
            if par == PISO:
                techo_v = final[(hj, TECHO, cl)].vector if (hj, TECHO, cl) in final else T.get((TECHO, cl), (0, [99999] * 26))[1]
                malos = [a for i, a in enumerate(ANIOS) if p.vector[i] > techo_v[i] + 1e-9]
                if malos:
                    raise ErrorGenerador(f"{e['codigo']}: piso {cl} sobre su techo en {malos[:3]}")
            if par == TECHO:
                piso_v = final[(hj, PISO, cl)].vector if (hj, PISO, cl) in final else T.get((PISO, cl), (0, [0.0] * 26))[1]
                malos = [a for i, a in enumerate(ANIOS) if piso_v[i] > p.vector[i] + 1e-9]
                if malos:
                    raise ErrorGenerador(f"{e['codigo']}: techo {cl} bajo su piso en {malos[:3]}")
    solo = None
    if args.solo:
        sufs = {x.strip() for x in args.solo.split(",") if x.strip()}
        solo = {e["codigo"] for e in E if "-" in e["codigo"] and e["codigo"].split("-", 1)[1] in sufs}
        if not solo:
            raise ErrorGenerador(f"--solo {args.solo}: ningún escenario con esos estados")
    antes = LIBRO.read_bytes()
    nuevo, resumen = construir_libro(antes, E, P, quitar=args.quitar, solo=solo)
    comparar_libros(antes, nuevo, {e["codigo"] for e in E})
    doc_nuevo = libro_documental(E, P, info)
    difiere = []                        # sandbox PNUD: comparación por contenido (valores de las hojas)
    if contenido_hojas(nuevo, art6=False) != contenido_hojas(antes, art6=False):
        difiere.append(LIBRO.name)
    if contenido_hojas(antes, art6=True) != contenido_xlsx(doc_nuevo):
        difiere.append("hojas Art6_*")
    if args.comprobar:
        print("Coincide con lo que generaría el script" if not difiere else "Difiere: " + ", ".join(difiere))
        return 1 if difiere else 0
    if LIBRO.name in difiere or "hojas Art6_*" in difiere:
        respaldo = RESPALDOS
        respaldo.mkdir(parents=True, exist_ok=True)
        copia = respaldo / f"B1_Scenario_Config.xlsx.bak_escenarios_art6_{datetime.now():%H%M%S}"
        shutil.copy2(LIBRO, copia)
        LIBRO.write_bytes(nuevo if LIBRO.name in difiere else antes)
        print(f"Libro escrito ({len(E) if not args.quitar else 0} escenarios del Artículo 6); respaldo en {copia}")
    if args.quitar:
        return 0
    if "hojas Art6_*" in difiere:
        LIBRO.write_bytes(unir_hojas_art6(LIBRO.read_bytes(), doc_nuevo))
        print(f"Hojas documentales escritas al final de {LIBRO.relative_to(ENERGY)}")
    for e in E:
        (ENERGY / "B1_Output_Params" / e["codigo"]).mkdir(parents=True, exist_ok=True)
    sb_md5 = {f"{c}_0.txt": md5((SANDBOX / "Executables" / f"{c}_0" / f"{c}_0.txt").read_bytes())
              for c in ("NDCCON", "A", "B", "C")}
    registro = {"fecha": datetime.now().isoformat(timespec="seconds"), "rama": rama, "commit": GR.commit_actual(ENERGY),
                "base": BASE, "base_congelada": huella_base, "libro_md5": md5(LIBRO.read_bytes()),
                "art6_configuracion_contenido_md5": contenido_hojas(LIBRO.read_bytes(), art6=True),
                "sandbox_datafiles_md5": sb_md5, "supuestos": info, "escenarios": [e["codigo"] for e in E],
                "activos_solo": sorted(solo) if solo else "todos",
                "filas_por_escenario": resumen,
                "palancas": {p.id: {"grupo": p.grupo, "parametro": p.parametro, "clave": p.clave,
                                    "2026": p.vector[IX[2026]], "2030": p.vector[IX[2030]],
                                    "2035": p.vector[IX[2035]]} for p in P},
                "insumos_md5": {n: md5((AQUI / n).read_bytes()) for n in
                                ("parametros_hidro_por_clase.csv", "act_2026_observado.csv")}}
    REGISTRO.write_text(json.dumps(registro, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{len(E)} escenarios y {len(P)} palancas; registro en {REGISTRO.name}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except ErrorGenerador as ex:
        print("ERROR", ex)
        sys.exit(2)
