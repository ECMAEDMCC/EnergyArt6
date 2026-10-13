# -*- coding: utf-8 -*-
"""Construye Art6_calibracion/correcciones_K.csv (K01-K16, G01 y RE01) desde sus fuentes; K01-K16 se aplican iguales a
INCA6, CONA6, INCA6G y CONA6G (salvo las filas marcadas), G01 solo a INCA6G y CONA6G (sensibilidad de gas) y RE01 solo a
INCA6G (respaldo de emergencia).

Decisiones del 13-10-2026:
K16.  El piso de despacho de PPFOI 2026-2035 pasa del valor CLG (2022 congelado, 9,3785 PJ hasta 2030 y 9,0 PJ en 2035,
      sin fuente) al mínimo de la historia anclada 2015-2025: el piso K13b del año de menor generación observada (2021,
      BEN 2025, tablas 2.20, 2.28 y 3.12-3.15), igual a la actividad de PPFOI del clon en ese año. Los pisos CLG de PPDSL
      (0,4201 PJ) y PPNGS (1,9665 PJ) no cambian.
RE01. Respaldo de emergencia en INCA6G, que sin él no tiene solución: tope de PPFOI liberado desde 2026, CapitalCost y
      FixedCost de PPFOI por 1.000 y STOELE cerrado en 2026-2035. Supuesto de cierre, como la térmica arrendada que el
      país contrata hoy; las mismas filas van en los estados de riesgo de generar_escenarios_art6.py (grupo RESP).
      RE01 no fue una decisión: se introdujo en la corrida K16, fuera del acuerdo original (respaldo de
      fuel oil solo en los estados de riesgo). Se aprobó el 13-10-2026, después de la verificación
      independiente, conservarlo rotulado «con cierre de fuel oil de emergencia», con el peso del respaldo al lado.

Rótulos: «CLG» es el Modelo Climate Lead Group, el modelo publicado en ECU_NDC (clg-admin/ECU_NDC, main f19d6b1), cuyos
escenarios BAU, BAU1NDC, NDCINC y NDCCON no se tocan. En las fuentes y notas del CSV «CLG» reemplaza al antiguo «oficial».

Decisiones del 12-10-2026 que entran en esta versión:
(a) K14: la actividad 2026-2035 de PPCRU, PPRNGS, PPRDSL (térmica de autoproductores) y PPILPG queda anclada a la media
    observada 2023-2025 (K13: BEN 2025, tablas 3.14-3.15 y 2.20; para PPILPG, GLP de la tabla 3.14).
(b) G01: INCA6G y CONA6G son INCA6 y CONA6 con PPNGS en su piso CLG 2026-2035 (techo = piso).
(c) La inversión mínima solar del condicional de 2025 (67 MW) se conserva (sin fila).
(d) K13g: PPDDSL 2011-2014 toma el InputActivityRatio diésel del propio modelo (el de PPDSL en A2 NDCCON); K08 lo usa.
(e) La capa CENTRALES de ARCONEL (uso autorizado; referencia del oficio [verificar]) entra solo en K11 (fecha de entrada
    de Huascachaca y de la eólica distribuida) y K12 (fecha de entrada de la solar del SNI en 2013-2014); en K04, K06 y
    K03 sirve de contraste. Se lee del extracto Art6_calibracion/centrales_series.csv (extraer_centrales.py).

Decisiones del 12-10-2026 (tarde) que entran en la versión K15:
K15a. CF 2026-2035 de las térmicas de la red continuo con lo observado: medio 2023-2025 del CF del datafile del clon
      (K13d o CLG) cuando el CF CLG de 2026 cae más de 5 % frente a ese medio (PPFOI y PPDSL; PPNGS, PPDFOI y PPDDSL se
      revisan y no tienen esa discontinuidad).
K15b. Inversión térmica nueva limitada a la expansión declarada (plan ajustado, Anexo A del Acuerdo MAE-MAE-2026-0066-AM,
      y plan de 920 MW de CELEC), conciliada central por central en expansion_termica_declarada.csv. B1 escribe
      TotalAnnualMaxCapacityInvestment como bloque vacío fijo (B1_Base_Scenarios_Adj_Parallel.py, l. 1311) y glpsol
      rechaza un segundo bloque del mismo parámetro, así que el tope entra, sin tocar B1, como TotalAnnualMaxCapacity
      2026-2035 = residual K15c de 2025 + expansión declarada acumulada (99999, el valor por omisión, en 2010-2025).
      Es el tope acumulado equivalente; calibracion_propia.py comprueba después si la inversión de cada año cabe en lo
      declarado para ese año.
K15c. Potencia residual 2023-2035 de las térmicas contrastada con Energy.xlsx (Capacidad_GW 2023 por sistema y
      combustible), con las altas de 2024-2025 de la capa CENTRALES y de CENACE (termica_altas_2024_2025.csv) y con los
      totales de la Estadística Anual y Multianual 2025 de ARCONEL; se corrige donde difiere más que el redondeo.
K14 revisado. PPCRU ancla los autoproductores (BEN 3.14-3.15) y el crudo de centrales eléctricas (BEN 3.12-3.13) porque
      las centrales de crudo que la capa identifica (CPT Crudo Diesel CELEC y Cuyabeno Crudo Diesel CELEC) son de
      sistemas no incorporados.
Biomasa y biogás: potencia efectiva de la Estadística Anual y Multianual 2025 de ARCONEL, tablas 18-19
(capacidad_bio_fuentes.csv).

Arreglo de la verificación (12-10-2026, noche), versión K15 con arriendos:
- La potencia arrendada es temporal. Los arriendos de 2026 del plan de 920 MW (expansion_termica_declarada.csv, filas
  «arriendo») y las centrales arrendadas de 2024-2025 (termica_altas_2024_2025.csv: Karpowership, Erin Sultan, Murat Bey
  III, GASVESUBIO y El Descanso emergente) entran como potencia residual con vigencia (K15c), no como inversión: la
  potencia de un año es MW x fracción del año en que el contrato está vigente. Santa Elena va del 12-11-2026 al
  10-11-2028 (S26-081). Para los demás no hay fecha de fin publicada (S26-116): caso central, fin el 31-12-2028
  (único contrato con fecha, fin del bloque 2 de generación firme del PME en marzo de 2028); variante, fin el
  31-03-2032 (fin del bloque 1, PME 2023-2032, cap. 4, tabla 4-9, p. 134), que variante_suficiencia.py resuelve aparte.
  El tope K15b es la potencia del año (residual K15c + construcción del modelo en 2010-2022) más la expansión propia
  declarada acumulada.
- K14 revisado (crudo): PPCRU ancla la media 2023-2025 del crudo de autoproductores (BEN 3.14-3.15) y, del crudo de
  centrales eléctricas (BEN 3.12-3.13), solo el nivel de 2025 (0,6035 PJ), el único año en que CENACE confirma que el
  S.N.I. no quemó crudo; el exceso de 2023-2024 no tiene central atribuible y queda al despacho. CPT Crudo Diesel CELEC y
  Cuyabeno Crudo Diesel CELEC (47,16 MW, capa CENTRALES) se suman a la potencia de PPCRU desde 2023 porque Energy.xlsx no
  las incluye [inferencia]. El IAR de PPCRU desde 2026 es el de la generación que K14 ancla.
- K15d y K15e (gas): el techo CLG de PPNGS (2,9497 PJ, la generación del S.N.I. a gas de 2022 congelada, sin fuente de
  suministro) se reemplaza por 1,02 x la media observada 2023-2025 más la energía media declarada de la expansión a gas
  desde su año de entrada (Anexo A: Machala tercera unidad 510 GWh, 2026; Machala ciclo combinado 690 GWh, 2027;
  Bloque de Ciclo Combinado I 3.000 GWh, 2028, con gas propio según el PME, cap. 4, p. 129), y el CF de PPNGS es el que
  deja a la potencia declarada producir ese techo (regla del modelo). Solo INCA6 y CONA6: G01 sigue fijando el techo de
  INCA6G y CONA6G en el piso CLG.

Uso, desde ECU_NDC/Energy:
    python Art6_calibracion/construir_correcciones.py              escribe correcciones_K.csv y correcciones_K_insumos.json
    python Art6_calibracion/construir_correcciones.py --comprobar  no escribe; sale con 1 si el CSV en disco no coincide

Después: python Art6_calibracion/generar_referencias.py (aplica el CSV al libro y a las carpetas A2 de INCA6 y CONA6).

Criterio (12-10-2026): la historia 2010-2025 se ancla a lo observado; ningún valor se ajusta para que una métrica cierre,
cada uno sale de una fuente o de una regla del propio modelo; nunca CF = 1; 2010 intacto en actividad y emisiones;
2026-2035 conserva la proyección oficial salvo donde una corrección histórica tiene continuidad física (capacidad
existente que sigue operando).

Fuentes (leídas por el script; las publicaciones están transcritas en CSV con su página):
- A2_Output_Params/NDCCON y NDCINC: vectores oficiales.
- Executables/NDCCON_0 y NDCINC_0: datafile y _Output.csv oficiales (no regresión IGUAL), para los vectores que B1
  reescribe desde las hojas (eólica y solar del condicional), los valores de 2010 que se conservan y la capacidad que el
  modelo oficial construye (hidro por clase, eólica 2010, solar 2012).
- ECU_NDC/Data/Energy.xlsx: hoja Proyectos_Plantas (Mazar, Coca Codo, Sopladora, Villonaco), hoja Capacidad_GW (eólica y
  solar distribuidas) y hoja FP (factor de planta eólico 2013-2014, contraste).
- BEN 2025 (MAE): tabla 2.20, p. 66 (ben2025_tabla_2_20.csv); tabla 2.23, p. 67 (ben2025_tabla_2_23.csv); tabla 2.25,
  p. 69 (ben2025_tabla_2_25_sgda.csv).
- CENACE, Informes Anuales 2016 y 2017 (Parte 1): cenace_produccion_neta_2016_2017.csv; fechas de entrada en el CSV
  y en capacidad_bio_fuentes.csv.
- capacidad_bio_fuentes.csv: potencia de biomasa (PPIBIM) y biogás (PPBGS).
- IPCC 2006, vol. 2, cap. 2, cuadro 2.2 (industrias de la energía): factores de CO2, CH4 y N2O; PCG AR5 28 y 265.
- BEN 2025 (MAE): tabla 2.28, p. 71, combustibles para generación (ben2025_tabla_2_28.csv); tablas 3.12-3.13, p. 86
  (centrales eléctricas) y 3.14-3.15, p. 87 (autoproductores), carga por combustible y electricidad producida
  (ben2025_tablas_3_12_a_3_15.csv); factores de conversión del anexo, p. 178 (10^3 kWh = 0,6196 BEP).
- ECU_NDC/Data/Energy.xlsx, hoja Generación_PJ: generación distribuida con diésel y fuel oil (K13).
- ARCONEL, capa CENTRALES (centrales_series.csv y centrales_extracto.csv, con el md5 del .shp y del .dbf): K11 y K12.
Ninguna corrección usa Capacidad_instalada_por_fuente_ARCONEL.xlsx (extracto sin año).

K13 (12-10-2026, tarde) reemplaza a K09a y K09b: el rendimiento (InputActivityRatio) de las térmicas fósiles 2015-2025
sale del BEN (combustible fósil / generación fósil de cada grupo de centrales) y la generación de cada térmica sale del
combustible del BEN dividido por ese rendimiento, así el combustible por tipo del modelo es el del BEN; desde 2026 el
rendimiento es el observado de 2023-2025 (continuidad del parque existente) y la actividad conserva la proyección oficial.
"""
import argparse
import os
import csv
import hashlib
import io
import json
import math
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from verificar_correcciones import leer_datafile  # noqa: E402  (parser del datafile de B1)

ENERGIA = Path(__file__).resolve().parents[1]
AQUI = Path(__file__).resolve().parent
A2 = ENERGIA / "A2_Output_Params" / "NDCCON"
EX = ENERGIA / "Executables"
ENERGY_XLSX = ENERGIA / "Insumos_Art6" / "Energy.xlsx"   # copia sin cambios de Data/Energy.xlsx del Modelo CLG
BEN_CSV = AQUI / "ben2025_tabla_2_20.csv"
BEN_POT = AQUI / "ben2025_tabla_2_23.csv"
BEN_SGDA = AQUI / "ben2025_tabla_2_25_sgda.csv"
CENACE_CSV = AQUI / "cenace_produccion_neta_2016_2017.csv"
CAP_CSV = AQUI / "capacidad_bio_fuentes.csv"
BEN_228 = AQUI / "ben2025_tabla_2_28.csv"
BEN_3 = AQUI / "ben2025_tablas_3_12_a_3_15.csv"
TJ_KBEP = 3.6 / 0.6196  # TJ por kBEP: BEN 2025, anexo de factores de conversión, p. 178 (10^3 kWh = 0,6196 BEP)
SALIDA = AQUI / "correcciones_K.csv"
INSUMOS = AQUI / "correcciones_K_insumos.json"
CENTRALES_SERIES = AQUI / "centrales_series.csv"
CENTRALES_EXTRACTO = AQUI / "centrales_extracto.csv"
ALTAS = AQUI / "termica_altas_2024_2025.csv"
EXPANSION = AQUI / "expansion_termica_declarada.csv"
CONCILIACION = AQUI / "conciliacion_termica_2025.csv"
# fuentes de K15 guardadas localmente (documentos públicos; no se distribuyen): carpeta en la variable de
# entorno ART6_FUENTES (por omisión, Fuentes en la carpeta de trabajo). Solo se usa su md5 en los insumos.
FUENTES_LOCALES = Path(os.environ.get("ART6_FUENTES", "Fuentes"))
FUENTES_K15 = {
    "Anexo A del ajuste del PME": FUENTES_LOCALES / "Situacion_2026" / "MAE_Anexos_Ajuste_PME_2023-2032_2026-06.pdf",
    "PME 2023-2032, cap. 4": FUENTES_LOCALES / "CAP 4.pdf",
    "CELEC, Boletín 042": FUENTES_LOCALES / "Situacion_2026" / "CELEC_boletines_2026"
                          / "CELEC_Boletin042_plan_920MW_termicos_2026-04-20.html",
    "CELEC, Boletín 122": FUENTES_LOCALES / "Situacion_2026" / "CELEC_boletines_2026"
                          / "CELEC_Boletin122_Santa_Elena_45MW_rentados_2026-10-03.html",
    "CENACE, Informe Anual 2025, Parte 2": FUENTES_LOCALES / "Situacion_2026"
                                           / "CENACE_Informe_Anual_2025_Parte2_2026-04.pdf",
    "ARCONEL, Estadística Anual y Multianual 2025": FUENTES_LOCALES / "Situacion_2026"
                                                    / "ARCONEL_EstadisticaAnualMultianualSectorElectrico2025_2026-04.pdf"}
F_ARC25 = ("ARCONEL, Estadística Anual y Multianual del Sector Eléctrico Ecuatoriano 2025 (abril de 2026), tablas 18-23, "
           "p. 33-34")
F_ENERGY_CAP = "ECU_NDC/Data/Energy.xlsx, hoja Capacidad_GW (2023, por sistema y combustible)"
# K15c: fila de Energy.xlsx (Capacidad_GW) que corresponde a cada térmica del modelo
K15C_MAPA = {"PPFOI": ("SNI", "Fuel Oil"), "PPDSL": ("SNI", "Diesel"), "PPNGS": ("SNI", "Gas Natural"),
             "PPCRU": ("Autoproductores", "Crudo"), "PPRDSL": ("Autoproductores", "Diesel"),
             "PPRNGS": ("Autoproductores", "Gas Natural"), "PPDDSL": ("Distribuida", "Diesel"),
             "PPDFOI": ("Distribuida", "Fuel Oil")}
RED_K15 = ("PPFOI", "PPDSL", "PPNGS")  # térmicas de la red (S.N.I.) que desplazan A, B y C
K15C_ANIOS = list(range(2023, 2036))
ESC = "INCA6;CONA6;INCA6G;CONA6G"   # K01-K14: las cuatro referencias corregidas
ESC_INC, ESC_CON = "INCA6;INCA6G", "CONA6;CONA6G"  # filas propias de la familia incondicional o condicional
ESC_G = "INCA6G;CONA6G"             # G01: sensibilidad de gas
ESC_CEN = "INCA6;CONA6"             # K15d-e: caso central (en INCA6G y CONA6G manda G01)
# Vigencia de los arriendos sin fecha de fin publicada (S26-116): caso central y variante
FIN_ARRIENDO = {"2028": date(2028, 12, 31), "2032": date(2032, 3, 31)}
VARIANTE_CENTRAL = "2028"
F_FIN_ARRIENDO = {
    "2028": "fin el 31-12-2028: único contrato con fecha publicada (Santa Elena, 10-11-2028, S26-081) y fin del bloque 2 de "
            "generación firme en marzo de 2028 (PME 2023-2032, cap. 4, tabla 4-9, p. 134)",
    "2032": "fin el 31-03-2032: fin del bloque 1 de generación firme (PME 2023-2032, cap. 4, tabla 4-9, p. 134)"}
F_PME49 = "MEM, PME 2023-2032, cap. 4, p. 134, tabla 4-9 (Documents/Articulo 6/CAP 4.pdf)"
F_ANEXO_A = ("MAE, Anexos del ajuste PME 2023-2032, Anexo A, p. 1-2 (Acuerdo MAE-MAE-2026-0066-AM, RO 319; S26-109): "
             "energía media de Machala Gas Tercera Unidad (77 MW, 510 GWh/año, 2026), Machala Gas Ciclo Combinado (110 MW, "
             "690 GWh/año, 2027) y Bloque de Ciclo Combinado I (400 MW, 3.000 GWh/año, 2028)")
REF ={"INCA6": "NDCINC", "CONA6": "NDCCON", "INCA6G": "NDCINC", "CONA6G": "NDCCON"}
AUTOPRODUCTORES_K14 = ("PPCRU", "PPRNGS", "PPRDSL")  # decisión (a); PPILPG va aparte (su actividad es combustible)
F_CAPA = ("ARCONEL, capa CENTRALES; uso autorizado; referencia del oficio [verificar] (Art6_calibracion/"
          "centrales_series.csv y centrales_extracto.csv, generados por extraer_centrales.py)")
PJ = 277.78  # GWh por PJ (regla del proyecto: 1 PJ = 277,78 GWh)
ANIOS = list(range(2010, 2036))
HIST = list(range(2015, 2026))  # años con BEN 2025
PROY = list(range(2026, 2036))
COLUMNAS = ["codigo", "escenarios", "hoja_o_carpeta", "parametro", "tecnologia", "anios", "valores", "unidad",
            "fuente", "nota"]
UL, LL = "TotalTechnologyAnnualActivityUpperLimit", "TotalTechnologyAnnualActivityLowerLimit"
EAR, MININV = "EmissionActivityRatio", "TotalAnnualMinCapacityInvestment"
# IPCC 2006, vol. 2, cap. 2, cuadro 2.2 (combustión estacionaria en industrias de la energía): CO2 en t/TJ, CH4 y N2O
# en kg/TJ. PCG AR5: CH4 28 y N2O 265 (los del inventario y de la Segunda NDC).
IPCC = {"E1FO1": ("fuel oil residual", 77.4, 3.0, 0.6), "E1NGS": ("gas natural", 56.1, 1.0, 0.1),
        "E1DSL": ("gas/diésel oil", 74.1, 3.0, 0.6), "E1CRU": ("petróleo crudo", 73.3, 3.0, 0.6)}
PCG = {"CH4": 28, "N2O": 265}
TERMICAS_K08 = ["PPFOI", "PPDFOI", "PPDSL", "PPRDSL", "PPDDSL", "PPNGS", "PPRNGS", "PPCRU"]
F_BEN = "BEN 2025 (MAE), tabla 2.20, p. 66 (Art6_calibracion/ben2025_tabla_2_20.csv)"
F_BEN23 = "BEN 2025 (MAE), tabla 2.23, p. 67, potencia efectiva (Art6_calibracion/ben2025_tabla_2_23.csv)"
F_BEN25 = "BEN 2025 (MAE), tabla 2.25, p. 69, SGDA (Art6_calibracion/ben2025_tabla_2_25_sgda.csv)"
F_PP = "ECU_NDC/Data/Energy.xlsx, hoja Proyectos_Plantas"
F_CAPX = "ECU_NDC/Data/Energy.xlsx, hoja Capacidad_GW"
F_CEN16 = ("CENACE, Informe Anual 2016 (Parte 1), tabla 3.3, p. 61-62 "
           "(Art6_calibracion/cenace_produccion_neta_2016_2017.csv)")
F_CEN16_FECHA = ("CENACE, Informe Anual 2016 (Parte 1): tabla 2.4, p. 42 (Coca Codo Sinclair vende energía desde mayo de "
                 "2016; enero-abril 0,00 GWh), p. 18 (caudales en Coca Codo desde mayo de 2016) y tabla 3.3, p. 62 "
                 "(Sopladora U2 y U3 desde mayo y U1 desde junio de 2016)")
F_CAP = "Art6_calibracion/capacidad_bio_fuentes.csv"
F_A2 = "A2_Output_Params/NDCCON"
F_OF = "datafile y _Output.csv CLG de Executables/{NDCCON,NDCINC}_0 (no regresión IGUAL)"
F_BEN228 = ("BEN 2025 (MAE), tabla 2.28, p. 71, combustibles para generación eléctrica "
            "(Art6_calibracion/ben2025_tabla_2_28.csv)")
F_BEN3 = ("BEN 2025 (MAE), tablas 3.12-3.13, p. 86 (centrales eléctricas) y 3.14-3.15, p. 87 (autoproductores): carga "
          "por combustible y electricidad producida (Art6_calibracion/ben2025_tablas_3_12_a_3_15.csv)")
F_BEP = "BEN 2025, anexo de factores de conversión (OLADE 2017), p. 178: 10^3 kWh = 0,6196 BEP, 1 kBEP = 5,8102 TJ"
F_GENX = "ECU_NDC/Data/Energy.xlsx, hoja Generación_PJ (generación distribuida, diésel y fuel oil)"
# K13: combustible de cada térmica fósil. Grupo A = centrales eléctricas (tablas 3.12-3.13), B = autoproductores
# (3.14-3.15). El modelo no tiene térmica de fuel oil de autoproductores ni de crudo de centrales: el fuel oil y el
# crudo reducido de B van a PPFOI y el crudo de A va a PPCRU, cada uno con el rendimiento de su grupo.
COMB_K13 = {"PPCRU": "E1CRU", "PPNGS": "E1NGS", "PPRNGS": "E1NGS", "PPDSL": "E1DSL", "PPRDSL": "E1DSL",
            "PPDDSL": "E1DSL", "PPFOI": "E1FO1", "PPDFOI": "E1FO1"}
FOSILES_KBEP = ["petroleo_kBEP", "gas_natural_kBEP", "gas_licuado_kBEP", "diesel_oil_kBEP", "fuel_oil_kBEP",
                "crudo_reducido_kBEP"]


def md5(ruta):
    return hashlib.md5(Path(ruta).read_bytes()).hexdigest()


def r4(x):
    return round(float(x) + 0.0, 4)


def techo4(x):
    """Redondeo hacia arriba en el cuarto decimal: el CF así fijado nunca deja la capacidad por debajo del piso."""
    return math.ceil(float(x) * 1e4 - 1e-7) / 1e4


def piso4(x):
    """Redondeo hacia abajo en el cuarto decimal: un techo acotado por la capacidad nunca la excede."""
    return math.floor(float(x) * 1e4 + 1e-7) / 1e4


def fecha_ddmm(texto, defecto=None):
    """Fecha «dd-mm-aaaa» (o la de omisión si el texto está vacío)."""
    texto = str(texto or "").strip()
    if not texto:
        return defecto
    d, m, a = (int(x) for x in texto.split("-"))
    return date(a, m, d)


def fraccion_anio(inicio, fin, anio):
    """Fracción del año en que una central con contrato entre inicio y fin (inclusive) está disponible."""
    a, b = max(inicio, date(anio, 1, 1)), min(fin, date(anio, 12, 31))
    if b < a:
        return 0.0
    return ((b - a).days + 1) / ((date(anio, 12, 31) - date(anio, 1, 1)).days + 1)


def clg(texto):
    """Rótulo del 12-10-2026: el modelo publicado se llama Modelo Climate Lead Group (CLG), no «oficial»."""
    texto = re.sub(r"\b[Mm]odelo oficial\b", "Modelo CLG", texto)
    texto = re.sub(r"\b[Oo]ficiales\b", "CLG", texto)
    return re.sub(r"\b[Oo]ficial\b", "CLG", texto)


def fmt(x, dec=4):
    if dec == "sig":
        v = float(x) + 0.0
        return "0" if v == 0 else f"{v:.10g}"
    v = round(float(x) + 0.0, dec)
    return "0" if v == 0 else f"{v:.{dec}f}".rstrip("0").rstrip(".")


# ------------------------------------------------------------------------------------------------ lectura de fuentes
def leer_a2(parametro, carpeta=A2):
    return pd.read_csv(carpeta / f"{parametro}.csv", dtype={"FUEL": str, "EMISSION": str, "YEAR": str},
                       keep_default_na=False, low_memory=False)


def vector(parametro, tec, fuel=None, emision=None, carpeta=A2):
    d = leer_a2(parametro, carpeta)
    d = d[d.TECHNOLOGY == tec]
    if fuel is not None:
        d = d[d.FUEL == fuel]
    if emision is not None:
        d = d[d.EMISSION == emision]
    if d.empty:
        return {}
    if (d.YEAR == "").all():
        return {None: float(d.Value.iloc[0])}
    s = {int(y): float(v) for y, v in zip(d.YEAR, d.Value)}
    if len(s) != len(d):
        raise SystemExit(f"{parametro} {tec}: años repetidos en A2")
    return s


class Oficial:
    """Datafile y salida de un escenario oficial (lo que B1 escribe y lo que el modelo construye)."""

    def __init__(self, esc):
        self.esc = esc
        self.ruta_df = EX / f"{esc}_0" / f"{esc}_0.txt"
        self.ruta_out = EX / f"{esc}_0" / f"{esc}_0_Output.csv"
        self.df = leer_datafile(self.ruta_df)
        o = pd.read_csv(self.ruta_out, usecols=["Technology", "Year", "NewCapacity", "TotalCapacityAnnual"],
                        low_memory=False)
        self.nueva = o.groupby(["Technology", "Year"]).NewCapacity.max().fillna(0).to_dict()
        self.cap = o.groupby(["Technology", "Year"]).TotalCapacityAnnual.max().fillna(0).to_dict()

    def v(self, par, tec, anio, idx2=None):
        if par == "CapacityFactor":
            k = (par, tec, "All", str(anio))
        elif par in (EAR, "InputActivityRatio", "OutputActivityRatio"):
            k = (par, tec, idx2, "1", str(anio))
        else:
            k = (par, tec, str(anio))
        return self.df[k][0] if k in self.df else None


def fila_excel(ws, sistema, comb):
    """Valores por año de la fila (Sistema, Combustible) de Capacidad_GW y referencia de la fila."""
    cab = None
    for fila in ws.iter_rows():
        vals = [c.value for c in fila]
        if "Sistema" in vals and cab is None:
            cab = {c.column: c.value for c in fila if isinstance(c.value, int)}
            continue
        if cab and len(vals) > 2 and vals[1] == sistema and vals[2] == comb:
            return {cab[c.column]: float(c.value) for c in fila if c.column in cab}, f"fila {fila[0].row}"
    raise SystemExit(f"Capacidad_GW: falta la fila {sistema} / {comb}")


def fila_fp(ws, sistema, tec):
    cab = None
    for fila in ws.iter_rows():
        vals = [c.value for c in fila]
        if cab is None and any(isinstance(v, int) and v == 2010 for v in vals):
            cab = {c.column: c.value for c in fila if isinstance(c.value, int)}
            continue
        if cab and sistema in vals and tec in vals:
            return {cab[c.column]: float(c.value) for c in fila if c.column in cab}, f"fila {fila[0].row}"
    raise SystemExit(f"FP: falta la fila {sistema} / {tec}")


def fila_gen(ws, sistema, tec):
    """Generación (PJ) por año de la primera fila (Sistema, Tecnología) de Generación_PJ (bloque NDC Incondicional +
    Condicional; las filas de generación distribuida son iguales en el bloque Incondicional) y referencia de la fila."""
    cab = None
    for fila in ws.iter_rows():
        vals = [c.value for c in fila]
        if cab is None and "Tecnología" in vals:
            cab = {c.column: c.value for c in fila if isinstance(c.value, int)}
            continue
        if cab and len(vals) > 5 and vals[4] == sistema and vals[5] == tec:
            return {cab[c.column]: float(c.value) for c in fila if c.column in cab}, f"fila {fila[0].row}"
    raise SystemExit(f"Generación_PJ: falta la fila {sistema} / {tec}")


def fuentes():
    wb = openpyxl.load_workbook(ENERGY_XLSX, data_only=True, read_only=False)
    pp = wb["Proyectos_Plantas"]
    proy = {ref: pp[ref].value for ref in ("C4", "F4", "G4", "H4", "C5", "F5", "H5", "C6", "F6", "C8", "F8", "G8", "H8")}
    if proy["C4"] != "Central Hidroeléctrica Mazar" or proy["C8"] != "Proyecto Hidroeléctrico Sopladora" or \
            proy["C6"] != "Central Hidroeléctrica Coca Codo Sinclair" or proy["C5"] != "Proyecto Eólico Villonaco":
        raise SystemExit(f"Proyectos_Plantas cambió de orden: {proy}")
    capx = {k: fila_excel(wb["Capacidad_GW"], *k) for k in
            (("Distribuida", "Eólica"), ("SNI", "Eólica"), ("Distribuida", "Solar"), ("SNI", "Solar"))
            + tuple(K15C_MAPA.values())}
    fp_eol = fila_fp(wb["FP"], "Generación en el SNI", "Eólico")
    gen_dist = {c: fila_gen(wb["Generación_PJ"], "Generación distribuida", c) for c in ("Diesel", "Fuel Oil")}
    wb.close()
    ben = pd.read_csv(BEN_CSV, index_col="anio")
    pot = pd.read_csv(BEN_POT, index_col="anio")
    sgda = pd.read_csv(BEN_SGDA, index_col="anio").sgda_MW
    for t in (ben, pot):
        if list(t.index) != HIST:
            raise SystemExit("las tablas del BEN deben traer 2015-2025")
    cen = pd.read_csv(CENACE_CSV)
    cen = cen.groupby(["anio", "central"]).gwh_neta.sum().round(2).to_dict()
    cap = {"PPIBIM": {}, "PPBGS": {}}
    for _, f in pd.read_csv(CAP_CSV).iterrows():
        for y in range(int(f.anio_desde), int(f.anio_hasta) + 1):
            if y in cap[f.tecnologia]:
                raise SystemExit(f"{CAP_CSV.name}: {f.tecnologia} {y} repetido")
            cap[f.tecnologia][y] = {"MW": float(f.MW), "estado": f.estado, "tipo": f.tipo_potencia}
    if sorted(cap["PPIBIM"]) != list(range(2015, 2036)) or sorted(cap["PPBGS"]) != list(range(2016, 2036)):
        raise SystemExit(f"{CAP_CSV.name}: PPIBIM debe cubrir 2015-2035 y PPBGS 2016-2035")
    of = {e: Oficial(e) for e in ("NDCCON", "NDCINC")}
    capa = pd.read_csv(CENTRALES_SERIES)
    capa = {g: d.set_index("anio") for g, d in capa.groupby("grupo")}
    return proy, capx, fp_eol, cen, cap, ben, pot, sgda, of, gen_dist, capa


# ------------------------------------------------------------------------------------------------ construcción
class Filas:
    def __init__(self):
        self.filas = []

    def add(self, codigo, destino, parametro, tecnologia, anios, valores, unidad, fuente, nota, dec=4, esc=ESC):
        if anios is None:
            anios_txt, valores_txt = "sin_anio", fmt(valores[0], dec)
        else:
            anios = list(anios)
            anios_txt = f"{anios[0]}-{anios[-1]}" if anios == list(range(anios[0], anios[-1] + 1)) and len(anios) > 1 \
                else ";".join(str(a) for a in anios)
            valores_txt = ";".join(fmt(v, dec) for v in valores)
        self.filas.append({"codigo": codigo, "escenarios": esc, "hoja_o_carpeta": destino, "parametro": parametro,
                           "tecnologia": tecnologia, "anios": anios_txt, "valores": valores_txt, "unidad": unidad,
                           "fuente": clg(fuente), "nota": clg(nota)})


def iguales_historia(of, par, tec, anios, idx2=None):
    a = [of["NDCCON"].v(par, tec, y, idx2) for y in anios]
    b = [of["NDCINC"].v(par, tec, y, idx2) for y in anios]
    if a != b:
        raise SystemExit(f"{par} {tec}: NDCCON y NDCINC difieren en {anios}; revise el diseño")
    return a


def rendimientos_observados(ben, gen_dist):
    """K13: rendimiento y generación de las ocho térmicas fósiles (y combustible de PPILPG) 2015-2025 desde el BEN.

    Rendimiento del grupo g (A centrales eléctricas, B autoproductores) en el año y:
        IAR_g = combustible fósil del grupo (tablas 3.12 y 3.14, kBEP x 5,8102 TJ/kBEP)
                / generación fósil del grupo (electricidad de 3.13 y 3.15 menos hidro, menos eólica, solar y biogás en A
                  y menos biomasa en B, todas en GWh de la tabla 2.20)
    La suma de la generación fósil de A y B reproduce la térmica de la tabla 2.20 (control). Supuesto: dentro de un
    grupo y un año todos los combustibles tienen el mismo rendimiento (el BEN no publica la generación por combustible).
    Generación de cada térmica = combustible del BEN que le corresponde / IAR de su grupo; PPDDSL y PPDFOI (generación
    distribuida) toman la generación de Energy.xlsx y PPDSL y PPFOI el resto del diésel y del fuel oil de las centrales.
    """
    b28 = pd.read_csv(BEN_228, index_col="anio")
    b3 = pd.read_csv(BEN_3)
    A = b3[b3.grupo == "centrales"].set_index("anio")
    B = b3[b3.grupo == "autoproductores"].set_index("anio")
    for t in (b28, A, B):
        if list(t.index) != HIST:
            raise SystemExit("K13: las tablas 2.28 y 3.12-3.15 deben traer 2015-2025")
    control = {}
    for c28, c3 in (("petroleo", "petroleo_kBEP"), ("gas_natural", "gas_natural_kBEP"),
                    ("gas_licuado", "gas_licuado_kBEP"), ("diesel_oil", "diesel_oil_kBEP"), ("fuel_oil", "fuel_oil_kBEP"),
                    ("crudo_reducido", "crudo_reducido_kBEP"), ("bagazo", "productos_cana_kBEP"),
                    ("biogas", "otras_secundarias_kBEP")):
        dif = b28[c28] - A[c3] - B[c3]
        control[c28] = round(float(dif.abs().max()), 3)
        if dif.abs().max() > 2.0:
            raise SystemExit(f"K13: tabla 2.28 ({c28}) no es la suma de 3.12 y 3.14: {dif.round(2).to_dict()}")
    g_a = A.electricidad_GWh - A.hidroenergia_GWh - ben.eolica - ben.solar - ben.biogas
    g_b = B.electricidad_GWh - B.hidroenergia_GWh - ben.biomasa
    dif_t = (g_a + g_b - ben.termica)
    if dif_t.abs().max() > 30:
        raise SystemExit(f"K13: la generación fósil de 3.13 + 3.15 no reproduce la térmica de 2.20: {dif_t.to_dict()}")
    pj = TJ_KBEP / 1000  # PJ por kBEP
    i_a = A[FOSILES_KBEP].sum(axis=1) * pj / (g_a / PJ)
    i_b = B[FOSILES_KBEP].sum(axis=1) * pj / (g_b / PJ)
    comb, gen = {t: {} for t in COMB_K13}, {t: {} for t in COMB_K13}
    for y in HIST:
        ca = {k: A.loc[y, k] * pj for k in FOSILES_KBEP}
        cb = {k: B.loc[y, k] * pj for k in FOSILES_KBEP}
        ia, ib = i_a[y], i_b[y]
        comb["PPCRU"][y] = ca["petroleo_kBEP"] + cb["petroleo_kBEP"]
        gen["PPCRU"][y] = ca["petroleo_kBEP"] / ia + cb["petroleo_kBEP"] / ib
        comb["PPNGS"][y], gen["PPNGS"][y] = ca["gas_natural_kBEP"], ca["gas_natural_kBEP"] / ia
        comb["PPRNGS"][y], gen["PPRNGS"][y] = cb["gas_natural_kBEP"], cb["gas_natural_kBEP"] / ib
        comb["PPRDSL"][y], gen["PPRDSL"][y] = cb["diesel_oil_kBEP"], cb["diesel_oil_kBEP"] / ib
        gen["PPDDSL"][y] = gen_dist["Diesel"][0][y]
        comb["PPDDSL"][y] = gen["PPDDSL"][y] * ia
        comb["PPDSL"][y] = ca["diesel_oil_kBEP"] - comb["PPDDSL"][y]
        gen["PPDSL"][y] = comb["PPDSL"][y] / ia
        gen["PPDFOI"][y] = gen_dist["Fuel Oil"][0][y]
        comb["PPDFOI"][y] = gen["PPDFOI"][y] * ia
        fo_a = ca["fuel_oil_kBEP"] + ca["crudo_reducido_kBEP"] - comb["PPDFOI"][y]
        fo_b = cb["fuel_oil_kBEP"] + cb["crudo_reducido_kBEP"]
        comb["PPFOI"][y] = fo_a + fo_b
        gen["PPFOI"][y] = fo_a / ia + fo_b / ib
    iar = {t: {} for t in COMB_K13}
    for t in COMB_K13:
        for y in HIST:
            if gen[t][y] <= 0 or comb[t][y] <= 0:
                raise SystemExit(f"K13: {t} {y} con generación o combustible no positivo")
            iar[t][y] = r4(comb[t][y] / gen[t][y])
        u = [2023, 2024, 2025]
        v = r4(sum(comb[t][y] for y in u) / sum(gen[t][y] for y in u))
        iar[t].update({y: v for y in PROY})
        if not all(2.0 <= x <= 4.0 for x in iar[t].values()):
            raise SystemExit(f"K13: {t} con rendimiento fuera de 25-50 %: {iar[t]}")
    # K14 revisado: desde 2026 PPCRU genera la media 2023-2025 del crudo de autoproductores (B) y el crudo de centrales
    # eléctricas (A) de 2025, el único año que se puede atribuir a sistemas no incorporados; su IAR es el de esa mezcla
    u = [2023, 2024, 2025]
    cru_cb = sum(B.loc[y, "petroleo_kBEP"] * pj for y in u) / 3
    cru_gb = sum(B.loc[y, "petroleo_kBEP"] * pj / i_b[y] for y in u) / 3
    cru_ca = A.loc[2025, "petroleo_kBEP"] * pj
    cru_ga = cru_ca / i_a[2025]
    iar_cru_proy = r4((cru_cb + cru_ca) / (cru_gb + cru_ga))
    iar_cru_antes = iar["PPCRU"][2026]
    iar["PPCRU"].update({y: iar_cru_proy for y in PROY})
    glp ={y: B.loc[y, "gas_licuado_kBEP"] * pj for y in HIST}
    cru_grupo = {"centrales": {y: round(float(A.loc[y, "petroleo_kBEP"] * pj / i_a[y]), 5) for y in HIST},
                 "autoproductores": {y: round(float(B.loc[y, "petroleo_kBEP"] * pj / i_b[y]), 5) for y in HIST},
                 "kBEP_centrales": {y: float(A.loc[y, "petroleo_kBEP"]) for y in HIST}}
    der = {"TJ_por_kBEP": TJ_KBEP, "control_2_28_menos_3_12_3_14_kBEP": control,
           "generacion_PPCRU_por_grupo_PJ": cru_grupo,
           "IAR_PPCRU_2026_2035": {"K14_revisado": iar_cru_proy, "media_A_y_B_2023_2025": iar_cru_antes,
                                   "combustible_B_medio_PJ": round(float(cru_cb), 5),
                                   "generacion_B_media_PJ": round(float(cru_gb), 5),
                                   "combustible_A_2025_PJ": round(float(cru_ca), 5),
                                   "generacion_A_2025_PJ": round(float(cru_ga), 5)},
           "control_generacion_fosil_menos_termica_2_20_GWh": {y: round(float(dif_t[y]), 1) for y in HIST},
           "generacion_fosil_GWh": {"centrales": {y: float(g_a[y]) for y in HIST},
                                    "autoproductores": {y: float(g_b[y]) for y in HIST}},
           "IAR_grupo": {"centrales": {y: round(float(i_a[y]), 5) for y in HIST},
                         "autoproductores": {y: round(float(i_b[y]), 5) for y in HIST}},
           "rendimiento_grupo_pct": {"centrales": {y: round(100 / float(i_a[y]), 2) for y in HIST},
                                     "autoproductores": {y: round(100 / float(i_b[y]), 2) for y in HIST}},
           "combustible_PJ": {t: {y: round(v, 5) for y, v in d.items()} for t, d in comb.items()},
           "generacion_PJ": {t: {y: round(v, 5) for y, v in d.items()} for t, d in gen.items()},
           "IAR": iar, "GLP_autoproductores_PJ": {y: round(v, 5) for y, v in glp.items()},
           "generacion_GLP_GWh": {y: round(glp[y] / i_b[y] * PJ, 1) for y in HIST},
           "generacion_distribuida_Energy_xlsx": {c: {"fila": v[1], "PJ": {y: v[0][y] for y in HIST}}
                                                   for c, v in gen_dist.items()}}
    return gen, comb, iar, glp, der


def construccion_historica():
    """Potencia que el modelo construye en 2010-2022 en los clones (NewCapacity de Executables/CONA6_0 e INCA6_0, iguales
    en la historia) y que sigue en operación en 2023-2035 (vida útil de 31,5 años). La construye el LP en 2010-2014,
    donde la historia conserva la actividad y el CF del CLG; no depende de K15c, que solo cambia 2023-2035."""
    out = {}
    for e in ("CONA6", "INCA6"):
        ruta = EX / f"{e}_0" / f"{e}_0_Output.csv"
        if not ruta.exists():
            raise SystemExit(f"K15c: falta {ruta}; corra B1 una vez antes de construir K15c")
        o = pd.read_csv(ruta, usecols=["Technology", "Year", "NewCapacity"], low_memory=False)
        o = o[o.Technology.isin(list(K15C_MAPA)) & (o.Year <= 2022)]
        n = o.groupby(["Technology", "Year"]).NewCapacity.max().fillna(0)
        out[e] = {t: round(float(n.loc[t].sum()) if t in n.index.get_level_values(0) else 0.0, 6) for t in K15C_MAPA}
        out[e + "_por_anio_MW"] = {t: {int(y): round(float(v) * 1000, 3) for y, v in n.loc[t].items() if v > 0}
                                   for t in K15C_MAPA if t in n.index.get_level_values(0)}
    if any(abs(out["CONA6"][t] - out["INCA6"][t]) > 1e-6 for t in K15C_MAPA):
        raise SystemExit(f"K15c: la construcción histórica difiere entre CONA6 e INCA6: {out}")
    return out["CONA6"], out["CONA6_por_anio_MW"]


def arriendos_declarados():
    """Arriendos de 2026 del plan de 920 MW (expansion_termica_declarada.csv, tratamiento «arriendo»)."""
    exp_ = pd.read_csv(EXPANSION, dtype=str, keep_default_na=False)
    out = []
    for r in exp_[exp_.tratamiento == "arriendo"].itertuples():
        if r.tecnologia_modelo not in RED_K15 or float(r.MW_en_tope) != 0:
            raise SystemExit(f"K15c: arriendo {r.id} con tecnología o tope inconsistentes")
        out.append({"id": r.id, "central": r.proyecto, "tecnologia": r.tecnologia_modelo,
                    "MW": float(r.MW_residual_arriendo), "inicio": fecha_ddmm(r.fecha_inicio, date(int(r.anio), 1, 1)),
                    "fin_publicado": fecha_ddmm(r.fecha_fin), "fuente_fechas": r.fuente_fechas})
    return out


def potencia_arrendada(arr, variante):
    """{tecnología: {año: GW}} de la potencia arrendada vigente en 2026-2035: MW x fracción del año entre el inicio y el
    fin del contrato (el publicado o, sin fecha publicada, el de la variante)."""
    out = {}
    for a in arr:
        fin = a["fin_publicado"] or FIN_ARRIENDO[variante]
        for y in PROY:
            out.setdefault(a["tecnologia"], {}).setdefault(y, 0.0)
            out[a["tecnologia"]][y] += a["MW"] / 1000 * fraccion_anio(a["inicio"], fin, y)
    return out


def potencia_termica_k15c(capx, ext, capa):
    """K15c: potencia de las térmicas en 2023-2035 = Energy.xlsx (Capacidad_GW 2023, por sistema y combustible) + altas
    de 2024-2025 (termica_altas_2024_2025.csv: central de la capa CENTRALES, combustible de CENACE) + centrales que
    Energy.xlsx omite (mismo CSV, clase «omitida en Energy.xlsx»). En 2023-2025 cuenta la potencia al 31 de diciembre
    (convención de Energy.xlsx y de ARCONEL). Desde 2026 el parque propio de 2025 sigue y la potencia arrendada (las
    centrales arrendadas de 2024-2025 y los arriendos de 2026 del plan de 920 MW) cuenta solo mientras el contrato está
    vigente, como MW x fracción del año (potencia_arrendada). La potencia del modelo es residual + lo que el LP construyó
    en 2010-2022 y sigue en operación (construccion_historica), así que la residual K15c es la fuente menos esa
    construcción: la potencia total del clon en 2023-2025 es la de la fuente. Se corrige donde la potencia total del
    modelo en 2025 (residual CLG + construcción histórica) difiere de la fuente más de 1 MW (más que el redondeo de
    Energy.xlsx). Devuelve la residual por tecnología (GW) del caso central, la construcción histórica, las filas de
    conciliación y las derivaciones (con la residual de la variante de arriendos hasta 2032)."""
    hist, hist_anio = construccion_historica()
    altas = pd.read_csv(ALTAS, dtype=str, keep_default_na=False)
    usados, filas_altas, arr = {}, [], []
    for _, a in altas.iterrows():
        r = ext[(ext.NombreCent == a.central_capa) & (ext.TipoCentra == "Térmica")]
        if len(r) != 1:
            raise SystemExit(f"K15c: {a.central_capa} aparece {len(r)} veces en el extracto de la capa")
        r = r.iloc[0]
        anio_capa = int(r.FechaInici[:4])
        if anio_capa != int(a.anio_modelo):
            raise SystemExit(f"K15c: {a.central_capa} inicia en {anio_capa} en la capa y el CSV dice {a.anio_modelo}")
        if a.modalidad not in ("arrendamiento", "propia"):
            raise SystemExit(f"K15c: {a.central_capa} sin modalidad (arrendamiento o propia)")
        mw = float(r.PotenciaNo)
        inicio = date(anio_capa, int(r.FechaInici[4:6]), int(r.FechaInici[6:]))
        filas_altas.append({"central": a.central_capa, "tecnologia": a.tecnologia_modelo, "anio": anio_capa,
                            "MW_nominal": mw, "MW_efectiva": float(r.PotenciaEf), "incluida": a.incluida,
                            "fecha_inicio_capa": inicio.strftime("%d-%m-%Y"), "activa_capa": r.EstaActiva,
                            "combustible": a.combustible, "clase_registro": a.clase_registro,
                            "modalidad": a.modalidad, "fecha_fin_publicada": a.fecha_fin or None,
                            "fuente_modalidad_y_fin": a.fuente_modalidad_y_fin})
        if a.incluida == "sí":
            usados.setdefault(a.tecnologia_modelo, []).append((anio_capa, mw / 1000, a.modalidad, a.clase_registro))
            if a.modalidad == "arrendamiento":
                arr.append({"id": "capa", "central": a.central_capa, "tecnologia": a.tecnologia_modelo, "MW": mw,
                            "inicio": inicio, "fin_publicado": fecha_ddmm(a.fecha_fin),
                            "fuente_fechas": a.fuente_modalidad_y_fin})
    arr += arriendos_declarados()
    arrendada = {v: potencia_arrendada(arr, v) for v in FIN_ARRIENDO}
    serie, serie_var, filas = {}, {}, []
    der = {"altas": filas_altas, "construccion_historica_GW": hist, "construccion_historica_por_anio_MW": hist_anio,
           "arriendos": [{**{k: (v.strftime("%d-%m-%Y") if isinstance(v, date) else v) for k, v in a.items()},
                          **{f"MW_equivalentes_fin_{v}": {y: round(a["MW"] * fraccion_anio(
                              a["inicio"], a["fin_publicado"] or FIN_ARRIENDO[v], y), 3) for y in PROY}
                             for v in FIN_ARRIENDO}} for a in arr],
           "vigencia_arriendos": {v: F_FIN_ARRIENDO[v] for v in FIN_ARRIENDO}, "variante_central": VARIANTE_CENTRAL}
    for t, clave in K15C_MAPA.items():
        e23 = capx[clave][0][2023]
        s = {y: e23 + sum(g for a_, g, _, _ in usados.get(t, []) if a_ <= y) for y in (2023, 2024, 2025)}
        propia = e23 + sum(g for a_, g, m_, _ in usados.get(t, []) if a_ <= 2025 and m_ == "propia")
        s_var = {}
        for v in FIN_ARRIENDO:
            s_var[v] = {**s, **{y: propia + arrendada[v].get(t, {}).get(y, 0.0) for y in PROY}}
        s = s_var[VARIANTE_CENTRAL]
        clg_ = vector("ResidualCapacity", t)
        if {y: clg_[y] for y in K15C_ANIOS} != {y: vector("ResidualCapacity", t, carpeta=ENERGIA / "A2_Output_Params" /
                                                            "NDCINC")[y] for y in K15C_ANIOS}:
            raise SystemExit(f"K15c: la residual de {t} difiere entre NDCCON y NDCINC")
        total_antes = clg_[2025] + hist[t]
        dif = total_antes - s[2025]
        corrige = abs(dif) > 0.001
        if corrige:
            if s[2023] - hist[t] < 0:
                raise SystemExit(f"K15c: {t} construyó en 2010-2022 más que la potencia de la fuente")
            serie[t] = {y: round(s[y] - hist[t], 6) for y in K15C_ANIOS}
            serie_var[t] = {v: {y: round(s_var[v][y] - hist[t], 6) for y in K15C_ANIOS} for v in FIN_ARRIENDO}
        elif any(abs(s[y] - s[2025]) > 1e-9 for y in PROY):
            raise SystemExit(f"K15c: {t} coincide en 2025 pero tiene arriendos que vencen; prevea su fila")
        u_ = usados.get(t, [])
        filas.append({"tecnologia": t, "fila_Energy_xlsx": f"{clave[0]} / {clave[1]} ({capx[clave][1]})",
                      "Energy_xlsx_2023_MW": round(e23 * 1000, 2),
                      "omitidas_en_Energy_xlsx_MW": round(sum(g for _, g, _, c in u_ if c != "alta 2024-2025") * 1000, 2),
                      "altas_2024_MW": round(sum(g for a_, g, _, c in u_ if a_ == 2024 and c == "alta 2024-2025") * 1000, 2),
                      "altas_2025_MW": round(sum(g for a_, g, _, c in u_ if a_ == 2025 and c == "alta 2024-2025") * 1000, 2),
                      "de_ello_arrendada_2025_MW": round(sum(g for _, g, m_, _ in u_ if m_ == "arrendamiento") * 1000, 2),
                      "fuente_2025_MW": round(s[2025] * 1000, 2), "CLG_2025_MW": round(clg_[2025] * 1000, 2),
                      "CLG_2022_MW": round(clg_[2022] * 1000, 2),
                      "construida_por_el_modelo_2010_2022_MW": round(hist[t] * 1000, 2),
                      "potencia_total_antes_2025_MW": round(total_antes * 1000, 2),
                      "desvio_CLG_pct": round(dif / s[2025] * 100, 2),
                      "decision": "corrige 2023-2035" if corrige else "conserva (coincide con la fuente)",
                      "K15c_2023_MW": round(s[2023] * 1000, 2), "K15c_2024_MW": round(s[2024] * 1000, 2),
                      "K15c_2025_MW": round(s[2025] * 1000, 2),
                      **{f"K15c_{y}_MW_fin_{v}": round(s_var[v][y] * 1000, 2) for v in FIN_ARRIENDO
                         for y in (2026, 2027, 2028, 2029, 2032, 2035)},
                      "residual_K15c_2025_MW": round((s[2025] - hist[t]) * 1000, 2) if corrige else None})
    clg_lpg = vector("ResidualCapacity", "PPILPG")
    filas.append({"tecnologia": "PPILPG", "fila_Energy_xlsx": "sin fila de GLP en Energy.xlsx",
                  "CLG_2025_MW": round(clg_lpg[2025] * 1000, 2), "CLG_2022_MW": round(clg_lpg[2022] * 1000, 2),
                  "decision": "conserva (sin fuente por combustible para contrastar; su actividad es contable)"})
    # contraste de totales con ARCONEL 2025 y la capa (potencia nominal con nominal)
    arc = pd.read_csv(AQUI / "arconel2025_potencia_termica.csv")

    def a25(concepto, anio=2025):
        x = arc[(arc.concepto == concepto) & (arc.anio == anio)]
        return float(x.potencia_nominal_MW.iloc[0])
    k15 = {f["tecnologia"]: f.get("K15c_2025_MW") for f in filas}
    k15["PPILPG"] = round(clg_lpg[2025] * 1000, 2)
    victoria = ext[(ext.NombreCent == "Victoria II") & (ext.TipoCentra == "Térmica")]
    victoria = float(victoria.PotenciaNo.iloc[0]) if len(victoria) else 0.0
    capa25 = capa["termica_SNI_generadora"].potencia_nominal_MW
    capa_d25 = capa["termica_SNI_distribuidora"].potencia_nominal_MW
    totales = [
        ("Térmicas de la red S.N.I. (PPFOI + PPDSL + PPNGS)", sum(k15[t] for t in RED_K15),
         float(capa25[2025]) - victoria,
         f"capa CENTRALES, térmicas del S.N.I. de generadoras activas con inicio hasta 2025 ({float(capa25[2025]):.2f} MW) "
         f"menos Victoria II ({victoria:g} MW, sin producción desde el 06-05-2017 según la propia capa)"),
        ("Generación distribuida (PPDDSL + PPDFOI)", k15["PPDDSL"] + k15["PPDFOI"], float(capa_d25[2025]),
         "capa CENTRALES, térmicas del S.N.I. de distribuidoras (Aníbal Santos, Álvaro Tinajero, G. Hernández, Catamayo, "
         "Lligua)"),
        ("Autoproductores (PPCRU + PPRDSL + PPRNGS + PPILPG)", k15["PPCRU"] + k15["PPRDSL"] + k15["PPRNGS"] + k15["PPILPG"],
         a25("Térmica fósil de autogeneradoras"), "ARCONEL 2025, tabla 23, p. 34, autogeneradoras térmicas (nominal)"),
        ("Térmica fósil total del modelo", sum(v for v in k15.values() if v),
         a25("No renovable (térmica fósil) nacional"), "ARCONEL 2025, tabla 20, p. 33, no renovable (nominal)")]
    for nombre, modelo, fuente_v, f in totales:
        filas.append({"tecnologia": nombre, "fuente_2025_MW": round(fuente_v, 2), "K15c_2025_MW": round(modelo, 2),
                      "desvio_CLG_pct": None, "decision": f"contraste: {f}",
                      "desvio_K15c_frente_a_fuente_pct": round((modelo - fuente_v) / fuente_v * 100, 2)})
    der["series_GW"] = serie
    der["series_GW_por_vigencia_de_arriendos"] = serie_var
    der["victoria_II_MW"] = victoria
    return serie, hist, filas, der


def construir():
    # sandbox PNUD: los extractos por central de la capa CENTRALES (uso autorizado) no se publican
    faltan = [p.name for p in (CENTRALES_EXTRACTO, ALTAS) if not p.exists()]
    if faltan:
        raise SystemExit("construir_correcciones.py necesita los extractos por central de la capa CENTRALES de "
                         "ARCONEL (%s), que no se distribuyen: el repositorio público no lleva filas por central "
                         "(decisión del 13-10-2026). Su resultado, correcciones_K.csv, está versionado y es el insumo "
                         "de generar_referencias.py y de las pruebas. Para regenerarlo hay que correr antes "
                         "extraer_centrales.py con la capa local (uso autorizado)." % ", ".join(faltan))
    proy, capx, fp_eol, cen, capbio, ben, pot, sgda, of, gen_dist, capa = fuentes()
    F = Filas()
    ext = pd.read_csv(CENTRALES_EXTRACTO, dtype=str)

    def capa_central(nombre, tipo):
        """(fecha de inicio dd-mm-aaaa, potencia nominal, potencia efectiva, tipo de empresa) de una central de la capa."""
        r = ext[(ext.NombreCent == nombre) & (ext.TipoCentra == tipo)]
        if len(r) != 1:
            raise SystemExit(f"capa CENTRALES: {nombre} ({tipo}) aparece {len(r)} veces en el extracto")
        r = r.iloc[0]
        f = r.FechaInici
        return f"{f[6:]}-{f[4:6]}-{f[:4]}", float(r.PotenciaNo), float(r.PotenciaEf), r.TipoEmpres
    insumos = {"Energy.xlsx": {"md5": md5(ENERGY_XLSX), **{k: v for k, v in proy.items()},
                               "Capacidad_GW": {f"{k[0]} {k[1]} ({v[1]})": v[0] for k, v in capx.items()},
                               "FP_eolico_SNI": {"fila": fp_eol[1], "valores": fp_eol[0]}},
               "CENACE_csv": {"md5": md5(CENACE_CSV), "GWh_neta": {f"{a} {c}": v for (a, c), v in cen.items()}},
               "capacidad_bio_csv": {"md5": md5(CAP_CSV),
                                     "MW": {t: {y: d["MW"] for y, d in s.items()} for t, s in capbio.items()}},
               "BEN_csv_md5": {p.name: md5(p) for p in (BEN_CSV, BEN_POT, BEN_SGDA, BEN_228, BEN_3)},
               "capa_CENTRALES": {"cita": F_CAPA, "centrales_series_md5": md5(CENTRALES_SERIES),
                                  "centrales_extracto_md5": md5(CENTRALES_EXTRACTO),
                                  **pd.read_csv(CENTRALES_EXTRACTO)[["md5_shp", "md5_dbf", "fecha_dbf"]].iloc[0]
                                  .to_dict()},
               "oficiales_md5": {e: {"datafile": md5(o.ruta_df), "output": md5(o.ruta_out)} for e, o in of.items()},
               "A2_NDCCON_md5": {p: md5(A2 / f"{p}.csv") for p in
                                 (UL, LL, "CapacityFactor", "ResidualCapacity", "AvailabilityFactor",
                                  "CapacityToActivityUnit", "OperationalLife", MININV, "InputActivityRatio", EAR)},
               "derivaciones": {}}
    der = insumos["derivaciones"]
    ul = {}   # techos finales por tecnología antes de K09 (para la escala hidro)
    ll = {}
    # ---------------- K15c potencia residual de las térmicas 2023-2035 (las filas se escriben después de K14) ---------
    rc15, hist15, conc15, der15 = potencia_termica_k15c(capx, ext, capa)
    der["K15c"] = der15
    insumos["conciliacion_termica_2025"] = conc15

    def rc_k15(t):
        """Potencia con la que se fijan los CF: la residual CLG de A2 y, en 2023-2035 donde K15c corrige, la residual
        K15c más lo que el modelo construyó en 2010-2022 (= potencia de la fuente)."""
        return {**vector("ResidualCapacity", t), **{y: v + hist15[t] for y, v in rc15.get(t, {}).items()}}

    # ---------------- K01 Mazar en PPHDAMAM 2011-2025 -------------------------------------------------------
    d_mazar = r4(proy["G4"] / 1000 / PJ)  # MWh -> PJ
    ul_of, ll_of = vector(UL, "PPHDAMAM"), vector(LL, "PPHDAMAM")
    yrs = list(range(2011, 2026))
    ul["PPHDAMAM"] = {**ul_of, **{y: r4(ul_of[y] + d_mazar) for y in yrs}}
    ll["PPHDAMAM"] = {**ll_of, **{y: r4(ll_of[y] / ul_of[y] * ul["PPHDAMAM"][y]) for y in yrs}}
    der["K01"] = {"mazar_PJ": d_mazar}
    f_k01 = f"{F_PP} G4 = {proy['G4']} MWh y H4 = {proy['H4']} (Mazar, 170 MW); techo oficial de {F_A2}"
    F.add("K01a", "A2", UL, "PPHDAMAM", yrs, [ul["PPHDAMAM"][y] for y in yrs], "PJ", f_k01,
          f"Techo oficial + {d_mazar} PJ (663,5 GWh de Mazar) en 2011-2025; desde 2026 el oficial ya lo incluye. "
          "2015-2025 los reemplaza K09h")
    F.add("K01b", "A2", LL, "PPHDAMAM", yrs, [ll["PPHDAMAM"][y] for y in yrs], "PJ", f_k01,
          "Piso = 0,98 x techo K01a (razón piso/techo oficial de PPHDAMAM, la misma que el oficial aplica en 2026 "
          "con Mazar: +2,3409 PJ). 2015-2025 los acota K09l")

    # ---------------- K02 y K03 Coca Codo Sinclair y Sopladora en PPHRORAB --------------------------------------
    cap_k = (proy["F6"] + proy["F8"]) / 1000          # Coca Codo + Sopladora, GW
    ul_of, ll_of = vector(UL, "PPHRORAB"), vector(LL, "PPHRORAB")
    cf, af = vector("CapacityFactor", "PPHRORAB"), vector("AvailabilityFactor", "PPHRORAB")
    cau = vector("CapacityToActivityUnit", "PPHRORAB")[None]
    mi_of = vector(MININV, "PPHRORAB")
    techo_2029 = ul_of[2029]
    sopl16 = r4(cen[(2016, "Sopladora")] / PJ)
    ul_k02 = {2016: r4(ul_of[2016] + sopl16)}
    for y in range(2017, 2026):
        ul_k02[y] = r4(min(techo_2029, cap_k * cf[y] * af[y] * cau))
    for y in (2026, 2027, 2028):
        ul_k02[y] = r4(techo_2029)
    yrs = list(range(2016, 2029))
    ul["PPHRORAB"] = {**ul_of, **ul_k02}
    ll["PPHRORAB"] = {**ll_of, **{y: r4(ll_of[y] / ul_of[y] * ul_k02[y]) for y in yrs}}
    lim16 = cap_k * cf[2016] * af[2016] * cau
    der["K02_K03"] = {"capacidad_GW": cap_k, "techo_2029_PJ": techo_2029,
                      "sopladora_2016_CENACE_GWh_neta": cen[(2016, "Sopladora")],
                      "coca_codo_CENACE_GWh_neta": {y: cen[(y, "Coca Codo Sinclair")] for y in (2016, 2017)},
                      "sopladora_implicita_GWh": {y: round((ul_k02[y] - ul_of[y]) * PJ, 1) for y in yrs},
                      "limite_2016_1987MW_CF_oficial_GWh": round(lim16 * PJ, 1),
                      "techo_K02a_2016_GWh": round(ul_k02[2016] * PJ, 1),
                      "inversion_minima_oficial_GW": {y: mi_of[y] for y in (2015, 2016, 2017, 2026, 2027, 2028)}}
    if lim16 < ul_k02[2016]:
        raise SystemExit("K03: con 1,987 GW y el CF oficial de 2016 no se alcanza el techo de K02a")
    f_k02 = (f"{F_CEN16}: Sopladora U1 + U2 + U3 = {cen[(2016, 'Sopladora')]:.2f} GWh netos en 2016; "
             f"{F_PP} F6 = {proy['F6']} MW, F8 = {proy['F8']} MW, H8 = {proy['H8']}; CF, AF y techo 2029 de {F_A2}")
    F.add("K02a", "A2", UL, "PPHRORAB", yrs, [ul_k02[y] for y in yrs], "PJ", f_k02,
          f"Coca Codo + Sopladora. 2016: techo oficial ({ul_of[2016] * PJ:.1f} GWh, igual a la producción neta de "
          f"Coca Codo en CENACE) + {sopl16} PJ de Sopladora ({cen[(2016, 'Sopladora')]:.2f} GWh netos). 2017-2025: "
          "min(33,7883; 1,987 GW x CF oficial x 0,95 x 31,54); en 2017 da "
          f"{(ul_k02[2017] - ul_of[2017]) * PJ:.1f} GWh de Sopladora frente a {cen[(2017, 'Sopladora')]:.2f} GWh netos "
          "de CENACE (Informe Anual 2017, tabla 3.3, p. 73). 2026-2028: 33,7883 (techo 2029). 2016-2025 los acota K09h")
    F.add("K02b", "A2", LL, "PPHRORAB", yrs, [ll["PPHRORAB"][y] for y in yrs], "PJ", f_k02,
          "Piso = 0,99 x techo K02a (razón piso/techo oficial de PPHRORAB). 2016-2025 los acota K09l")
    mi_yrs = [2015, 2016, 2017, 2026, 2027, 2028]
    mi_val = [0.0, cap_k, 0.0, 0.0, 0.0, 0.0]
    f_ccs, f_sop = capa_central("Coca Codo Sinclair", "Hidráulica"), capa_central("Sopladora", "Hidráulica")
    der["K02_K03"]["capa_CENTRALES"] = {"Coca Codo Sinclair": f_ccs, "Sopladora": f_sop}
    F.add("K03", "A2", MININV, "PPHRORAB", mi_yrs, mi_val, "GW",
          f"{F_CEN16_FECHA}; BEN 2025, tabla 2.23, p. 67 (hidráulica efectiva 2.402 MW en 2015 y 4.418 MW en 2016); "
          f"{F_PP} F6 = {proy['F6']} MW y F8 = {proy['F8']} MW; inversión mínima oficial de {F_A2} (0,5 GW 2015-2017 y "
          f"0,1623 GW 2026-2028); {F_CAPA}: fecha de inicio de Coca Codo Sinclair {f_ccs[0]} ({f_ccs[1]:g} MW) y de "
          f"Sopladora {f_sop[0]} ({f_sop[1]:g} MW)",
          "Coca Codo Sinclair (1.500 MW) y Sopladora (487 MW) entran juntas en 2016: 0 en 2015, 1,987 GW en 2016 y 0 en "
          "2017; 0 en 2026-2028 (quita la reentrada de Sopladora). Con 1,987 GW el CF oficial de 2016 (0,2393) ya "
          f"alcanza el techo K02a ({lim16 * PJ:.1f} GWh frente a {ul_k02[2016] * PJ:.1f}), así que la corrección de CF "
          "de 2016 de la versión anterior (K03b) se retira. Primera energía vendida en mayo de 2016 (CENACE); la capa "
          f"CENTRALES da como fecha de inicio {f_ccs[0]} (Coca Codo Sinclair) y {f_sop[0]} (Sopladora): las dos en "
          "2016, lo que confirma el año de K03")

    # ---------------- K04 PPIBIM y K05 PPBIM -------------------------------------------------------------------------
    af_b = vector("AvailabilityFactor", "PPIBIM")
    cau_b = vector("CapacityToActivityUnit", "PPIBIM")[None]
    yrs = list(range(2015, 2036))
    cap_b = {y: r4(capbio["PPIBIM"][y]["MW"] / 1000) for y in yrs}
    prom = ben.loc[2019:2025, "biomasa"].mean()
    piso_b = {y: r4(ben.loc[y, "biomasa"] / PJ) if y <= 2025 else r4(prom / PJ) for y in yrs}
    razon = vector(UL, "PPIBIM")[2022] / vector(LL, "PPIBIM")[2022]
    techo_b = {y: r4(piso_b[y] * razon) for y in yrs}
    cf_b = {y: techo4(piso_b[y] / (cap_b[y] * af_b[y] * cau_b)) for y in yrs}
    if max(cf_b.values()) >= 1:
        raise SystemExit("K04: CF >= 1")
    bio_capa = {n: capa_central(n, "Biomasa") for n in ("San Carlos", "Ecudos A-G", "Ecoelectric")}
    bio_nom, bio_ef = sum(v[1] for v in bio_capa.values()), sum(v[2] for v in bio_capa.values())
    der["K04"] = {"capacidad_GW": cap_b, "estado_capacidad": {y: capbio["PPIBIM"][y]["estado"] for y in yrs},
                  "capacidad_oficial_2014_GW": vector("ResidualCapacity", "PPIBIM")[2014],
                  "promedio_2019_2025_GWh": round(prom, 2), "razon_techo_piso_oficial": razon, "CF": cf_b,
                  "contraste_capa_CENTRALES": {
                      "centrales": bio_capa, "nominal_MW": round(bio_nom, 2), "efectiva_MW": round(bio_ef, 2),
                      "lectura": "Tres autogeneradoras del SNI (San Carlos, Ecudos A-G y Ecoelectric), todas con fecha de "
                                 "inicio anterior a 2015 y activas: confirma el parque, no la serie. Su potencia actual "
                                 "(155,3 MW nominales y 147,7 efectivos) se compara nominal con nominal (ARCONEL 2025, "
                                 "tabla 18: 156,05 MW) y efectiva con efectiva (tabla 19: 148,50 MW); la serie anual es la "
                                 "de la Estadística Anual y Multianual 2025"}}
    f_k04 = (f"{F_CAP} (PPIBIM: ARCONEL, Estadística Anual y Multianual 2025, tablas 18-19, p. 33, potencia efectiva "
             f"2016-2025); {F_BEN} (biomasa)")
    F.add("K04a", "A2", "ResidualCapacity", "PPIBIM", yrs, [cap_b[y] for y in yrs], "GW",
          f"{f_k04}; contraste: {F_CAPA}",
          "Potencia efectiva de ARCONEL (Estadística Anual y Multianual 2025, tabla 19): 136,40 MW en 2016-2023, 147,70 "
          "MW en 2024 y 148,50 MW en 2025 (desde 2026 se mantiene); 2015 repite 2016 porque la serie empieza en 2016 y la "
          f"capa CENTRALES no registra centrales de biomasa nuevas en 2015-2016. La capa confirma el parque (tres "
          f"autogeneradoras del SNI con inicio entre 1968 y 2005, {bio_nom:.1f} MW nominales y {bio_ef:.1f} efectivos "
          "hoy), no la serie. 2010-2014 conservan 0,2125 GW del CLG sin fuente (límite conocido, ver LEEME)")
    F.add("K04b", "A2", LL, "PPIBIM", yrs, [piso_b[y] for y in yrs], "PJ", f"{F_BEN} (biomasa)",
          "Piso = biomasa del BEN por año 2015-2025; 1,3947 PJ (387,4 GWh, promedio 2019-2025) desde 2026")
    F.add("K04c", "A2", UL, "PPIBIM", yrs, [techo_b[y] for y in yrs], "PJ", f"{F_BEN} (biomasa)",
          f"Techo = {razon:.4f} x piso (razón techo/piso oficial de PPIBIM)")
    F.add("K04d", "A2", "CapacityFactor", "PPIBIM", yrs, [cf_b[y] for y in yrs], "fracción",
          f"Regla CF = generación observada / (capacidad x AF x CapacityToActivityUnit) con K04a y K04b; {F_A2} "
          "(AF 0,95, CAU 31,536)",
          "Generación / (capacidad del año x 8.760 h x 0,95): el modelo multiplica CF por AF. Nunca CF = 1; redondeo "
          "hacia arriba en el cuarto decimal")
    yrs = list(range(2011, 2036))
    F.add("K05", "A2", "ResidualCapacity", "PPBIM", yrs, [0.0] * len(yrs), "GW",
          f"A1_Inputs A-O_Parametrization, Secondary Techs I55 (modo Flat, sin fuente) y techo de actividad 0 en {F_A2}",
          "Capacidad 0 desde 2011: los 100 MW no generan (techo 0) y no tienen fuente. 2010 intacto")

    # ---------------- K06 PPBGS ---------------------------------------------------------------------------------------
    yrs = list(range(2011, 2036))
    cap_g = {y: (0.0 if y < 2016 else capbio["PPBGS"][y]["MW"] / 1000) for y in yrs}
    bg_capa = {n: capa_central(n, "Biogas") for n in ("El Inga", "Pichacay")}
    der["K06_capa_CENTRALES"] = {
        "centrales": bg_capa, "nominal_MW": round(sum(v[1] for v in bg_capa.values()), 2),
        "efectiva_MW": round(sum(v[2] for v in bg_capa.values()), 2),
        "lectura": "Dos generadoras del SNI: El Inga (Gasgreen) con inicio el 26-02-2016 (CENACE: cobra desde febrero de "
                   "2016) y Pichacay (EMAC-BGP) el 26-07-2017 (CENACE: cobra desde junio de 2017). Potencia actual 11,12 MW "
                   "nominales, igual a la nominal de ARCONEL 2024-2025 (Estadística Anual y Multianual 2025, tabla 18). "
                   "La capa confirma el parque, no la serie: la serie anual es la efectiva de la tabla 19"}
    F.add("K06a", "A2", "ResidualCapacity", "PPBGS", yrs, [cap_g[y] for y in yrs], "GW",
          f"{F_CAP} (PPBGS: ARCONEL, Estadística Anual y Multianual 2025, tablas 18-19, p. 33, potencia efectiva "
          f"2016-2025; fechas de entrada en CENACE 2016, tabla 2.11, p. 48, y CENACE 2017, tabla 2.11, p. 54); {F_BEN} "
          f"(biogás desde 2016); contraste: {F_CAPA}",
          "0 en 2011-2015 (el BEN no registra biogás antes de 2016). Potencia efectiva de ARCONEL (tabla 19): 1,76 MW "
          "en 2016 (Gasgreen, El Inga, cobra desde febrero de 2016); 6,50 MW en 2017-2020 (Pichacay entra en junio de "
          "2017); 7,20 MW en 2021-2023; 9,97 MW en 2024-2025 (desde 2026 se mantiene). La capa CENTRALES confirma las "
          f"fechas (El Inga {bg_capa['El Inga'][0]}, Pichacay {bg_capa['Pichacay'][0]}) y la nominal de 2024-2025 "
          "(11,12 MW, igual a la tabla 18); confirma el parque, no la serie. 2010 intacto", dec=5)
    piso_g = {y: 0.0 for y in ANIOS}
    for y in ANIOS:
        if 2016 <= y <= 2025:
            piso_g[y] = r4(ben.loc[y, "biogas"] / PJ)
        elif y >= 2026:
            piso_g[y] = r4(ben.loc[2025, "biogas"] / PJ)
    techo_g = {y: r4(piso_g[y] * razon) for y in range(2016, 2036)}
    F.add("K06b", "A2_nueva", LL, "PPBGS", ANIOS, [piso_g[y] for y in ANIOS], "PJ", f"{F_BEN} (biogás)",
          "PPBGS no tiene piso en el oficial: filas nuevas. 0 en 2010-2015 (equivale a no tener piso); BEN 2016-2025; "
          "43 GWh/año (BEN 2025) desde 2026")
    F.add("K06c", "A2", UL, "PPBGS", list(range(2016, 2036)), [techo_g[y] for y in range(2016, 2036)], "PJ",
          f"{F_BEN} (biogás); techo oficial 0 en {F_A2}",
          f"Techo = {razon:.4f} x piso (misma razón que PPIBIM); el oficial tiene techo 0 todos los años")
    cfg, afg = vector("CapacityFactor", "PPBGS"), vector("AvailabilityFactor", "PPBGS")
    caug = vector("CapacityToActivityUnit", "PPBGS")[None]
    yrs_cf = [y for y in range(2016, 2036) if cap_g[y] * cfg[y] * afg[y] * caug < techo_g[y] - 1e-9]
    cf_g = {y: techo4(techo_g[y] / (cap_g[y] * afg[y] * caug)) for y in yrs_cf}
    if cf_g and max(cf_g.values()) >= 1:
        raise SystemExit("K06: CF >= 1")
    if yrs_cf:
        F.add("K06d", "A2", "CapacityFactor", "PPBGS", yrs_cf, [cf_g[y] for y in yrs_cf], "fracción",
              f"Regla del modelo CF = techo / (capacidad x AF x CapacityToActivityUnit) con K06a y K06c; {F_A2} "
              "(AF 1,0, CAU 31,536)",
              "Solo en los años en que el CF oficial no alcanza el techo con la potencia del año; en los demás queda el "
              "oficial. Redondeo hacia arriba en el cuarto decimal; nunca CF = 1")
    der["K06"] = {"capacidad_GW": {y: cap_g[y] for y in range(2016, 2036)}, "CF_K06d": cf_g}

    # ---------------- K07 PPHD ------------------------------------------------------------------------------------------
    f_k07 = (f"{F_A2}/CapacityToActivityUnit.csv y OperationalLife.csv: 20 centrales PP con 31,536 o 31,54 "
             "(1 GW x 8.760 h x 3.600 s = 31,536 PJ/año); PPHD no tiene fila (default 1)")
    F.add("K07a", "A2_nueva", "CapacityToActivityUnit", "PPHD", None, [31.536], "PJ/GW", f_k07,
          "Con el default 1 el modelo construye 19,6 GW ficticios de PPHD cada año; no cambia generación ni emisiones")
    F.add("K07b", "A2_nueva", "OperationalLife", "PPHD", None, [31.536], "años", f_k07,
          "Convención del modelo para las centrales PP (igual a PPIHD) [verificar vida útil real de la hidro "
          "distribuida]; sin efecto si no se construye capacidad nueva")

    # ---------------- K13 rendimientos térmicos observados (derivación; las filas se escriben después de K08) ---------
    gen13, comb13, iar13, glp13, der13 = rendimientos_observados(ben, gen_dist)
    for c, t in (("Diesel", "PPDDSL"), ("Fuel Oil", "PPDFOI")):  # la fila de Energy.xlsx está entre piso y techo
        ul_t, ll_t = vector(UL, t), vector(LL, t)
        fuera = [y for y in range(2015, 2023) if not ll_t[y] - 1e-3 <= gen_dist[c][0][y] <= ul_t[y] + 1e-3]
        if fuera:
            raise SystemExit(f"K13: Generación_PJ {c} distribuida fuera del piso y techo oficiales de {t} en {fuera}")
    der["K13"] = der13

    # ---------------- K08 y K08g factores de emisión 2011-2035 ---------------------------------------------------------
    iar = leer_a2("InputActivityRatio")
    sg = pd.read_excel(ENERGIA / "B1_Scenario_Config.xlsx", sheet_name="SmartGrid")
    yrs = list(range(2011, 2036))
    comprobacion = {}
    iar_of = {}
    # K13g (decisión d): PPDDSL 2011-2014 lleva en el CLG el IAR copiado del gas natural (4,3645 y 4,0603); toma el IAR
    # diésel del propio modelo, el de PPDSL en A2 NDCCON (2,781 en 2011-2012 y 2,5349 en 2013-2014)
    K13G = range(2011, 2015)
    iar_k13g = {y: vector("InputActivityRatio", "PPDSL", fuel="E1DSL")[y] for y in K13G}
    if iar_k13g != {y: vector("InputActivityRatio", "PPDSL", fuel="E1DSL",
                              carpeta=ENERGIA / "A2_Output_Params" / "NDCINC")[y] for y in K13G}:
        raise SystemExit("K13g: el IAR de PPDSL 2011-2014 difiere entre NDCCON y NDCINC")
    gases = [("CO2e_PP", "K08"), ("CO2_PP", "K08g"), ("CH4_PP", "K08g"), ("N2O_PP", "K08g")]
    for t in TERMICAS_K08:
        d = iar[(iar.TECHNOLOGY == t) & (iar.MODE_OF_OPERATION.astype(str) == "1")]
        fuels = sorted(set(d.FUEL))
        if len(fuels) != 1 or fuels[0] not in IPCC or fuels[0] != COMB_K13[t]:
            raise SystemExit(f"{t}: combustible {fuels} no previsto")
        comb = fuels[0]
        nombre, f_co2, f_ch4, f_n2o = IPCC[comb]
        f_co2e = f_co2 + PCG["CH4"] * f_ch4 / 1000 + PCG["N2O"] * f_n2o / 1000
        i_of = {int(y): float(v) for y, v in zip(d.YEAR, d.Value)}
        for y in ANIOS:  # el IAR del datafile oficial es el de A2 (B1 no lo reescribe)
            if abs(of["NDCCON"].v("InputActivityRatio", t, y, comb) - i_of[y]) > 1e-9 or \
                    abs(of["NDCINC"].v("InputActivityRatio", t, y, comb) - i_of[y]) > 1e-9:
                raise SystemExit(f"{t}: el IAR del datafile difiere del de A2 en {y}")
        iar_of[t] = i_of
        # 2011-2014 oficial (K13g en PPDDSL); 2015-2035 K13a
        i = {y: (iar13[t][y] if y >= 2015 else iar_k13g[y] if t == "PPDDSL" and y in K13G else i_of[y]) for y in ANIOS}
        valores = {"CO2e_PP": {y: f_co2e * i[y] / 1000 for y in yrs}, "CO2_PP": {y: f_co2 * i[y] / 1000 for y in yrs},
                   "CH4_PP": {y: f_ch4 * i[y] / 1e6 for y in yrs}, "N2O_PP": {y: f_n2o * i[y] / 1e6 for y in yrs}}
        comprobacion[t] = {"combustible": comb, "factores_t_TJ": {"CO2": f_co2, "CH4_kg": f_ch4, "N2O_kg": f_n2o,
                                                                    "CO2e_AR5": round(f_co2e, 4)},
                           "IAR_usado": {y: i[y] for y in (2011, 2015, 2017, 2019, 2022, 2025, 2026, 2035)},
                           "IAR_oficial": {y: i_of[y] for y in (2011, 2015, 2017, 2019, 2022, 2025, 2026, 2035)},
                           "factor_implicito_oficial_t_TJ": {}}
        for gas, cod in gases:
            texto_sg = f"{t} ; {gas}"
            fila = sg[(sg.Scenario == "NDCCON") & (sg.Parameter == EAR) &
                      (sg.Tech_Set.astype(str).str.replace(" ", "") == texto_sg.replace(" ", ""))]
            e10 = iguales_historia(of, EAR, t, [2010], gas)[0]
            ofi = {y: of["NDCCON"].v(EAR, t, y, gas) for y in (2012, 2017, 2018, 2019)}
            if gas == "CO2e_PP":
                comprobacion[t]["factor_implicito_oficial_t_TJ"] = {y: round(ofi[y] / i_of[y] * 1000, 2) for y in ofi}
            codigo = f"{cod}-{t}" if gas == "CO2e_PP" else f"{cod}-{t}-{gas.split('_')[0]}"
            divisor = "1000" if gas in ("CO2e_PP", "CO2_PP") else "10^6"
            f_gas = {"CO2e_PP": f"{f_co2e:.4f} t CO2-eq/TJ (CO2 {f_co2} t/TJ + 28 x CH4 {f_ch4} kg/TJ + 265 x N2O "
                                f"{f_n2o} kg/TJ)", "CO2_PP": f"{f_co2} t CO2/TJ", "CH4_PP": f"{f_ch4} kg CH4/TJ",
                     "N2O_PP": f"{f_n2o} kg N2O/TJ"}[gas]
            i_1114 = ("K13g-PPDDSL (IAR diésel de PPDSL en " + F_A2 + ")" if t == "PPDDSL" else "oficial de " + F_A2)
            fuente = (f"IPCC 2006, vol. 2, cap. 2, cuadro 2.2 (industrias de la energía, {nombre}): {f_gas}; PCG AR5 "
                      f"(CH4 28, N2O 265); InputActivityRatio por año ({comb}): {i_1114} en 2011-2014 y "
                      f"K13a-{t} en 2015-2035 ({F_BEN3}; {F_BEN228})")
            nota = (f"EAR = factor x IAR(año) / {divisor} en 2011-2035, con el IAR que queda en el datafile del clon "
                    f"({'K13g' if t == 'PPDDSL' else 'oficial'} en 2011-2014, observado del BEN en 2015-2025 y medio "
                    "observado 2023-2025 desde 2026). El oficial conserva el factor de 2010-2018 con PCG AR4 (77,65 t/TJ "
                    "en fuel oil) y no lo actualiza cuando cambia el IAR (2011, 2017 y desde 2019)")
            if len(fila):
                vals = [e10] + [valores[gas][y] for y in yrs]
                F.add(codigo, "A2_quita_SmartGrid", EAR, texto_sg, ANIOS, vals, "Mt/PJ", fuente,
                      nota + ". La fila de SmartGrid del clon se retira porque B1 redondea esas filas a cuatro "
                      f"decimales (CH4 y N2O quedan en 0); 2010 lleva el valor oficial del datafile ({e10:g})", dec="sig")
            else:
                F.add(codigo, "A2", EAR, texto_sg, yrs, [valores[gas][y] for y in yrs], "Mt/PJ", fuente, nota,
                      dec="sig")
    der["K08"] = comprobacion

    # ---------------- K13 filas: rendimiento, actividad y CF de las térmicas fósiles; combustible de PPILPG ------------
    f_k13 = f"{F_BEN3}; {F_BEN228}; {F_BEN} (eólica, solar, biogás, biomasa y térmica); {F_BEP}"
    cf_k13 = {}
    for t in TERMICAS_K08:
        comb = COMB_K13[t]
        grupo = {"PPCRU": "crudo de centrales (A) y de autoproductores (B), cada uno con el rendimiento de su grupo",
                 "PPNGS": "gas natural de centrales (A)", "PPRNGS": "gas natural de autoproductores (B)",
                 "PPRDSL": "diésel de autoproductores (B)",
                 "PPDSL": "diésel de centrales (A) menos el de la generación distribuida (PPDDSL)",
                 "PPDDSL": "generación distribuida con diésel de Energy.xlsx con el rendimiento de A",
                 "PPDFOI": "generación distribuida con fuel oil de Energy.xlsx con el rendimiento de A",
                 "PPFOI": "fuel oil y crudo reducido de centrales (A, menos PPDFOI) y de autoproductores (B), cada uno "
                          "con el rendimiento de su grupo"}[t]
        fuente_t = f_k13 + (f"; {F_GENX}" if t in ("PPDDSL", "PPDFOI", "PPDSL", "PPFOI") else "")
        F.add(f"K13a-{t}", "A2", "InputActivityRatio", f"{t} ; {comb}", HIST + PROY,
              [iar13[t][y] for y in HIST + PROY], "PJ/PJ", fuente_t,
              f"IAR = combustible / generación de {t} ({grupo}); IAR de grupo = combustible fósil / generación fósil del "
              "grupo (centrales: 2,65-2,87; autoproductores: 2,87-3,05; rendimiento 33 % a 38 %). Supuesto: un solo "
              "rendimiento por grupo y año para todos los combustibles. Desde 2026 el medio observado 2023-2025 "
              f"(combustible / generación de esos tres años): {iar13[t][2026]} frente a {iar_of[t][2026]} del oficial. "
              + ("En PPCRU, desde 2026, el de la generación que ancla K14 (media 2023-2025 del crudo de autoproductores "
                 "y crudo de centrales eléctricas de 2025). " if t == "PPCRU" else "")
              + ("2011-2014: K13g" if t == "PPDDSL" else "2011-2014 conserva el oficial (el BEN 2025 empieza en 2015)"))
        if t == "PPDDSL":
            F.add("K13g-PPDDSL", "A2", "InputActivityRatio", "PPDDSL ; E1DSL", list(K13G), [iar_k13g[y] for y in K13G],
                  "PJ/PJ", f"{F_A2}/InputActivityRatio.csv: PPDSL ; E1DSL, 2011-2014 (IAR diésel del propio modelo)",
                  "PPDDSL (generación distribuida con diésel) lleva en el oficial de 2011-2014 el IAR del gas natural "
                  "(4,3645 y 4,0603, los de PPNGS); toma el IAR diésel del propio modelo, el de PPDSL (2,781 en 2011-2012 "
                  "y 2,5349 en 2013-2014), porque las fuentes locales no traen el combustible para generación de esos "
                  "años (el BEN 2025 empieza en 2015). K08-PPDDSL usa este IAR en 2011-2014. Decisión del "
                  "12-10-2026")
        g = gen13[t]
        F.add(f"K13b-{t}", "A2", LL, t, HIST, [0.98 * g[y] for y in HIST], "PJ", fuente_t,
              "Piso = 0,98 x generación observada (combustible del BEN / IAR K13a); reemplaza a K09a-K09b (PPFOI como "
              "residuo de la térmica)")
        F.add(f"K13c-{t}", "A2", UL, t, HIST, [1.02 * g[y] for y in HIST], "PJ", fuente_t,
              "Techo = 1,02 x generación observada; desde 2026 " +
              ("la ancla K14 a la media observada 2023-2025" if t in AUTOPRODUCTORES_K14 else
               "conserva la proyección oficial"))
        cf_of, rc = vector("CapacityFactor", t), rc_k15(t)
        af, cau = vector("AvailabilityFactor", t), vector("CapacityToActivityUnit", t)[None]
        for p_ in ("CapacityFactor", "ResidualCapacity", "AvailabilityFactor"):
            if vector(p_, t, carpeta=ENERGIA / "A2_Output_Params" / "NDCINC") != vector(p_, t):
                raise SystemExit(f"K13: {p_} de {t} difiere entre NDCCON y NDCINC")
        nec = {y: 1.02 * g[y] / (rc[y] * af[y] * cau) for y in HIST}
        yc = [y for y in HIST if nec[y] > cf_of[y] + 1e-9]
        cf_k13[t] = {y: techo4(nec[y]) for y in yc}
        if any(v >= 1 for v in cf_k13[t].values()):
            raise SystemExit(f"K13: {t} necesita CF >= 1 con su capacidad residual: {cf_k13[t]}")
        if yc:
            F.add(f"K13d-{t}", "A2", "CapacityFactor", t, yc, [cf_k13[t][y] for y in yc], "fracción",
                  f"Regla del modelo CF = techo K13c / (capacidad residual x AF x CapacityToActivityUnit); {F_A2} "
                  f"(ResidualCapacity, AF {af[2015]:g}, CAU {cau:g})" +
                  ("; residual K15c en 2023-2025" if t in rc15 else ""),
                  "Solo en los años en que el CF oficial (generación oficial / capacidad, un artefacto de calibración) no "
                  "alcanza el techo K13c con la capacidad residual; así el modelo no construye potencia térmica ficticia "
                  "en 2015-2025. Redondeo hacia arriba en el cuarto decimal; nunca CF = 1. Desde 2026 el CF oficial")
    for cod, par, fac in (("K13e-PPILPG-piso", LL, 0.98), ("K13f-PPILPG-techo", UL, 1.02)):
        F.add(cod, "A2", par, "PPILPG", HIST, [fac * glp13[y] for y in HIST], "PJ",
              f"{F_BEN3} (gas licuado de autoproductores, tabla 3.14); {F_BEP}",
              f"{fac} x GLP para generación del BEN (PJ de combustible). PPILPG no tiene IAR ni OutputActivityRatio y su "
              "EAR de CO2e_PP (0,0632665 Mt/PJ, el factor del GLP por unidad de combustible) trata la actividad como "
              "combustible; el oficial usa 1,20-1,30 PJ y el BEN da 0,30-0,77 PJ. Desde 2026 la ancla K14 a la media "
              "observada 2023-2025")
    der["K13"]["CF_K13d"] = cf_k13

    # ---------------- K14 actividad 2026-2035 de los autoproductores anclada a lo observado (decisión a) ----------------
    # La proyección CLG repite desde 2026 el valor CLG de 2022, que frente a la media observada 2023-2025 (K13) deja a
    # PPCRU en 0,72 veces, a PPRNGS en 0,38, a PPRDSL en 1,37 y a PPILPG en 3,61. Piso 0,98 y techo 1,02 de esa media,
    # la misma banda de K13b-c; CF solo donde el CLG no alcanza el techo con la capacidad residual.
    U3 = [2023, 2024, 2025]
    der_k14 = {}
    # K14 revisado (arreglo de la verificación, 12-10-2026, noche): del crudo de centrales eléctricas (BEN 3.12-3.13)
    # solo se ancla el nivel de 2025, el único año en que CENACE confirma que el S.N.I. no quemó crudo (Informe Anual 2025,
    # Parte 1, figuras 1.7 y 1.10); el exceso de 2023-2024 no tiene central atribuible y queda al despacho. Las centrales
    # de crudo de la capa (nombre con «crudo»; la capa no trae combustible) se contrastan con esa generación
    crudo = ext[(ext.TipoCentra == "Térmica") & ext.NombreCent.str.contains("crudo", case=False)]
    if crudo.empty or set(crudo.Sistema) != {"No Incorporado"}:
        raise SystemExit(f"K14: las centrales de crudo de la capa no son todas de sistemas no incorporados: "
                         f"{crudo[['NombreCent', 'Sistema']].values.tolist()}")
    nom_c, ef_c = crudo.PotenciaNo.astype(float).sum(), crudo.PotenciaEf.astype(float).sum()
    g_a = der13["generacion_PPCRU_por_grupo_PJ"]["centrales"]
    g_b = der13["generacion_PPCRU_por_grupo_PJ"]["autoproductores"]
    max_nom, max_ef = nom_c / 1000 * 31.536, ef_c / 1000 * 31.536
    a_ancla = g_a[2025]
    m_cru = sum(g_b[y] for y in U3) / 3 + a_ancla
    der_k14["PPCRU_crudo_de_centrales"] = {
        "centrales_capa": [{"central": r.NombreCent, "empresa": r.NombreEmpr, "tipo_empresa": r.TipoEmpres,
                            "sistema": r.Sistema, "inicio": f"{r.FechaInici[6:]}-{r.FechaInici[4:6]}-{r.FechaInici[:4]}",
                            "nominal_MW": float(r.PotenciaNo), "efectiva_MW": float(r.PotenciaEf)}
                           for r in crudo.itertuples()],
        "generacion_grupo_A_PJ": {y: g_a[y] for y in U3}, "generacion_grupo_B_PJ": {y: g_b[y] for y in U3},
        "media_A_PJ": round(sum(g_a[y] for y in U3) / 3, 5), "media_B_PJ": round(sum(g_b[y] for y in U3) / 3, 5),
        "A_anclada_PJ": round(a_ancla, 5), "ancla_PPCRU_PJ": round(m_cru, 5),
        "A_al_despacho_respecto_de_la_media_PJ": round(sum(g_a[y] for y in U3) / 3 - a_ancla, 5),
        "CF_de_la_A_de_cada_anio_sobre_la_potencia_efectiva_de_la_capa": {y: round(g_a[y] / max_ef, 3) for y in U3},
        "generacion_maxima_capa_nominal_PJ": round(max_nom, 4), "generacion_maxima_capa_efectiva_PJ": round(max_ef, 4),
        "anios_A_sobre_maximo_efectivo": [y for y in U3 if g_a[y] > max_ef],
        "lectura": "Del crudo de centrales eléctricas del BEN (3.12-3.13; ARCONEL 2025, tabla 41, crudo de generadoras) "
                   "solo se ancla el de 2025 (0,6035 PJ de generación): ese año CENACE no registra crudo en el S.N.I. "
                   "(Informe Anual 2025, Parte 1, figuras 1.7 y 1.10, p. 34 y 40), así que se quemó fuera del S.N.I. La "
                   "media 2023-2025 (1,3169 PJ) equivale a CF 1,00 sobre la potencia efectiva de las dos centrales de "
                   "crudo de la capa (CPT Crudo Diesel CELEC y Cuyabeno Crudo Diesel CELEC, 41,74 MW), y 2024 a 1,55: el "
                   "exceso de 2023-2024 sobre 2025 no tiene central atribuible y queda al despacho de la red. La capa "
                   "rotula esas dos centrales como autogeneradoras de Petroecuador operadas por CELEC; ARCONEL contaría "
                   "su crudo como de generadoras si clasifica por el operador [inferencia; verificar con ARCONEL o CELEC "
                   "Termopichincha qué centrales quemaron crudo en 2023-2024]. Quedan sin examinar con datos de "
                   "combustible el Sistema Oriental de CELEC Termopichincha (Secoya, 10 MW, No Incorporado) y las "
                   "térmicas de Termopichincha del S.N.I. en Sucumbíos y Orellana (Jivino I a III, Celso Castellanos, "
                   "Payamino y Dayuma): ninguna fuente local trae su combustible"}
    f_crudo = ("ARCONEL, capa CENTRALES; uso autorizado; referencia del oficio [verificar]: CPT Crudo Diesel CELEC y "
               "Cuyabeno Crudo Diesel CELEC, sistema No Incorporado; CENACE, Informe Anual 2025, Parte 1, figuras 1.7 y "
               "1.10, p. 34 y 40 (sin crudo en el S.N.I. en 2025); ARCONEL, Estadística Anual y Multianual 2025, tabla 41, "
               "p. 44 (crudo de generadoras)")
    for t in AUTOPRODUCTORES_K14:
        m = m_cru if t == "PPCRU" else sum(gen13[t][y] for y in U3) / 3
        ll_of, ul_of = vector(LL, t), vector(UL, t)
        cf_of, rc = vector("CapacityFactor", t), rc_k15(t)
        af, cau = vector("AvailabilityFactor", t), vector("CapacityToActivityUnit", t)[None]
        for p_ in (LL, UL, "CapacityFactor", "ResidualCapacity", "AvailabilityFactor"):
            if {y: v for y, v in vector(p_, t, carpeta=ENERGIA / "A2_Output_Params" / "NDCINC").items() if y in PROY} != \
                    {y: v for y, v in vector(p_, t).items() if y in PROY}:
                raise SystemExit(f"K14: {p_} de {t} 2026-2035 difiere entre NDCCON y NDCINC")
        nec = {y: 1.02 * m / (rc[y] * af[y] * cau) for y in PROY}
        yc = [y for y in PROY if nec[y] > cf_of[y] + 1e-9]
        cf14 = {y: techo4(nec[y]) for y in yc}
        if any(v >= 1 for v in cf14.values()):
            raise SystemExit(f"K14: {t} necesita CF >= 1 con su capacidad residual: {cf14}")
        der_k14[t] = {"generacion_observada_PJ": {y: round(gen13[t][y], 5) for y in U3}, "media_2023_2025_PJ": round(m, 5),
                      "media_GWh": round(m * PJ, 1), "piso_CLG_2026_PJ": ll_of[2026], "techo_CLG_2026_PJ": ul_of[2026],
                      "CLG_2026_sobre_media": round(ul_of[2026] / m, 3), "CF_K14c": cf14,
                      "capacidad_residual_GW": rc[2026], "CF_CLG": cf_of[2026]}
        grupo = {"PPCRU": "media 2023-2025 del crudo de autoproductores (B, "
                          f"{der_k14['PPCRU_crudo_de_centrales']['media_B_PJ']:.4f} PJ) más el crudo de centrales "
                          f"eléctricas (A) de 2025 ({a_ancla:.4f} PJ), el único año en que CENACE confirma que el S.N.I. "
                          "no quemó crudo; el exceso de A en 2023-2024 no tiene central atribuible y queda al despacho",
                 "PPRNGS": "gas natural de autoproductores (B)", "PPRDSL": "diésel de autoproductores (B)"}[t]
        nota_m_obs = ("" if t != "PPCRU" else
                      f"; la media observada de A y B (sin el ajuste) era {sum(gen13[t][y] for y in U3) / 3:.4f} PJ")
        f_k14 = (f"{F_BEN3} (tabla 3.14 combustible y 3.15 electricidad de autoproductores; 3.12-3.13 para el crudo de "
                 f"centrales); {F_BEN} (biomasa e hidro de autoproductores que se restan); {F_BEN228}; {F_BEP}" +
                 (f"; {f_crudo}" if t == "PPCRU" else ""))
        nota_m = (f"{'ancla' if t == 'PPCRU' else 'media observada 2023-2025'} de {t} ({grupo}) = {m:.4f} PJ "
                  f"({m * PJ:.1f} GWh; K13: combustible del BEN / IAR del grupo{nota_m_obs}); el CLG repite desde 2026 "
                  f"su valor de 2022 ({ul_of[2026]:g} PJ de techo, {ul_of[2026] / m:.2f} veces la ancla). Decisión del "
                  "12-10-2026" + (" y arreglo de la verificación del 12-10-2026 (crudo de centrales)"
                                              if t == "PPCRU" else ""))
        F.add(f"K14a-{t}", "A2", LL, t, PROY, [0.98 * m] * len(PROY), "PJ", f_k14, "Piso 2026-2035 = 0,98 x " + nota_m)
        F.add(f"K14b-{t}", "A2", UL, t, PROY, [1.02 * m] * len(PROY), "PJ", f_k14, "Techo 2026-2035 = 1,02 x " + nota_m)
        if yc:
            F.add(f"K14c-{t}", "A2", "CapacityFactor", t, yc, [cf14[y] for y in yc], "fracción",
                  f"Regla del modelo CF = techo K14b / (capacidad residual x AF x CapacityToActivityUnit); {F_A2} "
                  f"(ResidualCapacity {rc[2026]:g} GW{' de K15c' if t in rc15 else ''}, AF {af[2026]:g}, "
                  f"CAU {cau:g})",
                  f"El CF oficial desde 2026 ({cf_of[2026]:g}, generación oficial / capacidad) no alcanza el techo K14b con "
                  "la capacidad residual y obligaría a construir potencia ficticia; la capacidad existente opera con el "
                  "CF que pide la actividad observada. Redondeo hacia arriba en el cuarto decimal; nunca CF = 1")
    m_glp = sum(glp13[y] for y in U3) / 3
    ul_glp, rc_glp = vector(UL, "PPILPG"), vector("ResidualCapacity", "PPILPG")
    lim_glp = rc_glp[2026] * vector("CapacityFactor", "PPILPG")[2026] * vector("AvailabilityFactor", "PPILPG")[2026] * \
        vector("CapacityToActivityUnit", "PPILPG")[None]
    if lim_glp < 1.02 * m_glp:
        raise SystemExit("K14: la capacidad de PPILPG no alcanza el techo")
    der_k14["PPILPG"] = {"GLP_observado_PJ": {y: round(glp13[y], 5) for y in U3}, "media_2023_2025_PJ": round(m_glp, 5),
                         "techo_CLG_2026_PJ": ul_glp[2026], "CLG_2026_sobre_media": round(ul_glp[2026] / m_glp, 3),
                         "limite_capacidad_PJ": round(lim_glp, 4)}
    nota_glp = (f"media observada 2023-2025 del GLP para generación de autoproductores = {m_glp:.4f} PJ de combustible "
                "(BEN tabla 3.14). PPILPG no tiene InputActivityRatio ni OutputActivityRatio: su actividad es el "
                "combustible (EAR de CO2e_PP 0,0632665 Mt/PJ, el factor del GLP), no abastece la demanda eléctrica y su "
                f"generación (unos {sum(der13['generacion_GLP_GWh'][y] for y in U3) / 3:.0f} GWh/año con el IAR de los "
                "autoproductores) no "
                f"entra al balance. El CLG usa {ul_glp[2026]:g} PJ de techo desde 2026 ({ul_glp[2026] / m_glp:.2f} veces la "
                "media). Decisión del 12-10-2026")
    for cod, par, fac in (("K14d-PPILPG-piso", LL, 0.98), ("K14e-PPILPG-techo", UL, 1.02)):
        F.add(cod, "A2", par, "PPILPG", PROY, [fac * m_glp] * len(PROY), "PJ",
              f"{F_BEN3} (gas licuado de autoproductores, tabla 3.14); {F_BEP}",
              f"{'Piso' if fac < 1 else 'Techo'} 2026-2035 = {fac} x " + nota_glp)
    der["K14"] = der_k14

    # ---------------- K15c filas: potencia residual de las térmicas 2023-2035 ------------------------------------------
    nombres_altas = {t: [f"{a['central']} ({a['MW_nominal']:g} MW, {a['anio']}"
                         f"{', arrendada' if a['modalidad'] == 'arrendamiento' else ''}"
                         f"{', omitida en Energy.xlsx' if a['clase_registro'] != 'alta 2024-2025' else ''})"
                         for a in der15["altas"] if a["tecnologia"] == t and a["incluida"] == "sí"] for t in K15C_MAPA}
    arr_t = {t: [a for a in der15["arriendos"] if a["tecnologia"] == t] for t in K15C_MAPA}
    for t, s in rc15.items():
        c = next(f for f in conc15 if f["tecnologia"] == t)
        altas_t = "; ".join(nombres_altas[t]) or "ninguna"
        nota_arr = ""
        if arr_t[t]:
            nota_arr = (f" Desde 2026 la potencia arrendada ({len(arr_t[t])} contratos: "
                        + "; ".join(f"{a['central']} {a['MW']:g} MW, inicio {a['inicio']}, fin "
                                    f"{a['fin_publicado'] or 'sin fecha publicada'}" for a in arr_t[t])
                        + ") cuenta como MW x fracción del año en que el contrato está vigente; sin fecha de fin publicada "
                          f"(S26-116), caso central con {F_FIN_ARRIENDO[VARIANTE_CENTRAL]}; la variante con "
                          f"{F_FIN_ARRIENDO['2032']} la resuelve variante_suficiencia.py.")
        F.add(f"K15c-{t}", "A2", "ResidualCapacity", t, K15C_ANIOS, [s[y] for y in K15C_ANIOS], "GW",
              f"{F_ENERGY_CAP}: {c['fila_Energy_xlsx']} = {c['Energy_xlsx_2023_MW']:g} MW en 2023; {ALTAS.name} ({F_CAPA}; "
              f"combustible de CENACE, Informe Anual 2025, Parte 2, tablas 3.1 y 3.2, p. 116 y 120)"
              + (f"; {EXPANSION.name} (arriendos de 2026 del plan de 920 MW: CELEC EP, boletines 042, 109 y 122, S26-079 "
                 f"a S26-081); vigencia: {F_PME49} y S26-116" if arr_t[t] else "")
              + f"; contraste con {F_ARC25} (arconel2025_potencia_termica.csv)",
              f"Potencia del modelo en 2025 = residual CLG {c['CLG_2025_MW']:g} MW + "
              f"{c['construida_por_el_modelo_2010_2022_MW']:g} MW que el LP construye en 2010-2022 y siguen en operación "
              f"= {c['potencia_total_antes_2025_MW']:g} MW, frente a {c['fuente_2025_MW']:g} MW de la fuente "
              f"({c['desvio_CLG_pct']:+.1f} %): se corrige 2023-2035. Potencia de la fuente: 2023 = Energy.xlsx (la fuente "
              f"del CLG, que A2 no siguió) más las centrales que la capa registra y Energy.xlsx omite; 2024-2025 suman "
              f"las altas con fecha de la capa ({altas_t}), con la potencia al 31 de diciembre; desde 2026 el parque "
              f"propio de 2025.{nota_arr} Residual = potencia de la fuente menos la construcción del modelo en 2010-2022, "
              "de modo que la potencia total del clon en 2023-2025 es la de la fuente. Potencia nominal, como la de "
              "Energy.xlsx. 2015-2022 conservan la residual CLG. Decisión K15c del 12-10-2026", dec=6)

    # ---------------- K15a CF 2026-2035 continuo con lo observado (térmicas de la red) --------------------------------
    der_k15a = {}
    for t in ("PPFOI", "PPDSL", "PPNGS", "PPDFOI", "PPDDSL"):
        cf_of = vector("CapacityFactor", t)
        cf_obs = {y: cf_k13[t].get(y, cf_of[y]) for y in U3}
        m = sum(cf_obs.values()) / 3
        razon = cf_of[2026] / m
        aplica = razon < 0.95
        der_k15a[t] = {"CF_2023_2025_datafile": cf_obs, "media": round(m, 5), "CF_CLG_2026": cf_of[2026],
                       "CLG_2026_sobre_media": round(razon, 4), "aplica": aplica,
                       "potencia_2025_GW": rc_k15(t)[2025]}
        if aplica:
            v = techo4(m)
            F.add(f"K15a-{t}", "A2", "CapacityFactor", t, PROY, [v] * len(PROY), "fracción",
                  f"CF del datafile del clon en 2023-2025 (K13d: {', '.join(f'{y} {cf_obs[y]:g}' for y in U3)}), con la "
                  f"residual {'K15c' if t in rc15 else 'CLG'}; {F_BEN3}; {F_BEN228}",
                  f"CF 2026-2035 = medio observado 2023-2025 ({v:g}); el CF CLG de 2026 ({cf_of[2026]:g}) es "
                  f"{razon:.2f} veces ese medio, una discontinuidad que obligaba a construir potencia nueva para la misma "
                  "actividad. Decisión K15a del 12-10-2026 (se citó 0,5707 en PPFOI y 0,3040 en "
                  "PPDSL con la residual CLG; con la residual K15c la misma regla da este valor). Nunca CF = 1")
    der["K15a"] = der_k15a

    # ---------------- K15b inversión térmica nueva limitada a la expansión declarada ----------------------------------
    exp_ = pd.read_csv(EXPANSION, dtype={"id": str})
    if exp_.id.duplicated().any():
        raise SystemExit("K15b: identificadores repetidos en expansion_termica_declarada.csv")
    malos = exp_[exp_.tratamiento.isin(["suma", "arriendo"]) & ~exp_.tecnologia_modelo.isin(RED_K15)]
    if len(malos) or ((exp_.tratamiento != "suma") & (exp_.MW_en_tope != 0)).any() or \
            ((exp_.tratamiento != "arriendo") & (exp_.MW_residual_arriendo != 0)).any():
        raise SystemExit("K15b: filas de expansion_termica_declarada.csv con tecnología o tope inconsistentes")
    der_k15b = {"declarada_por_anio_MW": {}, "tope_GW": {}, "potencia_existente_GW": {},
                "variante_arriendos_2032": {}, "regla": "tope(año) = residual K15c(año) + construcción del modelo en "
                                                       "2010-2022 + expansión propia declarada acumulada"}
    rc_var = der15["series_GW_por_vigencia_de_arriendos"]
    for t in RED_K15:
        d = exp_[(exp_.tratamiento == "suma") & (exp_.tecnologia_modelo == t)]
        por_anio = {y: float(d[d.anio == y].MW_en_tope.sum()) for y in PROY}
        if (d.anio < 2026).any():
            raise SystemExit(f"K15b: {t} con expansión sumada antes de 2026")
        exist = rc_k15(t)
        acum = {y: sum(por_anio[k] for k in PROY if k <= y) / 1000 for y in PROY}
        tope = {y: exist[y] + acum[y] for y in PROY}
        der_k15b["declarada_por_anio_MW"][t] = por_anio
        der_k15b["tope_GW"][t] = {y: round(v, 6) for y, v in tope.items()}
        der_k15b["potencia_existente_GW"][t] = {y: round(exist[y], 6) for y in PROY}
        if t in rc_var:  # variante de arriendos hasta el 31-03-2032: residual y tope que variante_suficiencia.py escribe
            rv = rc_var[t]["2032"]
            der_k15b["variante_arriendos_2032"][t] = {
                "ResidualCapacity": {y: rv[y] for y in PROY},
                "TotalAnnualMaxCapacity": {y: round(rv[y] + hist15[t] + acum[y], 6) for y in PROY}}
        ids = ", ".join(f"{r.id} {r.proyecto} {r.MW_en_tope:g} MW ({r.anio})" for r in d.itertuples()) or "ninguna"
        F.add(f"K15b-{t}", "A2_nueva", "TotalAnnualMaxCapacity", t, ANIOS,
              [99999] * 16 + [tope[y] for y in PROY], "GW",
              f"{EXPANSION.name}: MAE, Anexos del ajuste PME 2023-2032, Anexo A, p. 1-2 (Acuerdo MAE-MAE-2026-0066-AM, RO 319, "
              f"S26-109); CELEC EP, Boletín 042 del 20-04-2026 (S26-079); PME 2023-2032, cap. 4, p. 129 (combustible del "
              f"ciclo combinado); potencia existente de cada año = residual K15c (con los arriendos vigentes) más lo "
              f"construido por el modelo en 2010-2022 ({hist15[t] * 1000:g} MW)",
              f"Tope de potencia total = potencia existente del año + expansión propia declarada acumulada (filas: {ids}); "
              "99999 (valor por omisión) en 2010-2025. Los arriendos no entran como inversión: están en la residual K15c "
              "solo mientras su contrato está vigente, así que el tope baja cuando vencen y la inversión permitida es "
              "solo la propia declarada. B1 escribe TotalAnnualMaxCapacityInvestment como bloque vacío fijo y glpsol "
              "rechaza un segundo bloque, así que el tope anual de inversión no entra sin cambiar B1. Sin expansión "
              f"declarada desde 2033. Topes (MW): {', '.join(f'{y} {tope[y] * 1000:.1f}' for y in PROY)}. Decisión K15b "
              "del 12-10-2026 y arreglo de la verificación (arriendos con vigencia)", dec=6)
    der["K15b"] = der_k15b

    # ---------------- K15d y K15e techo y CF de PPNGS con la oferta observada y la energía declarada (solo central) ----
    # El techo CLG de PPNGS desde 2026 (2,9497 PJ) es la generación del S.N.I. a gas de 2022 del CLG congelada (Energy.xlsx,
    # Generación_PJ), sin fuente de suministro: deja sin despacho la expansión a gas declarada. Se reemplaza por 1,02 x la
    # media observada 2023-2025 (la banda de K13c y K14b) más la energía media declarada de cada proyecto a gas desde su año
    # de entrada (Anexo A). El CF es el que deja a la potencia declarada producir ese techo (regla del modelo CF = techo /
    # (tope K15b x AF x CAU)), solo donde el CLG no lo alcanza. G01 sigue mandando en INCA6G y CONA6G.
    gas = exp_[(exp_.tratamiento == "suma") & (exp_.tecnologia_modelo == "PPNGS")]
    if gas.energia_media_GWh.isna().any():
        raise SystemExit("K15d: falta la energía media declarada de un proyecto a gas")
    m_ngs = sum(gen13["PPNGS"][y] for y in U3) / 3
    e_decl = {y: float(gas[gas.anio <= y].energia_media_GWh.sum()) / PJ for y in PROY}
    techo_ngs = {y: 1.02 * m_ngs + e_decl[y] for y in PROY}
    ul_ngs, ll_ngs = vector(UL, "PPNGS"), vector(LL, "PPNGS")
    cf_ngs, af_ngs = vector("CapacityFactor", "PPNGS"), vector("AvailabilityFactor", "PPNGS")
    cau_ngs = vector("CapacityToActivityUnit", "PPNGS")[None]
    tope_ngs = der_k15b["tope_GW"]["PPNGS"]
    nec_ngs = {y: techo_ngs[y] / (tope_ngs[y] * af_ngs[y] * cau_ngs) for y in PROY}
    yc_ngs = [y for y in PROY if nec_ngs[y] > cf_ngs[y] + 1e-9]
    cf15e = {y: techo4(nec_ngs[y]) for y in yc_ngs}
    if any(v >= 1 for v in cf15e.values()):
        raise SystemExit(f"K15e: PPNGS necesita CF >= 1: {cf15e}")
    if any(techo_ngs[y] < ll_ngs[y] for y in PROY):
        raise SystemExit("K15d: el techo de PPNGS queda bajo su piso CLG")
    ids_gas = "; ".join(f"{r.id} {r.proyecto} {r.MW_en_tope:g} MW, {r.energia_media_GWh:g} GWh/año ({r.anio})"
                        for r in gas.itertuples())
    der["K15d"] = {"media_observada_2023_2025_PJ": round(m_ngs, 5), "energia_declarada_acumulada_PJ":
                   {y: round(v, 5) for y, v in e_decl.items()}, "techo_PJ": {y: round(v, 4) for y, v in techo_ngs.items()},
                   "techo_CLG_PJ": {y: ul_ngs[y] for y in PROY}, "piso_CLG_PJ": {y: ll_ngs[y] for y in PROY},
                   "CF_CLG": {y: cf_ngs[y] for y in PROY}, "CF_K15e": cf15e,
                   "IAR_PPNGS_desde_2026": iar13["PPNGS"][2026],
                   "sobrestimacion_maxima_CO2e_ciclo_combinado_Machala_kt_anio": round(
                       690 / PJ * iar13["PPNGS"][2026] * (56.1 + 28 * 1.0 / 1000 + 265 * 0.1 / 1000), 1)}
    F.add("K15d-PPNGS", "A2", UL, "PPNGS", PROY, [techo_ngs[y] for y in PROY], "PJ",
          f"{F_BEN3}; {F_BEN228} (gas natural de centrales, K13); {F_ANEXO_A}; MEM, PME 2023-2032, cap. 4, p. 129, "
          f"hipótesis n (el ciclo combinado de 400 MW tendría su propio suministro de gas natural); {EXPANSION.name}",
          f"Techo 2026-2035 = 1,02 x media observada 2023-2025 de PPNGS ({m_ngs:.4f} PJ, BEN; la banda de K13c y K14b) + "
          f"energía media declarada acumulada de la expansión a gas desde su año de entrada ({ids_gas}; año completo "
          "desde el año de entrada [supuesto: el Anexo A no da el mes]). Reemplaza el techo CLG de "
          f"{ul_ngs[2026]:g} PJ, la generación del S.N.I. a gas de 2022 del CLG congelada (Energy.xlsx, Generación_PJ, "
          "filas 18 y 48), que no tiene fuente de suministro y dejaba sin despacho la expansión declarada. El piso CLG "
          f"({ll_ngs[2026]:g} PJ) se conserva. Límite: la turbina de vapor del ciclo combinado de Machala no quema "
          "combustible adicional y el modelo le aplica el IAR medio de PPNGS (sobrestima su CO2). Solo INCA6 y CONA6 "
          "(en INCA6G y CONA6G manda G01). Arreglo de la verificación del 12-10-2026", esc=ESC_CEN)
    if yc_ngs:
        F.add("K15e-PPNGS", "A2", "CapacityFactor", "PPNGS", yc_ngs, [cf15e[y] for y in yc_ngs], "fracción",
              f"Regla del modelo CF = techo K15d / (tope K15b x AF x CapacityToActivityUnit); {F_A2} (AF "
              f"{af_ngs[2026]:g}, CAU {cau_ngs:g}); {F_ANEXO_A}",
              f"El CF CLG de PPNGS ({cf_ngs[2026]:g}) es el del parque existente limitado por el gas de Machala; con él la "
              "potencia declarada (Machala tercera unidad, ciclo combinado de Machala y Bloque de Ciclo Combinado I) no "
              "puede producir la energía media que el Anexo A le asigna. Topes (MW): "
              f"{', '.join(f'{y} {tope_ngs[y] * 1000:.1f}' for y in yc_ngs)}. Redondeo hacia arriba en el cuarto decimal; "
              "nunca CF = 1. Solo INCA6 y CONA6. Arreglo de la verificación del 12-10-2026", esc=ESC_CEN)

    # ---------------- K09 historia 2015-2025 anclada al BEN ------------------------------------------------------------
    # (a) K09a y K09b (PPFOI = BEN térmica - pisos del resto) quedan reemplazadas por K13
    # (b) importaciones 2023-2025
    imp_ll, imp_ul = vector(LL, "IMPELE"), vector(UL, "IMPELE")
    razon_i = imp_ul[2022] / imp_ll[2022]
    yrs = [2023, 2024, 2025]
    F.add("K09c", "A2", LL, "IMPELE", yrs, [ben.loc[y, "importaciones"] / PJ for y in yrs], "PJ",
          f"{F_BEN} (interconexión/importaciones)", "Piso = importaciones del BEN; 2015-2022 el oficial ya coincide")
    F.add("K09d", "A2", UL, "IMPELE", yrs, [ben.loc[y, "importaciones"] / PJ * razon_i for y in yrs], "PJ",
          f"{F_BEN} (interconexión/importaciones)", f"Techo = {razon_i:.4f} x piso (razón oficial de IMPELE)")
    # (c) hidro: techo = BEN x 1,05 en proporción a los techos con K01 y K02, acotado por capacidad x CF x AF x CAU
    hidro = sorted(t for t in leer_a2(UL).TECHNOLOGY.unique() if t.startswith("PPH")) + ["PPIHD"]
    base_ul = {t: (ul[t] if t in ul else vector(UL, t)) for t in hidro}
    base_ll = {t: (ll[t] if t in ll else vector(LL, t)) for t in hidro}
    k = {y: ben.loc[y, "hidraulica"] * 1.05 / (sum(base_ul[t].get(y, 0.0) for t in hidro) * PJ) for y in HIST}

    def capacidad(t, y):
        """Potencia de la clase: la del modelo oficial (residual + la que construye), salvo PPHD (K07: residual 0,813
        GW, sin los 19,6 GW ficticios) y PPHRORAB (K03: 0 en 2015 y 1,987 GW desde 2016)."""
        if t == "PPHD":
            return vector("ResidualCapacity", t)[y]
        if t == "PPHRORAB":
            return 0.0 if y < 2016 else cap_k
        c = [of[e].cap.get((t, y), 0.0) for e in ("NDCCON", "NDCINC")]
        if abs(c[0] - c[1]) > 1e-6:
            raise SystemExit(f"{t} {y}: la potencia hidro difiere entre NDCCON y NDCINC")
        return c[0]

    cfh, ulb, llb, cf_nuevo, lim = {}, {}, {}, {}, {}
    p_ben = pot.hidraulica
    for t in hidro:
        cf_t, af_t = vector("CapacityFactor", t), vector("AvailabilityFactor", t)
        cau_t = 31.536 if t == "PPHD" else vector("CapacityToActivityUnit", t)[None]  # PPHD: K07a
        cfh[t], ulb[t], lim[t] = {}, {}, {}
        for y in HIST:
            b = base_ul[t].get(y, 0.0)
            if b <= 0:
                continue
            c = capacidad(t, y)
            if y >= 2023:  # CF del año con la generación observada repartida en proporción al techo de la clase
                cuota = b / sum(base_ul[s].get(y, 0.0) for s in hidro)
                cfn = techo4(ben.loc[y, "hidraulica"] / PJ * cuota / (c * af_t[y] * cau_t))
                if cfn >= 1:
                    raise SystemExit(f"K09cf: {t} {y} CF {cfn} >= 1")
                cf_nuevo.setdefault(t, {})[y] = cfn
                cfh[t][y] = cfn
            else:
                cfh[t][y] = cf_t[y]
            lim[t][y] = c * cfh[t][y] * af_t[y] * cau_t
            ulb[t][y] = min(r4(b * k[y]), piso4(lim[t][y]))
    k_piso = {y: min(1.0, ben.loc[y, "hidraulica"] / PJ / sum(ulb[t].get(y, 0.0) for t in hidro)) for y in HIST}
    for t in hidro:
        for y in ulb[t]:
            llb.setdefault(t, {})[y] = min(base_ll[t].get(y, 0.0), piso4(k_piso[y] * ulb[t][y]))
    der["K09_hidro"] = {
        "suma_techos_K01_K02_GWh": {y: round(sum(base_ul[t].get(y, 0.0) for t in hidro) * PJ, 1) for y in HIST},
        "BEN_hidro_GWh": {y: float(ben.loc[y, "hidraulica"]) for y in HIST}, "k": {y: round(k[y], 5) for y in HIST},
        "suma_techos_acotados_GWh": {y: round(sum(ulb[t].get(y, 0.0) for t in hidro) * PJ, 1) for y in HIST},
        "suma_limites_capacidad_GWh": {y: round(sum(lim[t].get(y, 0.0) for t in hidro) * PJ, 1) for y in HIST},
        "k_piso": {y: round(k_piso[y], 5) for y in HIST},
        "suma_pisos_antes_GWh": {y: round(sum(base_ll[t].get(y, 0.0) for t in hidro if y in ulb[t]) * PJ, 1)
                                 for y in HIST},
        "suma_pisos_despues_GWh": {y: round(sum(llb.get(t, {}).get(y, 0.0) for t in hidro) * PJ, 1) for y in HIST},
        "potencia_modelo_MW": {y: round(sum(capacidad(t, y) for t in hidro) * 1000, 1) for y in HIST},
        "potencia_BEN_2_23_MW": {y: float(p_ben[y]) for y in HIST},
        "CF_parque_BEN": {y: round(ben.loc[y, "hidraulica"] / (p_ben[y] * 8.76 * 0.95), 4) for y in HIST},
        "CF_parque_modelo": {y: round(ben.loc[y, "hidraulica"] / (sum(capacidad(t, y) for t in hidro) * 1000 * 8.76
                                                                    * 0.95), 4) for y in HIST},
        "CF_nuevo_2023_2025": cf_nuevo,
        "acotados_por_capacidad": {t: [y for y in ulb[t] if ulb[t][y] < r4(base_ul[t][y] * k[y])] for t in hidro}}
    for t in hidro:
        yrs = sorted(ulb[t])
        if not yrs:
            continue
        F.add(f"K09h-{t}", "A2", UL, t, yrs, [ulb[t][y] for y in yrs], "PJ",
              f"{F_BEN} (hidráulica); techos de {F_A2} con K01 y K02; potencia del {F_OF}",
              "Techo = min(BEN hidro x 1,05 repartido en proporción a los techos de cada central; potencia x CF x AF x "
              "CAU). La potencia es la del modelo oficial, así el techo nunca obliga a construir capacidad (k y límites "
              "en correcciones_K_insumos.json)")
        if t in cf_nuevo:
            yc = sorted(cf_nuevo[t])
            F.add(f"K09cf-{t}", "A2", "CapacityFactor", t, yc, [cf_nuevo[t][y] for y in yc], "fracción",
                  f"{F_BEN} (hidráulica 2023-2025: 25.349; 22.614; 28.129 GWh); {F_BEN23} (2023-2025: 5.152; 5.371; "
                  f"5.387 MW); potencia de la clase del {F_OF}",
                  "CF = generación observada asignada a la clase en proporción a su techo / (potencia de la clase x 8.760 "
                  "h x AF). La regla con la potencia del BEN (G / (P_BEN x 8.760 x AF) = 0,591; 0,506; 0,627) da CF > 1 "
                  "en PPHRORAS porque la potencia hidro del modelo supera la del BEN en unos 1.050 MW; ver LEEME. Desde "
                  "2026 el CF vuelve al oficial")
        yl = [y for y in yrs if llb[t][y] < base_ll[t].get(y, 0.0) - 1e-9]
        if yl:
            F.add(f"K09l-{t}", "A2", LL, t, yl, [llb[t][y] for y in yl], "PJ", f"{F_BEN} (hidráulica)",
                  "Piso = min(piso oficial con K01 y K02; k_piso x techo K09h), k_piso = BEN hidro / suma de techos "
                  "(nunca > 1): la suma de pisos hidro no supera la hidro del BEN")

    # ---------------- K11 eólica existente -------------------------------------------------------------------------------
    cau_w, af_w = 31.536, vector("AvailabilityFactor", "PPWNDON")[2015]
    if vector("CapacityToActivityUnit", "PPWNDON")[None] != cau_w or \
            vector("CapacityToActivityUnit", "PPWNDD")[None] != cau_w:
        raise SystemExit("K11: CapacityToActivityUnit eólico inesperado")
    deol_x = capx[("Distribuida", "Eólica")][0]
    d_eol_xlsx = {y: 1000 * deol_x.get(y, deol_x[2023]) for y in ANIOS}  # MW; 2024-2035 arrastran 2023
    # Capa CENTRALES (decisión e3): la eólica no incorporada (Galápagos: San Cristóbal desde 2007 y Baltra desde el
    # 22-12-2014) da la distribuida de 2015-2035; 2011-2014 conservan Energy.xlsx porque el CLG pone la generación de San
    # Cristóbal de 2010-2012 en PPWNDON (los 2,29 MW que construye en 2010, año intacto)
    cap_d_capa = capa["eolica_no_incorporada"].potencia_efectiva_MW
    cap_sni_capa = capa["eolica_SNI"].potencia_efectiva_MW
    d_eol = {y: (float(cap_d_capa[y]) if y >= 2015 else d_eol_xlsx[y]) for y in ANIOS}
    b2010 = of["NDCCON"].nueva[("PPWNDON", 2010)]
    if abs(of["NDCINC"].nueva[("PPWNDON", 2010)] - b2010) > 1e-9:
        raise SystemExit("K11: la capacidad eólica de 2010 difiere entre NDCCON y NDCINC")
    g_d = {y: iguales_historia(of, UL, "PPWNDD", [y])[0] for y in HIST}
    # PPWNDON = potencia eólica del BEN 2.23 - distribuida - lo construido en 2010; si el BEN suma potencia en un año en
    # que la capa no tiene ninguna central nueva del SNI en operación, ese año repite el anterior (2022: el BEN agrega 29
    # MW de Huascachaca, que la capa registra con inicio el 31-03-2023, y la eólica del BEN 2.20 de 2022 es la de 2021)
    rc_on, anio_sin_central = {}, []
    for y in HIST:
        if y > HIST[0] and cap_sni_capa[y] == cap_sni_capa[y - 1] and pot.loc[y, "eolica"] > pot.loc[y - 1, "eolica"]:
            rc_on[y] = rc_on[y - 1]
            anio_sin_central.append(y)
        else:
            rc_on[y] = (pot.loc[y, "eolica"] - d_eol[y]) / 1000 - b2010
    huasca = capa_central("Huascachaca", "Eólica")
    g_on = {y: ben.loc[y, "eolica"] / PJ - g_d[y] for y in HIST}
    cf_on_raw = {y: g_on[y] / ((rc_on[y] + b2010) * af_w * cau_w) for y in HIST}
    rc_off = vector("ResidualCapacity", "PPWNDON")
    ul_off_on = {y: iguales_historia(of, UL, "PPWNDON", [y])[0] for y in (2013, 2014)}
    for y in (2013, 2014):
        cf_on_raw[y] = ul_off_on[y] / ((rc_off[y] + b2010) * af_w * cau_w)
    cf_on_med = sum(cf_on_raw[y] for y in (2023, 2024, 2025)) / 3
    g_on_ex = (rc_on[2025] + b2010) * cf_on_med * af_w * cau_w
    cf_d_raw = {y: g_d[y] / (d_eol[y] / 1000 * af_w * cau_w) for y in HIST}
    cf_d_med = sum(cf_d_raw[y] for y in (2023, 2024, 2025)) / 3
    g_d_ex = d_eol[2025] / 1000 * cf_d_med * af_w * cau_w
    for y, v in list(cf_on_raw.items()) + list(cf_d_raw.items()):
        if v >= 1:
            raise SystemExit(f"K11: CF >= 1 en {y}")
    der["K11"] = {"capacidad_2010_construida_GW": b2010, "distribuida_MW": {y: d_eol[y] for y in ANIOS},
                  "distribuida_Energy_xlsx_MW": {y: d_eol_xlsx[y] for y in ANIOS},
                  "capa_eolica_SNI_efectiva_MW": {y: float(cap_sni_capa[y]) for y in ANIOS},
                  "capa_eolica_no_incorporada_efectiva_MW": {y: float(cap_d_capa[y]) for y in ANIOS},
                  "Huascachaca_capa": huasca, "anios_que_repiten_el_anterior": anio_sin_central,
                  "residual_PPWNDON_GW": {y: round(rc_on[y], 6) for y in HIST},
                  "generacion_PPWNDD_oficial_PJ": g_d, "generacion_PPWNDON_PJ": {y: round(g_on[y], 6) for y in HIST},
                  "CF_PPWNDON": {y: round(v, 5) for y, v in sorted(cf_on_raw.items())},
                  "CF_FP_Energy_xlsx_2013_2014": {y: fp_eol[0].get(y) for y in (2013, 2014)},
                  "CF_PPWNDON_medio_2023_2025": round(cf_on_med, 5), "generacion_existente_PPWNDON_PJ": round(g_on_ex, 6),
                  "CF_PPWNDD": {y: round(v, 5) for y, v in cf_d_raw.items()},
                  "CF_PPWNDD_medio_2023_2025": round(cf_d_med, 5), "generacion_existente_PPWNDD_PJ": round(g_d_ex, 6),
                  "salto_2022_2023": f"BEN 2.23: 21 MW (2021) -> 50 (2022) -> 71 (2023). La capa CENTRALES registra "
                                     f"Huascachaca (Elecaustro, {huasca[1]:g} MW) con inicio el {huasca[0]} y la eólica "
                                     "del BEN 2.20 de 2022 (61 GWh) es la de 2021 (62 GWh): en 2022 PPWNDON repite la "
                                     "potencia de 2021 y Huascachaca entra en 2023 (año completo en el modelo; el CF de "
                                     "2023 absorbe los nueve meses de operación)"}
    f_eol = (f"{F_BEN23} (eólica: 21 MW 2015-2021, 50 en 2022, 71 en 2023-2025); {F_BEN} (eólica); {F_CAPA} "
             f"(Huascachaca, inicio {huasca[0]}; distribuida: San Cristóbal y Baltra); capacidad construida por el modelo "
             f"oficial en 2010 ({b2010 * 1000:.2f} MW, {F_OF})")
    yrs = PROY
    F.add("K11a", "A2", "ResidualCapacity", "PPWNDON", HIST + yrs, [rc_on[y] for y in HIST] + [rc_on[2025]] * len(yrs),
          "GW", f_eol,
          f"Potencia eólica del BEN menos la distribuida (PPWNDD, de la capa CENTRALES) menos los {b2010 * 1000:.2f} MW que "
          "el modelo oficial construye en 2010 (año intacto, vida útil 31,5 años); desde 2026 sigue la de 2025. 2022 "
          f"repite 2021: los 29 MW que el BEN 2.23 suma en 2022 son de Huascachaca, que la capa registra con inicio el "
          f"{huasca[0]} y que no genera en 2022 (BEN 2.20: 61 GWh en 2022 y 62 en 2021). El oficial tenía 21,15 MW de "
          "residual y construía 46 MW ficticios en 2013-2015 por su CF plano de 0,172", dec=6)
    cf_on = {y: techo4(cf_on_raw[y]) for y in [2013, 2014] + HIST}
    F.add("K11b", "A2", "CapacityFactor", "PPWNDON", [2013, 2014] + HIST + yrs,
          [cf_on[y] for y in [2013, 2014] + HIST] + [techo4(cf_on_med)] * len(yrs), "fracción",
          f"{f_eol}; 2013-2014: generación y residual oficiales de {F_A2} (iguales a la hoja FP de Energy.xlsx: "
          f"{fp_eol[0].get(2013):.4f} y {fp_eol[0].get(2014):.4f})",
          "CF = generación / (potencia x 8.760 h x AF 0,95). A2 aplanaba el CF eólico (0,172 en todos los años, "
          "proyección Flat de A1). Desde 2026 la existente sigue con el CF medio observado de 2023-2025; el CF es por "
          "tecnología y también lo usa la capacidad nueva del condicional, cuya generación sigue acotada por su techo")
    F.add("K11c-INCA6", "A2", MININV, "PPWNDON", [2013, 2014, 2015], [0, 0, 0], "GW", f_eol,
          "0 en 2013-2015: los 16,5 MW de Villonaco ya están en la capacidad residual desde 2013 (Energy.xlsx, "
          "Proyectos_Plantas F5 y H5); la inversión mínima oficial los duplicaba", esc=ESC_INC)
    mi_con = [of["NDCCON"].v(MININV, "PPWNDON", y) for y in ANIOS]
    mi_con = [0.0 if y in (2013, 2014, 2015) else v for y, v in zip(ANIOS, mi_con)]
    F.add("K11c-CONA6", "Electrical", MININV, "PPWNDON", ANIOS, mi_con, "GW",
          f"{f_eol}; vector oficial de la hoja Electrical (NDCCON) en el datafile",
          "0 en 2013-2015 (Villonaco ya está en la residual); 2026-2034 conserva la inversión mínima del condicional "
          "(Pimo 150 MW y Ducal Membrillo 80,5 MW, Energy.xlsx Proyectos_Plantas)", esc=ESC_CON)
    F.add("K11d", "A2", LL, "PPWNDON", HIST + yrs, [0.98 * g_on[y] for y in HIST] + [0.98 * g_on_ex] * len(yrs), "PJ",
          f_eol, "Piso = 0,98 x (eólica del BEN - generación oficial de PPWNDD) en 2015-2025; desde 2026, 0,98 x "
                 "generación de la existente (potencia 2025 x CF medio 2023-2025)")
    for e in ("INCA6", "CONA6"):  # cada fila vale también para su sensibilidad de gas (INCA6G, CONA6G)
        o = of[REF[e]]
        ul_ref = {y: o.v(UL, "PPWNDON", y) for y in ANIOS}
        nuevo = {y: 1.02 * g_on[y] for y in HIST}
        nuevo.update({y: ul_ref[y] - ul_ref[2026] + 1.02 * g_on_ex for y in yrs})
        nota = ("Techo = 1,02 x (eólica del BEN - PPWNDD) en 2015-2025; desde 2026, techo de " + REF[e] +
                " - techo de " + REF[e] + " en 2026 (incremento del condicional) + 1,02 x generación de la existente")
        if e == "CONA6":
            vals = [ul_ref[y] for y in range(2010, 2015)] + [nuevo[y] for y in HIST + yrs]
            F.add("K11e-CONA6", "Electrical", UL, "PPWNDON", ANIOS, vals, "PJ",
                  f"{f_eol}; techo oficial de NDCCON en el datafile (hoja Electrical: intact hasta 2026 y 1,0 PJ en 2035)",
                  nota + "; vector completo porque la hoja Electrical interpola desde 2026", esc=ESC_CON)
        else:
            F.add("K11e-INCA6", "A2", UL, "PPWNDON", HIST + yrs, [nuevo[y] for y in HIST + yrs], "PJ", f_eol, nota,
                  esc=ESC_INC)
    yrs_d = list(range(2011, 2036))
    F.add("K11f", "A2", "ResidualCapacity", "PPWNDD", yrs_d, [d_eol[y] / 1000 for y in yrs_d], "GW",
          f"{F_CAPX} (Distribuida/Eólica) en 2011-2014; {F_CAPA} (eólica no incorporada: San Cristóbal desde el "
          "01-10-2007 y Baltra desde el 22-12-2014, 4,65 MW efectivos) en 2015-2035",
          "El oficial tiene 258,6 MW de eólica distribuida en todos los años: es el CF de 2015 (0,2586) copiado como "
          "potencia en A1 (modo Flat). 2015: la capa da 4,65 MW (Energy.xlsx, 2,25): el BEN 2.23 de 2015 (21 MW) ya es "
          "Villonaco + San Cristóbal + Baltra. 2011-2014 conservan Energy.xlsx (0 y 2,25 MW) porque el CLG pone la "
          "generación de San Cristóbal de 2010-2012 en PPWNDON. 2010 intacto", dec=5)
    cf_d = {y: techo4(cf_d_raw[y]) for y in HIST}
    F.add("K11g", "A2", "CapacityFactor", "PPWNDD", HIST + yrs,
          [cf_d[y] for y in HIST] + [techo4(cf_d_med)] * len(yrs), "fracción",
          f"Generación oficial de PPWNDD ({F_A2}) y potencia de K11f ({F_CAPA})",
          "CF = generación oficial / (potencia x 8.760 h x 0,95); desde 2026 el medio de 2023-2025")
    F.add("K11h", "A2", LL, "PPWNDD", HIST + yrs, [0.98 * g_d[y] for y in HIST] + [0.98 * g_d_ex] * len(yrs), "PJ",
          f"Generación oficial de PPWNDD ({F_A2}); {F_BEN} (eólica, total)",
          "Piso = 0,98 x generación oficial de la eólica distribuida (parte del total del BEN)")
    F.add("K11i", "A2", UL, "PPWNDD", HIST + yrs, [1.02 * g_d[y] for y in HIST] + [1.02 * g_d_ex] * len(yrs), "PJ",
          f"Generación oficial de PPWNDD ({F_A2}); {F_BEN} (eólica, total)",
          "Techo = 1,02 x generación oficial de la eólica distribuida")

    # ---------------- K12 solar existente --------------------------------------------------------------------------------
    cau_s, af_s = 31.536, vector("AvailabilityFactor", "PPPVT")[2015]
    dsol_x = capx[("Distribuida", "Solar")][0]
    d_sol = {y: dsol_x[y] * 1000 for y in range(2015, 2023)}
    d_sol[2023] = dsol_x[2023] * 1000 - float(sgda[2023])  # 2023 de Energy.xlsx ya incluye los SGDA
    d_sol.update({2024: d_sol[2023], 2025: d_sol[2023]})
    b2012 = of["NDCCON"].nueva[("PPPVT", 2012)]
    if abs(of["NDCINC"].nueva[("PPPVT", 2012)] - b2012) > 1e-9:
        raise SystemExit("K12: la capacidad solar de 2012 difiere entre NDCCON y NDCINC")
    g_pvd = {y: iguales_historia(of, UL, "PPPVD", [y])[0] for y in HIST}
    rc_pvt = {y: (pot.loc[y, "solar"] - float(sgda[y]) - d_sol[y]) / 1000 - b2012 for y in HIST}
    g_pvt = {y: ben.loc[y, "solar"] / PJ - g_pvd[y] for y in HIST}
    cf_pvt_raw = {y: g_pvt[y] / ((rc_pvt[y] + b2012) * af_s * cau_s) for y in HIST}
    # Capa CENTRALES (decisión e3): la solar del SNI entra en 2013 (3,49 MW efectivos: Paragachi, Enersol, Pastocalle y
    # Mulaló) y en 2014 (23,60 MW al cierre del año); ninguna central solar del SNI entra en 2015-2025. 2013-2014: la
    # potencia del modelo es la de la capa (residual = capa - lo construido en 2012, nunca negativa) y el CF el que pide
    # la generación oficial con esa potencia
    csol = capa["solar_SNI"].potencia_efectiva_MW
    rc1314 = {y: max(0.0, float(csol[y]) / 1000 - b2012) for y in (2013, 2014)}
    ul1314 = {y: iguales_historia(of, UL, "PPPVT", [y])[0] for y in (2013, 2014)}
    cf1314 = {y: techo4(ul1314[y] / ((rc1314[y] + b2012) * af_s * cau_s)) for y in (2013, 2014)}
    cf13 = cf1314[2013]
    cf_pvt_med = sum(cf_pvt_raw[y] for y in (2023, 2024, 2025)) / 3
    g_pvt_ex = (rc_pvt[2025] + b2012) * cf_pvt_med * af_s * cau_s
    if max(cf_pvt_raw.values()) >= 1 or max(cf1314.values()) >= 1:
        raise SystemExit("K12: CF >= 1")
    der["K12"] = {"capacidad_2012_construida_GW": b2012, "distribuida_no_SGDA_MW": d_sol,
                  "SGDA_MW": {y: float(sgda[y]) for y in HIST},
                  "residual_PPPVT_GW": {y: round(rc_pvt[y], 6) for y in HIST},
                  "capa_solar_SNI_efectiva_MW": {y: float(csol[y]) for y in range(2012, 2026)},
                  "modelo_PPPVT_2015_2025_MW": {y: round((rc_pvt[y] + b2012) * 1000, 2) for y in HIST},
                  "residual_PPPVT_2013_2014_GW": rc1314, "residual_PPPVT_CLG_2013_2014_GW":
                      {y: vector("ResidualCapacity", "PPPVT")[y] for y in (2013, 2014)}, "CF_PPPVT_2013_2014": cf1314,
                  "generacion_PPPVD_oficial_PJ": g_pvd, "generacion_PPPVT_PJ": {y: round(g_pvt[y], 6) for y in HIST},
                  "CF_PPPVT": {y: round(v, 5) for y, v in cf_pvt_raw.items()}, "CF_PPPVT_2013": cf13,
                  "CF_PPPVT_medio_2023_2025": round(cf_pvt_med, 5), "generacion_existente_PPPVT_PJ": round(g_pvt_ex, 6),
                  "nota_SGDA": "La tabla 2.20 no incluye la energía de los SGDA: con 142 MW en 2025 la solar daría un "
                               "factor de planta de 0,03; sin los SGDA, 0,15 (inferido) [verificar con ARCONEL]"}
    f_sol = (f"{F_BEN23} (solar, incluye SGDA); {F_BEN25}; {F_BEN} (solar); {F_CAPX} (Distribuida/Solar "
             f"{capx[('Distribuida', 'Solar')][1]}: no SGDA; 2023 = 43 MW - 38,38 MW de SGDA; 2024-2025 arrastran 2023 "
             f"[verificar]); capacidad construida por el modelo oficial en 2012 ({b2012 * 1000:.2f} MW, {F_OF})")
    F.add("K12a", "A2", "ResidualCapacity", "PPPVT", [2013, 2014] + HIST + yrs,
          [rc1314[2013], rc1314[2014]] + [rc_pvt[y] for y in HIST] + [rc_pvt[2025]] * len(yrs),
          "GW", f"{f_sol}; 2013-2014: {F_CAPA} (solar del SNI por fecha de inicio)",
          f"Solar del BEN - SGDA - distribuida no SGDA - {b2012 * 1000:.2f} MW construidos por el modelo en 2012 (siguen en "
          "operación): unos 24 MW de centrales con título habilitante (ARCONEL, Rendición de Cuentas 2025, tabla 9: "
          "25,16 MW fotovoltaica con TH; capa CENTRALES: 24 centrales del SNI, 24,48 MW nominales y 23,60 efectivos, todas "
          "con inicio entre el 01-02-2013 y el 12-11-2014, ninguna nueva en 2015-2025). Desde 2026 sigue la de 2025. "
          f"2013-2014: capa CENTRALES ({float(csol[2013]):.2f} y {float(csol[2014]):.2f} MW efectivos al cierre del año) "
          f"menos lo construido en 2012, sin bajar de 0: 2013 queda en {b2012 * 1000:.2f} MW, por encima de la capa, porque "
          "la generación oficial de 2012 obliga a construirlos ese año (límite conocido)", dec=6)
    F.add("K12b", "A2", "CapacityFactor", "PPPVT", [2013, 2014] + HIST + yrs,
          [cf1314[2013], cf1314[2014]] + [techo4(cf_pvt_raw[y]) for y in HIST] + [techo4(cf_pvt_med)] * len(yrs),
          "fracción", f"{f_sol}; 2013-2014: techo oficial de {F_A2} y potencia de K12a; {F_CAPA} (solar del SNI)",
          "CF = generación / (potencia x 8.760 h x 0,95). 2013: el CF oficial de relleno (0,0152) obligaba a construir "
          "42,3 MW ficticios que seguían operando hasta 2035. 2014: la mayor parte de la potencia entra entre septiembre y "
          "noviembre, así que el CF anual es bajo. Desde 2026 la existente sigue con el CF medio 2023-2025")
    F.add("K12c-INCA6", "A2", LL, "PPPVT", HIST + yrs, [0.98 * g_pvt[y] for y in HIST] + [0.98 * g_pvt_ex] * len(yrs),
          "PJ", f"{f_sol}; potencia de K12a ({F_CAPA})",
          "Piso = 0,98 x (solar del BEN - PPPVD) en 2015-2025; desde 2026, 0,98 x generación de la existente",
          esc=ESC_INC)
    F.add("K12d-INCA6", "A2", UL, "PPPVT", HIST + yrs, [1.02 * g_pvt[y] for y in HIST] + [1.02 * g_pvt_ex] * len(yrs),
          "PJ", f"{f_sol}; potencia de K12a ({F_CAPA})", "Techo = 1,02 x (solar del BEN - PPPVD) en 2015-2025; desde "
          "2026, 1,02 x generación de la "
          "existente", esc=ESC_INC)
    o = of["NDCCON"]
    for cod, par, fac in (("K12c-CONA6", LL, 0.98), ("K12d-CONA6", UL, 1.02)):
        ref = {y: o.v(par, "PPPVT", y) for y in ANIOS}
        vals = [ref[y] for y in range(2010, 2015)] + [fac * g_pvt[y] for y in HIST] + \
            [ref[y] - ref[2026] + fac * g_pvt_ex for y in yrs]
        F.add(cod, "Electrical", par, "PPPVT", ANIOS, vals, "PJ",
              f"{f_sol}; potencia de K12a ({F_CAPA}); vector oficial de NDCCON en el datafile (hoja Electrical: intact "
              "hasta 2026)",
              f"{fac} x (solar del BEN - PPPVD) en 2015-2025; desde 2026, vector de NDCCON - su valor de 2026 (incremento "
              f"del condicional, iniciativa A1) + {fac} x generación de la existente. La inversión mínima del condicional "
              "de 2025 (67 MW) se conserva por decisión del 12-10-2026: no genera en 2025 y la capa "
              "CENTRALES no registra ninguna central solar del SNI con inicio en 2025", esc=ESC_CON)
    yrs_v = list(range(2019, 2036))
    F.add("K12e", "A2", "ResidualCapacity", "PPPVD", yrs_v,
          [float(sgda[y]) / 1000 for y in range(2019, 2026)] + [float(sgda[2025]) / 1000] * len(yrs), "GW", F_BEN25,
          "Los SGDA (autoabastecimiento) van a la solar distribuida. Su energía no está en la tabla 2.20 ni en la demanda "
          "del modelo, así que su potencia no despacha energía en el balance; la generación de PPPVD sigue la oficial", dec=6)
    F.add("K12f", "A2", LL, "PPPVD", HIST, [0.98 * g_pvd[y] for y in HIST], "PJ", f"Generación oficial de PPPVD ({F_A2})",
          "Piso = 0,98 x generación oficial de la solar distribuida no SGDA (parte del total del BEN)")
    F.add("K12g", "A2", UL, "PPPVD", HIST, [1.02 * g_pvd[y] for y in HIST], "PJ", f"Generación oficial de PPPVD ({F_A2})",
          "Techo = 1,02 x generación oficial de la solar distribuida no SGDA")

    # ---------------- G01 sensibilidad de gas (decisión b): PPNGS en su piso CLG 2026-2035 -----------------------------
    ll_g, ul_g = vector(LL, "PPNGS"), vector(UL, "PPNGS")
    if {y: ll_g[y] for y in PROY} != {y: vector(LL, "PPNGS", carpeta=ENERGIA / "A2_Output_Params" / "NDCINC")[y]
                                      for y in PROY}:
        raise SystemExit("G01: el piso de PPNGS 2026-2035 difiere entre NDCCON y NDCINC")
    for e in ("NDCCON", "NDCINC"):
        for y in PROY:
            if abs(of[e].v(LL, "PPNGS", y) - ll_g[y]) > 1e-12 or abs(of[e].v(UL, "PPNGS", y) - ul_g[y]) > 1e-12:
                raise SystemExit(f"G01: el datafile de {e} no tiene el piso o el techo de A2 para PPNGS en {y}")
    der["G01"] = {"piso_CLG_PJ": {y: ll_g[y] for y in PROY}, "techo_CLG_PJ": {y: ul_g[y] for y in PROY},
                  "IAR_CLG_2026": iar_of["PPNGS"][2026], "IAR_K13a_2026": iar13["PPNGS"][2026]}
    F.add("G01-PPNGS", "A2", UL, "PPNGS", PROY, [ll_g[y] for y in PROY], "PJ",
          f"Piso de PPNGS 2026-2035 de {F_A2} (igual en NDCINC y en los datafiles CLG)",
          f"Sensibilidad de gas (decisión del 12-10-2026): techo = piso oficial ({ll_g[2026]:g} PJ, "
          f"{ll_g[2026] * PJ:.0f} GWh). Con el IAR observado (K13a: {iar13['PPNGS'][2026]} frente a "
          f"{iar_of['PPNGS'][2026]} del oficial) el gas resulta más barato y en INCA6 y CONA6 PPNGS sube a su techo oficial "
          f"({ul_g[2026]:g} PJ); aquí despacha en su piso, como en NDCINC y NDCCON. INCA6 y CONA6 son el caso central",
          esc=ESC_G)
    # ---------------- K16 piso de PPFOI 2026-2035 (decisión del 13-10-2026) -----------------------------
    # El piso CLG de PPFOI 2026-2035 (9,3785 PJ hasta 2030, 9,0 PJ en 2035: el valor de 2022 del A2 congelado, sin
    # fuente) pasa al mínimo de la historia anclada 2015-2025: el piso K13b del año de menor generación observada (2021),
    # que es la actividad de PPFOI del clon en ese año. Los pisos CLG de PPDSL y PPNGS no cambian.
    piso13_foi = {y: round(0.98 * gen13["PPFOI"][y], 4) for y in HIST}
    y16 = min(HIST, key=lambda y: piso13_foi[y])
    k16 = piso13_foi[y16]
    out_c6 = EX / "CONA6_0" / "CONA6_0_Output.csv"
    act_c6 = pd.read_csv(out_c6, usecols=["Technology", "Year", "TotalTechnologyAnnualActivity"], low_memory=False)
    act_foi = act_c6[act_c6.Technology == "PPFOI"].groupby("Year").TotalTechnologyAnnualActivity.max()
    if abs(float(act_foi[y16]) - k16) > 5e-5:
        raise SystemExit(f"K16: la actividad de PPFOI de CONA6 en {y16} ({float(act_foi[y16])}) no es el piso K13b ({k16})")
    obs_min = {t: min(HIST, key=lambda y: gen13[t][y]) for t in ("PPDSL", "PPNGS")}
    der["K16"] = {"anio_minimo": y16, "piso_PJ": k16, "generacion_observada_K13_PJ": round(gen13["PPFOI"][y16], 5),
                  "piso_K13b_2015_2025_PJ": piso13_foi,
                  "actividad_CONA6_PPFOI_PJ": {int(y): round(float(v), 5) for y, v in act_foi.loc[2015:2025].items()},
                  "piso_CLG_2026_2035_PJ": {y: vector(LL, "PPFOI")[y] for y in PROY},
                  "CONA6_0_Output_md5": md5(out_c6),
                  "contraste_pisos_sin_cambio": {
                      t: {"piso_CLG_2026_PJ": vector(LL, t)[2026], "anio_minimo_observado": obs_min[t],
                          "piso_K13b_del_minimo_PJ": round(0.98 * gen13[t][obs_min[t]], 4),
                          "generacion_observada_2023_2025_PJ": {y: round(gen13[t][y], 4) for y in (2023, 2024, 2025)}}
                      for t in ("PPDSL", "PPNGS")}}
    f_k16 = (f"{F_BEN3}; {F_BEN228}; {F_BEN} (K13b-PPFOI, año {y16}); contraste: Executables/CONA6_0/"
             f"CONA6_0_Output.csv, TotalTechnologyAnnualActivity de PPFOI en {y16} = {k16:g} PJ (md5 en "
             "correcciones_K_insumos.json, K16)")
    n_k16 = (f"Piso de despacho de PPFOI 2026-2035 = mínimo de la historia anclada 2015-2025: piso K13b de {y16} "
             f"({k16:g} PJ, 0,98 x {gen13['PPFOI'][y16]:.4f} PJ de generación observada), igual a la actividad de PPFOI "
             f"del clon en ese año. Reemplaza el piso CLG ({vector(LL, 'PPFOI')[2026]:g} PJ en 2026-2030; en NDCCON una "
             "fila de Electrical lo interpola hasta 9,0 PJ en 2035: el valor de 2022 congelado, sin fuente). Los pisos CLG "
             "de PPDSL y PPNGS no cambian. Decisión del 13-10-2026")
    F.add("K16-PPFOI", "A2", LL, "PPFOI", PROY, [k16] * len(PROY), "PJ", f_k16, n_k16, esc=ESC_INC)
    # en CONA6 y CONA6G el piso 2031-2035 lo escribe la fila de Electrical copiada de NDCCON (intact hasta 2030, hito
    # 9,0 PJ en 2035): K16 reemplaza esa fila por el vector completo, con la historia que deja A2 (CLG en 2010-2014 y
    # K13b en 2015-2025) y el piso K16 desde 2026
    ll_hist = {**{y: vector(LL, "PPFOI")[y] for y in range(2010, 2015)}, **piso13_foi}
    F.add("K16-PPFOI-CON", "Electrical", LL, "PPFOI", ANIOS, [ll_hist[y] for y in range(2010, 2026)] + [k16] * len(PROY),
          "PJ", f_k16 + f"; {F_A2} (2010-2014)", n_k16 + ". En CONA6 y CONA6G reemplaza la fila de Electrical copiada de "
          "NDCCON por el vector completo: 2010-2014 de A2 (CLG), 2015-2025 de K13b y el piso K16 desde 2026", esc=ESC_CON)

    # ---------------- RE01 respaldo de emergencia en INCA6G (supuesto de cierre de la corrida K16; ver el docstring) --
    # INCA6G no tiene solución con el tope K15b (falta potencia desde 2029 sin la expansión a gas declarada). Supuesto de
    # cierre, como la térmica arrendada que el país contrata hoy: tope de PPFOI liberado desde 2026, CapitalCost y
    # FixedCost de PPFOI por 1.000 (solo se construye lo que falta) y STOELE cerrado en 2026-2035 (es un paso sin
    # pérdidas sin correlato físico que subestima las emisiones). Los mismos cuatro valores van en las filas RESP de
    # generar_escenarios_art6.py para los estados de riesgo.
    a2_inc = ENERGIA / "A2_Output_Params" / "NDCINC"
    caro = 1000.0
    costos = {}
    for par in ("CapitalCost", "FixedCost"):
        v_c, v_i = vector(par, "PPFOI"), vector(par, "PPFOI", carpeta=a2_inc)
        if v_c != v_i or sorted(v_i) != ANIOS:
            raise SystemExit(f"RE01: {par} de PPFOI difiere entre NDCCON y NDCINC o no tiene los 26 años")
        costos[par] = v_i
    sto = vector("TotalAnnualMaxCapacity", "STOELE", carpeta=a2_inc)
    if sorted(sto) != ANIOS:
        raise SystemExit("RE01: TotalAnnualMaxCapacity de STOELE no tiene los 26 años en A2 NDCINC")
    der["RE01"] = {"escenarios": "INCA6G", "multiplicador_costo_PPFOI": caro,
                   "CapitalCost_PPFOI_CLG": {y: costos["CapitalCost"][y] for y in PROY},
                   "FixedCost_PPFOI_CLG": {y: costos["FixedCost"][y] for y in PROY},
                   "TotalAnnualMaxCapacity_STOELE_CLG": {y: sto[y] for y in PROY}}
    f_re = ("Regla de la variante de potencia mínima (Art6_calibracion/variante_suficiencia.py; "
            "calibracion_g_suficiencia_variante.csv del 12-10-2026: INCA6G sin solución con el tope K15b)")
    n_re = ("Respaldo de emergencia, rotulado «con cierre de fuel oil de emergencia» (supuesto de cierre como la térmica "
            "arrendada que el país contrata hoy, introducido en la corrida K16 del 13-10-2026 fuera del acuerdo original de "
            "respaldo solo en los estados de riesgo; se aprobó el 13-10-2026, después de la verificación "
            "independiente, conservarlo rotulado): ")
    F.add("RE01-PPFOI-tope", "A2", "TotalAnnualMaxCapacity", "PPFOI", PROY, [99999] * len(PROY), "GW", f_re,
          n_re + "tope de capacidad de PPFOI liberado desde 2026 (99999, el valor por omisión) en lugar del tope K15b",
          esc="INCA6G")
    F.add("RE01-PPFOI-capital", "A2", "CapitalCost", "PPFOI", PROY, [costos["CapitalCost"][y] * caro for y in PROY],
          "costo del modelo", f_re, n_re + "CapitalCost de PPFOI 2026-2035 por 1.000: el modelo construye solo la "
          "potencia sin la cual no cubre la demanda", esc="INCA6G")
    F.add("RE01-PPFOI-fijo", "A2", "FixedCost", "PPFOI", PROY, [costos["FixedCost"][y] * caro for y in PROY],
          "costo del modelo", f_re, n_re + "FixedCost de PPFOI 2026-2035 por 1.000", esc="INCA6G")
    F.add("RE01-STOELE", "A2", "TotalAnnualMaxCapacity", "STOELE", PROY, [0] * len(PROY), "GW", f_re,
          n_re + "STOELE cerrado en 2026-2035 (tope de capacidad cero): es un atajo sin correlato físico que el modelo "
                 "usa cuando falta potencia y que subestima las emisiones", esc="INCA6G")
    insumos["K15_md5"] = {**{p_.name: md5(p_) for p_ in (ALTAS, EXPANSION, AQUI / "arconel2025_potencia_termica.csv",
                                                          CAP_CSV)},
                          **{k: (md5(v) if Path(v).exists() else "documento local no distribuido")
                             for k, v in FUENTES_K15.items()}}
    return F.filas, insumos


def renombrar_clg(x):
    """Rótulo del 12-10-2026: «oficial» pasa a «CLG» también en las claves y textos de los insumos."""
    if isinstance(x, dict):
        return {(re.sub(r"oficiales|oficial", "CLG", k) if isinstance(k, str) else k): renombrar_clg(v)
                for k, v in x.items()}
    if isinstance(x, list):
        return [renombrar_clg(v) for v in x]
    if isinstance(x, str):
        return clg(x)
    return x


def a_texto(filas):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNAS, lineterminator="\n")
    w.writeheader()
    w.writerows(filas)
    return buf.getvalue()


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--comprobar", action="store_true")
    a = p.parse_args()
    filas, insumos = construir()
    texto = a_texto(filas)
    if a.comprobar:
        igual = SALIDA.exists() and SALIDA.read_text(encoding="utf-8") == texto
        print("correcciones_K.csv coincide" if igual else "correcciones_K.csv difiere")
        return 0 if igual else 1
    SALIDA.write_text(texto, encoding="utf-8")
    insumos = renombrar_clg(insumos)
    INSUMOS.write_text(json.dumps(insumos, indent=1, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
    pd.DataFrame(insumos["conciliacion_termica_2025"]).to_csv(CONCILIACION, index=False)
    print(f"{len(filas)} filas en {SALIDA.name}; insumos y derivaciones en {INSUMOS.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
